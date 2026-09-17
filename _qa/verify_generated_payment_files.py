import json
from pathlib import Path
from openpyxl import load_workbook

base = Path(r"D:\QLHD\_qa")

dntt = load_workbook(base / "DNTT_N15K14.xlsx", data_only=True)
ws = dntt["DNTT"]
summary_rows = []
detail_rows = []
for row in range(12, ws.max_row + 1):
    if ws.cell(row, 1).value not in (None, ""):
        summary_rows.append(row)
    elif ws.cell(row, 2).value not in (None, "") and ws.cell(row, 2).value != "Tổng cộng":
        detail_rows.append(row)
total_row = next(row for row in range(12, ws.max_row + 1) if ws.cell(row, 2).value == "Tổng cộng")
dntt_result = {
    "summary_rows": len(summary_rows),
    "detail_rows": len(detail_rows),
    "total_row": total_row,
    "labor": ws.cell(total_row, 16).value,
    "travel": ws.cell(total_row, 17).value,
    "gross": ws.cell(total_row, 18).value,
    "tax": ws.cell(total_row, 19).value,
    "net": ws.cell(total_row, 20).value,
    "first_summaries": [
        [ws.cell(r, c).value for c in (1, 2, 3, 12, 13, 14, 15, 16, 17, 18, 19, 20)]
        for r in summary_rows[:3]
    ],
}
print(json.dumps({"DNTT": dntt_result}, ensure_ascii=True, default=str))
dntt.close()

dstk = load_workbook(base / "DSTK_N15K14.xlsx", data_only=True)
ws = dstk["DSTK"]
dstk_rows = [r for r in range(5, 20) if ws.cell(r, 2).value]
print(json.dumps({"DSTK": {"staff_rows": len(dstk_rows), "net": ws["J20"].value}}, ensure_ascii=True, default=str))
dstk.close()

dnck = load_workbook(base / "DNCK_N15K14.xlsx", data_only=True)
ws = dnck["DNCK"]
dnck_rows = [r for r in range(5, 19) if ws.cell(r, 2).value]
print(json.dumps({"DNCK": {"staff_rows": len(dnck_rows), "net": ws["G19"].value}}, ensure_ascii=True, default=str))
dnck.close()
