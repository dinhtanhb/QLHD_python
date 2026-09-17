import json
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(r"D:\VietHealth\INC3b - Documents\7. Theo doi hop dong\2. HDDV\HD Can thiep\Ca nhan\2026\Nhom 15\TQT\K5")
FILES = ["DNTT_N15K5.xlsx", "DSTK_N15K5.xlsx", "DNCK_N15K5.xlsx"]


def clean(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value) if not isinstance(value, (int, float, bool)) else value


for filename in FILES:
    path = BASE / filename
    wb = load_workbook(path, data_only=False, read_only=False)
    wb_values = load_workbook(path, data_only=True, read_only=False)
    print(json.dumps({"file": filename, "sheets": wb.sheetnames}, ensure_ascii=True))
    for ws in wb.worksheets:
        wsv = wb_values[ws.title]
        info = {
            "sheet": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged": [str(r) for r in ws.merged_cells.ranges],
            "print_area": str(ws.print_area),
            "print_title_rows": ws.print_title_rows,
            "freeze": str(ws.freeze_panes) if ws.freeze_panes else None,
            "landscape": ws.page_setup.orientation,
            "fit_to_width": ws.page_setup.fitToWidth,
            "fit_to_height": ws.page_setup.fitToHeight,
            "row_breaks": [b.id for b in ws.row_breaks.brk],
            "hidden_rows": [i for i, d in ws.row_dimensions.items() if d.hidden],
            "hidden_cols": [i for i, d in ws.column_dimensions.items() if d.hidden],
        }
        print(json.dumps(info, ensure_ascii=True))
        for row in range(1, ws.max_row + 1):
            formula_cells = []
            value_cells = []
            for col in range(1, ws.max_column + 1):
                cell = ws.cell(row, col)
                cached = wsv.cell(row, col).value
                if cell.value is not None:
                    formula_cells.append(f"{cell.coordinate}={clean(cell.value)}")
                    if cell.data_type == "f":
                        value_cells.append(f"{cell.coordinate}=>{clean(cached)}")
            if formula_cells:
                payload = {"row": row, "cells": formula_cells}
                if value_cells:
                    payload["cached"] = value_cells
                print(json.dumps(payload, ensure_ascii=True))
    wb.close()
    wb_values.close()
