import json
from pathlib import Path
from openpyxl import load_workbook

base = Path(r"D:\QLHD\quanly\document_templates\thanh_toan_cong_can_thiep")
for path in sorted(base.glob("*.xlsx")):
    wb = load_workbook(path, data_only=False)
    print(json.dumps({"file": path.name, "sheets": wb.sheetnames}, ensure_ascii=True))
    for ws in wb.worksheets:
        print(json.dumps({"sheet": ws.title, "rows": ws.max_row, "cols": ws.max_column, "merged": [str(x) for x in ws.merged_cells.ranges]}, ensure_ascii=True))
        for r in range(1, ws.max_row + 1):
            vals = [f"{ws.cell(r,c).coordinate}={ws.cell(r,c).value}" for c in range(1,ws.max_column+1) if ws.cell(r,c).value is not None]
            if vals:
                print(json.dumps({"row":r,"cells":vals},ensure_ascii=True))
    wb.close()
