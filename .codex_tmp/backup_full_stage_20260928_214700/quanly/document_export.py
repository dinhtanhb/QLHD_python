from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Q
from django.utils import timezone

from .financial import FinancialConfig, calculate_payment_breakdown, normalize_travel_location
from .models import (
    ChiTietGiaHanKhoiLuong,
    ChiTietKhoiLuongHopDong,
    ChiTietPhuLucPhanCong,
    ChiTietThanhToan,
    DeXuatHopDong,
    HopDong,
    NhatKyThucHien,
    PhanCongTre,
    PhuLucHopDong,
)


TEMPLATE_ROOT = Path(settings.BASE_DIR) / "quanly" / "document_templates"
CONTRACT_TEMPLATE = TEMPLATE_ROOT / "hop_dong" / "Mau_HopDong.docx"
ANNEX_TEMPLATE = TEMPLATE_ROOT / "phu_luc" / "Mau_PhuLucPhanCong.docx"
ACCEPTANCE_TEMPLATE = TEMPLATE_ROOT / "nghiem_thu_thanh_ly" / "Mau_BBNT.docx"
LIQUIDATION_TEMPLATE = TEMPLATE_ROOT / "nghiem_thu_thanh_ly" / "Mau_TLHD.docx"
EXTENSION_TIME_TEMPLATE = TEMPLATE_ROOT / "phu_luc_gia_han" / "Mau_GHHD_ThoiGian.docx"
EXTENSION_VOLUME_TEMPLATE = TEMPLATE_ROOT / "phu_luc_gia_han" / "Mau_GHHD_KhoiLuong.docx"
PAYMENT_REQUEST_TEMPLATE = TEMPLATE_ROOT / "thanh_toan_cong_can_thiep" / "Mau_DNTT.docx"
COMMITMENT_TEMPLATE = TEMPLATE_ROOT / "thanh_toan_cong_can_thiep" / "Mau_DNCK.xlsx"
PLACEHOLDER_PATTERN = re.compile(r"\{\{[^{}]+\}\}")


def _document(path):
    """Nạp python-docx khi thực sự xuất file, không chặn Django khởi động."""
    try:
        from docx import Document
        from docx.enum.section import WD_ORIENT
        from docx.shared import Mm
    except ImportError as exc:
        raise ValidationError("Thiếu thư viện python-docx. Hãy cài dependencies trong requirements.txt.") from exc
    document = Document(str(path))
    # Chuẩn hóa ở runtime để file xuất không phụ thuộc hoàn toàn vào thiết lập
    # cũ của template. Giữ hướng ngang nếu sau này bổ sung một mẫu ngang.
    for section in document.sections:
        is_landscape = section.orientation == WD_ORIENT.LANDSCAPE
        section.page_width = Mm(297 if is_landscape else 210)
        section.page_height = Mm(210 if is_landscape else 297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
        section.header_distance = Mm(10)
        section.footer_distance = Mm(10)
    return document


def _money(value):
    return f"{Decimal(value or 0):,.0f}".replace(",", ".")


def _payment_statement_total(payment_details, current_journal_ids, current_amount):
    """Cộng các khoản đã thanh toán trước đó và ĐNTT hiện tại, không tính trùng nhật ký."""
    historical_total = sum(
        (
            Decimal(item.thanh_tien or 0)
            for item in payment_details
            if item.nhat_ky_id not in current_journal_ids
        ),
        Decimal("0"),
    )
    return historical_total + Decimal(current_amount or 0)


def _number_to_words(value):
    """Đổi số nguyên tiền thành chữ ở mức đủ dùng cho mẫu hợp đồng."""
    units = ["", "nghìn", "triệu", "tỷ"]
    digits = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
    number = int(Decimal(value or 0))
    if number == 0:
        return "Không đồng"

    def group_words(group, full=False):
        hundred, remainder = divmod(group, 100)
        ten, one = divmod(remainder, 10)
        words = []
        if hundred:
            words.extend([digits[hundred], "trăm"])
        elif full and remainder:
            words.extend(["không", "trăm"])
        if ten > 1:
            words.extend([digits[ten], "mươi"])
            if one == 1:
                words.append("mốt")
            elif one == 5:
                words.append("lăm")
            elif one:
                words.append(digits[one])
        elif ten == 1:
            words.append("mười")
            if one == 5:
                words.append("lăm")
            elif one:
                words.append(digits[one])
        elif one:
            words.append(digits[one])
        return words

    groups = []
    while number:
        number, group = divmod(number, 1000)
        groups.append(group)
    words = []
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if not group:
            continue
        words.extend(group_words(group, full=index < len(groups) - 1))
        if index:
            words.append(units[index])
    return " ".join(words).capitalize() + " đồng"


def _date_parts(value):
    value = value or date.today()
    return {
        "Ngay": str(value.day),
        "Thang": str(value.month),
        "Nam": str(value.year),
        "Full": value.strftime("%d/%m/%Y"),
    }


def _export_simple_contract_record(hop_dong, record, template_path, context):
    if not template_path.exists():
        raise ValidationError(f"Chưa có template {template_path.name}.")
    document = _document(template_path)
    _replace_document(document, context)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _contract_journals(hop_dong):
    """Lấy nhật ký trực tiếp và nhật ký lịch sử thuộc phân bổ của hợp đồng."""
    scope = Q(hop_dong=hop_dong)
    if hop_dong.de_xuat_id:
        scope |= Q(
            hop_dong__isnull=True,
            phan_cong__phan_bo_id=hop_dong.de_xuat.phan_bo_id,
        )
    return list(
        NhatKyThucHien.objects.filter(scope)
        .select_related("phan_cong__tre", "phan_cong__phan_bo", "hop_dong")
        .order_by("ngay_thuc_hien", "id")
    )


def _contract_workload(hop_dong):
    """Trả về khối lượng hợp đồng theo hai nhóm PHCN/CS từ dữ liệu đã chốt."""
    result = {
        "PHCN": {"so_tre": 0, "so_buoi": 0, "don_gia": Decimal("0"), "dinh_muc": Decimal("0")},
        "CS": {"so_tre": 0, "so_buoi": 0, "don_gia": Decimal("0"), "dinh_muc": Decimal("0")},
    }
    items = list(hop_dong.chi_tiet_khoi_luong.all())
    for item in items:
        group = PhanCongTre.service_group(item.loai_dich_vu)
        if group not in result:
            continue
        result[group] = {
            "so_tre": item.so_tre,
            "so_buoi": item.so_buoi,
            "don_gia": Decimal(item.don_gia_cong or 0),
            "dinh_muc": Decimal(item.dinh_muc_di_lai or 0),
        }
    if items or not hop_dong.de_xuat_id:
        return result

    allocation = hop_dong.de_xuat.phan_bo
    result["PHCN"] = {
        "so_tre": allocation.so_tre_phcn,
        "so_buoi": allocation.so_buoi_phcn,
        "don_gia": Decimal(hop_dong.don_gia_cong or 0),
        "dinh_muc": Decimal(hop_dong.dinh_muc_di_lai_phcn or 0),
    }
    result["CS"] = {
        "so_tre": allocation.so_tre_cs,
        "so_buoi": allocation.so_buoi_cs,
        "don_gia": Decimal(hop_dong.don_gia_cong or 0),
        "dinh_muc": Decimal(hop_dong.dinh_muc_di_lai_cs or 0),
    }
    return result


def _journal_workload_rows(hop_dong):
    """Gom bảng nghiệm thu theo trẻ + dịch vụ, kèm số buổi theo địa điểm."""
    assignments = []
    if hop_dong.de_xuat_id:
        assignments = list(
            ChiTietPhuLucPhanCong.objects.filter(
                phu_luc__hop_dong=hop_dong,
                phu_luc__loai_phu_luc="KY_1",
            ).order_by("id")
        )
        if not assignments:
            assignments = list(
                PhanCongTre.objects.filter(
                    phan_bo_id=hop_dong.de_xuat.phan_bo_id,
                    ky_phan_cong=1,
                )
                .select_related("tre")
                .order_by("id")
            )
    journals = _contract_journals(hop_dong)
    rows = {}

    def ensure(key, ma_tre, ten_tre, service, planned=0):
        if key not in rows:
            rows[key] = {
                "ma_tre": ma_tre,
                "ten_tre": ten_tre,
                "service": service,
                "planned": planned or 0,
                "actual": 0,
                "home": 0,
                "school": 0,
                "other": 0,
            }
        elif planned:
            rows[key]["planned"] = max(rows[key]["planned"], planned)
        return rows[key]

    for assignment in assignments:
        if hasattr(assignment, "ma_tre"):
            ma_tre, ten_tre = assignment.ma_tre, assignment.ten_tre
            service, planned = assignment.loai_dich_vu, assignment.so_buoi_du_kien
        else:
            ma_tre, ten_tre = assignment.tre.ma_tre, assignment.tre.ho_ten
            service, planned = assignment.loai_dich_vu, assignment.so_buoi_du_kien
        ensure((ma_tre, service), ma_tre, ten_tre, service, planned)

    for journal in journals:
        ma_tre = journal.phan_cong.tre.ma_tre
        ten_tre = journal.phan_cong.tre.ho_ten
        service = journal.phan_cong.loai_dich_vu
        row = ensure((ma_tre, service), ma_tre, ten_tre, service)
        sessions = journal.so_buoi_thuc_hien or 0
        row["actual"] += sessions
        location = normalize_travel_location(
            journal.phan_cong.hinh_thuc_ct
            or journal.dia_diem_ct
            or journal.phan_cong.dia_diem_ct
        )
        if location == "nha":
            row["home"] += sessions
        elif location == "truong":
            row["school"] += sessions
        else:
            row["other"] += sessions

    return list(rows.values())


def _payment_history(hop_dong):
    """Gom các đợt thanh toán CBCT theo tháng để đưa vào phụ lục thanh lý."""
    details = list(
        ChiTietThanhToan.objects.filter(dot_thanh_toan__hop_dong=hop_dong)
        .select_related("dot_thanh_toan", "nhat_ky")
        .order_by("dot_thanh_toan__nam", "dot_thanh_toan__thang", "id")
    )
    grouped = {}
    for detail in details:
        dot = detail.dot_thanh_toan
        row = grouped.setdefault(
            dot.pk,
            {
                "date": (dot.nam, dot.thang),
                "labor": Decimal("0"),
                "tax": Decimal("0"),
                "travel": Decimal("0"),
            },
        )
        row["labor"] += detail.tien_cong
        row["tax"] += detail.thue_tncn
        row["travel"] += detail.tien_di_lai
    result = []
    for index, row in enumerate(grouped.values(), start=1):
        net_labor = row["labor"] - row["tax"]
        result.append({
            "label": f"Lần {index} - {row['date'][1]:02d}/{row['date'][0]}",
            "labor": row["labor"],
            "tax": row["tax"],
            "net_labor": net_labor,
            "travel": row["travel"],
            "received": net_labor + row["travel"],
        })
    return result


def _record_context(hop_dong, record):
    staff = hop_dong.can_bo
    if not staff:
        raise ValidationError("Biên bản nghiệm thu/thanh lý hiện chỉ áp dụng cho hợp đồng CBCT.")
    workload = _contract_workload(hop_dong)
    acceptance = getattr(hop_dong, "nghiem_thu", None)
    acceptance_date = getattr(acceptance, "ngay_nghiem_thu", None) or getattr(record, "ngay_nghiem_thu", None)
    payment_rows = _payment_history(hop_dong)
    paid_labor = sum((row["labor"] for row in payment_rows), Decimal("0"))
    paid_tax = sum((row["tax"] for row in payment_rows), Decimal("0"))
    paid_travel = sum((row["travel"] for row in payment_rows), Decimal("0"))
    acceptance_value = getattr(acceptance, "gia_tri_nghiem_thu", None)
    if acceptance_value is None:
        acceptance_value = getattr(record, "gia_tri_nghiem_thu", None)
    if acceptance_value is None:
        acceptance_value = getattr(record, "gia_tri_thanh_ly", None)
    if acceptance_value is None:
        raise ValidationError("Hồ sơ chưa có giá trị nghiệm thu/thanh lý được xác nhận.")
    liquidation_value = getattr(record, "gia_tri_thanh_ly", None)
    if liquidation_value is None:
        liquidation_value = acceptance_value
    contract_labor = sum(
        (Decimal(item["so_tre"]) * Decimal(item["so_buoi"]) * item["don_gia"] for item in workload.values()),
        Decimal("0"),
    )
    contract_travel = sum(
        (Decimal(item["so_tre"]) * Decimal(item["so_buoi"]) * item["dinh_muc"] for item in workload.values()),
        Decimal("0"),
    )
    remaining = calculate_payment_breakdown(
        max(Decimal("0"), contract_labor - paid_labor),
        max(Decimal("0"), contract_travel - paid_travel),
    )
    context = {
        "MaSoGVMN": staff.ma_can_bo,
        "HoTenGVMN": staff.ho_ten,
        "DanhXung": "Ông/Bà",
        "SoHopDong": hop_dong.so_hop_dong,
        "NgayKy_Ngay": hop_dong.ngay_ky.day if hop_dong.ngay_ky else "",
        "NgayKy_Thang": hop_dong.ngay_ky.month if hop_dong.ngay_ky else "",
        "NgayKy_Nam": hop_dong.ngay_ky.year if hop_dong.ngay_ky else "",
        "NgayNghiemThu_Ngay": acceptance_date.day if acceptance_date else "",
        "NgayNghiemThu_Thang": acceptance_date.month if acceptance_date else "",
        "NgayNghiemThu_Nam": acceptance_date.year if acceptance_date else "",
        "NgayNghiemThu": acceptance_date.strftime("%d/%m/%Y") if acceptance_date else "",
        "NgayThanhLy_Ngay": record.ngay_thanh_ly.day if getattr(record, "ngay_thanh_ly", None) else "",
        "NgayThanhLy_Thang": record.ngay_thanh_ly.month if getattr(record, "ngay_thanh_ly", None) else "",
        "NgayThanhLy_Nam": record.ngay_thanh_ly.year if getattr(record, "ngay_thanh_ly", None) else "",
        "DiaChi": staff.dia_chi or "",
        "Cccd": staff.cccd or "",
        "NgayCapCccd": staff.ngay_cap.strftime("%d/%m/%Y") if staff.ngay_cap else "",
        "NoiCapCccd": staff.noi_cap or "",
        "DienThoai": staff.dien_thoai or "",
        "Email": staff.email or "",
        "SoTrePHCN": workload["PHCN"]["so_tre"],
        "SoBuoiPHCN": workload["PHCN"]["so_buoi"],
        "SoTreCS": workload["CS"]["so_tre"],
        "SoBuoiCS": workload["CS"]["so_buoi"],
        "SoLuotDiLaiPHCN": workload["PHCN"]["so_tre"] * workload["PHCN"]["so_buoi"],
        "SoLuotDiLaiCS": workload["CS"]["so_tre"] * workload["CS"]["so_buoi"],
        "ThanhTienPHCN": _money(workload["PHCN"]["so_tre"] * workload["PHCN"]["so_buoi"] * workload["PHCN"]["don_gia"]),
        "HoTroDiLaiPHCN": _money(workload["PHCN"]["so_tre"] * workload["PHCN"]["so_buoi"] * workload["PHCN"]["dinh_muc"]),
        "ThanhTienCS": _money(workload["CS"]["so_tre"] * workload["CS"]["so_buoi"] * workload["CS"]["don_gia"]),
        "HoTroDiLaiCS": _money(workload["CS"]["so_tre"] * workload["CS"]["so_buoi"] * workload["CS"]["dinh_muc"]),
        "GiaTriNghiemThuBangChu": _number_to_words(acceptance_value),
        "GiaTriNghiemThu": _money(acceptance_value),
        "GiaTriThanhLyBangChu": _number_to_words(liquidation_value),
        "GiaTriThanhLy": _money(liquidation_value),
        "TongCong_TienCong": _money(paid_labor),
        "TongCong_Thue": _money(paid_tax),
        "TongCong_ThanhTien": _money(paid_labor - paid_tax),
        "TongCong_DiLai": _money(paid_travel),
        "TongCong_ThucNhan": _money(paid_labor - paid_tax + paid_travel),
        "ConLai_TienCong": _money(remaining["tien_cong"]),
        "ConLai_Thue": _money(remaining["thue_tncn"]),
        "ConLai_ThanhTien": _money(remaining["tien_cong"] - remaining["thue_tncn"]),
        "ConLai_DiLai": _money(remaining["tien_di_lai"]),
        "ConLai_ThucNhan": _money(remaining["thuc_linh"]),
        "HoTenGVMN": staff.ho_ten,
        "SoTaiKhoan": staff.tai_khoan or "",
        "NganHang": staff.ngan_hang or "",
        "ChiNhanh": staff.chi_nhanh or "",
    }
    for index in range(1, 16):
        row = payment_rows[index - 1] if index <= len(payment_rows) else None
        context.update({
            f"TenLan_{index}": row["label"] if row else "",
            f"TienCong_{index}": _money(row["labor"]) if row else "",
            f"Thue_{index}": _money(row["tax"]) if row else "",
            f"ThanhTien_{index}": _money(row["net_labor"]) if row else "",
            f"DiLai_{index}": _money(row["travel"]) if row else "",
            f"ThucNhan_{index}": _money(row["received"]) if row else "",
        })
    return context


def _fill_acceptance_table(document, rows):
    for table in document.tables:
        if not table.rows:
            continue
        header = " ".join(cell.text.replace("\n", " ") for cell in table.rows[0].cells)
        required = ("Mã", "Họ tên", "Loại CT", "Hồ sơ can thiệp")
        if not all(item in header for item in required) or len(table.rows) < 2:
            continue
        template_xml = deepcopy(table.rows[1]._tr)
        _remove_row(table.rows[1])
        for index, item in enumerate(rows, start=1):
            table._tbl.append(deepcopy(template_xml))
            row = table.rows[-1]
            actual = item["actual"]
            _replace_row(row, {
                "STT": index,
                "MaTre": item["ma_tre"],
                "HoTenTre": item["ten_tre"],
                "DichVu": _service_display(item["service"]),
                "HoSo": "Đầy đủ" if actual else "Chưa có nhật ký",
                "SoBuoiTaiNha": item["home"],
                "SoBuoiTaiTruong": item["school"],
                "GhiChuTre": "" if not item["other"] else f"Khác: {item['other']} buổi",
            })
        return


def _fill_liquidation_payment_table(document, payment_count):
    for table in document.tables:
        if not table.rows:
            continue
        header = " ".join(cell.text.replace("\n", " ") for cell in table.rows[0].cells)
        header_lower = header.lower()
        if "lần thanh toán" not in header_lower or "thực nhận" not in header_lower:
            continue
        total_index = next(
            (index for index, row in enumerate(table.rows) if "Cộng" in row.cells[0].text),
            len(table.rows) - 2,
        )
        # Dòng A/B/C... trong mẫu là dòng chú giải công thức, không phải kỳ thanh toán.
        data_start = 2 if len(table.rows) > 1 and table.rows[1].cells[0].text.strip() == "A" else 1
        for index in range(total_index - 1, data_start + payment_count - 1, -1):
            _remove_row(table.rows[index])
        return


def export_acceptance_record(hop_dong, record):
    if not ACCEPTANCE_TEMPLATE.exists():
        raise ValidationError(f"Chưa có template {ACCEPTANCE_TEMPLATE.name}.")
    document = _document(ACCEPTANCE_TEMPLATE)
    _fill_acceptance_table(document, _journal_workload_rows(hop_dong))
    _replace_document(document, _record_context(hop_dong, record))
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def export_liquidation_record(hop_dong, record):
    if not LIQUIDATION_TEMPLATE.exists():
        raise ValidationError(f"Chưa có template {LIQUIDATION_TEMPLATE.name}.")
    context = _record_context(hop_dong, record)
    document = _document(LIQUIDATION_TEMPLATE)
    payment_count = sum(1 for index in range(1, 16) if context.get(f"TenLan_{index}"))
    _fill_liquidation_payment_table(document, payment_count)
    _replace_document(document, context)
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def export_journal_payment_request(journals, ky=None, thang=None, nam=None, lan_tt=None, nguoi_de_nghi=None):
    """Xuất ĐNTT Word trực tiếp từ tập nhật ký đã lọc."""
    journals = list(journals)
    if not journals:
        raise ValidationError("Không có nhật ký phù hợp để xuất ĐNTT.")
    if not PAYMENT_REQUEST_TEMPLATE.exists():
        raise ValidationError("Chưa có template ĐNTT Word.")
    first = journals[0]
    staff = first.can_bo_hieu_luc
    if not staff:
        raise ValidationError("Nhật ký chưa xác định được CBCT để xuất ĐNTT.")
    contracts = [item.hop_dong_hieu_luc for item in journals if item.hop_dong_hieu_luc]
    first_contract = contracts[0] if contracts else None
    phcn = [x for x in journals if x.phan_cong.nhom_dich_vu == "PHCN"]
    cs = [x for x in journals if x.phan_cong.nhom_dich_vu == "CS"]
    def totals(rows):
        labor = sum((Decimal(x.so_buoi_thuc_hien) * Decimal(x.don_gia_cong) for x in rows), Decimal("0"))
        travel = sum((Decimal(x.so_luot_di_lai_cbct) * Decimal(x.dinh_muc_di_lai) for x in rows), Decimal("0"))
        return sum((x.so_buoi_thuc_hien for x in rows), 0), labor, sum((x.so_luot_di_lai_cbct for x in rows), 0), travel
    phcn_sessions, phcn_labor, phcn_trips, phcn_travel = totals(phcn)
    cs_sessions, cs_labor, cs_trips, cs_travel = totals(cs)
    labor_total = phcn_labor + cs_labor
    travel_total = phcn_travel + cs_travel
    breakdown = calculate_payment_breakdown(labor_total, travel_total)
    payment_round = int(lan_tt or getattr(first, "lan_thanh_toan", 1) or 1)
    contract_values = {item.pk: item.gia_tri_hop_dong for item in contracts}
    previous_payments = list(ChiTietThanhToan.objects.filter(dot_thanh_toan__hop_dong_id__in=contract_values).select_related("dot_thanh_toan").order_by("dot_thanh_toan__nam", "dot_thanh_toan__thang", "id"))
    current_journal_ids = {item.pk for item in journals if item.pk}
    paid_total = _payment_statement_total(
        previous_payments,
        current_journal_ids,
        breakdown["tong_truoc_thue"],
    )
    contract_total = sum(contract_values.values(), Decimal("0"))
    context = {
        "HoTenGVMN": staff.ho_ten, "DiaChi": staff.dia_chi or "", "DonViCongTac": staff.don_vi.ten_don_vi if staff.don_vi_id else "",
        "SoHopDong": first_contract.so_hop_dong if len(contract_values) == 1 else ("Theo danh sách hợp đồng" if contract_values else "Chưa có HĐ"),
        "NgayKy_Ngay": first_contract.ngay_ky.day if first_contract and first_contract.ngay_ky else "", "NgayKy_Thang": first_contract.ngay_ky.month if first_contract and first_contract.ngay_ky else "", "NgayKy_Nam": first_contract.ngay_ky.year if first_contract and first_contract.ngay_ky else "",
        "ThangCanThiep": thang or (first.ngay_thuc_hien.month if first.ngay_thuc_hien else ""), "NamCanThiep": nam or (first.ngay_thuc_hien.year if first.ngay_thuc_hien else ""), "LanThanhToan": payment_round,
        "SoBuoiPHCN_Thang": phcn_sessions, "TienCongPHCN_Thang": _money(phcn_labor), "SoBuoiDiLaiPHCN_Thang": phcn_trips, "TienDiLaiPHCN_Thang": _money(phcn_travel),
        "SoBuoiCS_Thang": cs_sessions, "TienCongCS_Thang": _money(cs_labor), "SoBuoiDiLaiCS_Thang": cs_trips, "TienDiLaiCS_Thang": _money(cs_travel),
        "GiaTriHopDong": _money(sum(contract_values.values(), Decimal("0"))),
        "TongCong_Thang": _money(breakdown["tong_truoc_thue"]), "ThueTNCN": _money(breakdown["thue_tncn"]), "TongThucNhan": _money(breakdown["thuc_linh"]),
        "SoTaiKhoan": staff.tai_khoan or "", "NganHang": staff.ngan_hang or "", "ChiNhanh": staff.chi_nhanh or "",
        "STT_Lan1": "1", "NoiDung_Lan1": "Thanh toán tiền công và đi lại PHCN", "SoTien_Lan1": _money(phcn_labor + phcn_travel),
        "STT_Lan2": "2", "NoiDung_Lan2": "Thanh toán tiền công và đi lại CS", "SoTien_Lan2": _money(cs_labor + cs_travel),
        "STT_Lan3": "", "NoiDung_Lan3": "", "SoTien_Lan3": "",
        "TongSoTienDaThanhToan": _money(paid_total),
        "SoTienConLai": _money(max(Decimal("0"), contract_total - paid_total)),
    }
    for index in range(4, 16):
        context.update({f"STT_Lan{index}": "", f"NoiDung_Lan{index}": "", f"SoTien_Lan{index}": ""})
    document = _document(PAYMENT_REQUEST_TEMPLATE)
    _replace_document(document, context)
    requester_name = getattr(nguoi_de_nghi, "ho_ten", nguoi_de_nghi) or staff.ho_ten
    if len(document.tables) > 1 and document.tables[1].rows:
        requester_cell = document.tables[1].rows[0].cells[-1]
        for paragraph in requester_cell.paragraphs:
            if staff.ho_ten in paragraph.text:
                for run in paragraph.runs:
                    run.text = run.text.replace(staff.ho_ten, requester_name)
    title = f"L{payment_round}_DNTT - {staff.ma_can_bo} {staff.ho_ten}"
    sentence = f"Tôi đề nghị Quý đơn vị thanh toán phí dịch vụ lần {payment_round:02d}, chi tiết như sau:"
    for paragraph in document.paragraphs:
        if "NTT -" in paragraph.text:
            paragraph.text = title
            from docx.shared import RGBColor
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
        elif "thanh toán phí dịch vụ lần 03" in paragraph.text:
            paragraph.text = sentence
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _service_display(value):
    return dict(PhanCongTre.LOAI_DV_CHOICES).get(value, value or "")


def _replace_text_in_paragraph(paragraph, context):
    full_text = paragraph.text
    if not full_text:
        return
    replaced = full_text
    for key, value in context.items():
        replaced = replaced.replace("{{" + key + "}}", str(value if value is not None else ""))
    if replaced == full_text:
        return
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(replaced)
        return
    runs[0].text = replaced
    for run in runs[1:]:
        run.text = ""


def _iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph
        for paragraph in section.footer.paragraphs:
            yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested_table in cell.tables:
                    for row2 in nested_table.rows:
                        for cell2 in row2.cells:
                            for paragraph in cell2.paragraphs:
                                yield paragraph


def _replace_document(document, context):
    for paragraph in _iter_paragraphs(document):
        _replace_text_in_paragraph(paragraph, context)


def _clone_table_row(table, source_row):
    new_row = deepcopy(source_row._tr)
    table._tbl.append(new_row)
    return table.rows[-1]


def _replace_row(row, context):
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            _replace_text_in_paragraph(paragraph, context)


def _remove_row(row):
    row._tr.getparent().remove(row._tr)


def _one_date(rows):
    dates = {row.ngay_phan_cong for row in rows if row.ngay_phan_cong}
    if len(dates) == 1:
        return next(iter(dates)).strftime("%d/%m/%Y")
    if len(dates) > 1:
        return "Nhiều ngày"
    return ""


def _allocation_context(hop_dong):
    allocation = hop_dong.de_xuat.phan_bo
    cong = Decimal(hop_dong.don_gia_cong or FinancialConfig.DON_GIA_CONG)
    phcn_work = Decimal(allocation.so_tre_phcn) * Decimal(allocation.so_buoi_phcn) * cong
    phcn_travel = Decimal(allocation.so_tre_phcn) * Decimal(allocation.so_buoi_phcn) * Decimal(
        allocation.dinh_muc_di_lai_phcn or 0
    )
    cs_work = Decimal(allocation.so_tre_cs) * Decimal(allocation.so_buoi_cs) * cong
    cs_travel = Decimal(allocation.so_tre_cs) * Decimal(allocation.so_buoi_cs) * Decimal(
        allocation.dinh_muc_di_lai_cs or 0
    )
    contract_date = _date_parts(hop_dong.ngay_ky)
    start_date = _date_parts(hop_dong.tu_ngay)
    end_date = _date_parts(hop_dong.den_ngay)
    can_bo = hop_dong.can_bo
    total = phcn_work + phcn_travel + cs_work + cs_travel
    values = {
        "SoHopDong": hop_dong.so_hop_dong,
        "MaSoGVMN": hop_dong.can_bo.ma_can_bo,
        "HoTenGVMN": hop_dong.can_bo.ho_ten,
        "DanhXung": "Ông/Bà",
        "CCCD": can_bo.cccd or "",
        "NgayCapCCCD": can_bo.ngay_cap.strftime("%d/%m/%Y") if can_bo.ngay_cap else "",
        "NoiCapCCCD": can_bo.noi_cap or "",
        "DiaChi": can_bo.dia_chi or "",
        "DienThoai": can_bo.dien_thoai or "",
        "Email": can_bo.email or "",
        "NganHang": can_bo.ngan_hang or "",
        "SoTaiKhoan": can_bo.tai_khoan or "",
        "TongTrePHCN": allocation.so_tre_phcn,
        "TongBuoiPHCN": allocation.so_buoi_phcn,
        "ThanhTienPHCN": _money(phcn_work),
        "HoTroDiLaiPHCN": _money(phcn_travel),
        "TongTreCSXH": allocation.so_tre_cs,
        "TongBuoiCSXH": allocation.so_buoi_cs,
        "ThanhTienCSXH": _money(cs_work),
        "HoTroDiLaiCSXH": _money(cs_travel),
        "TongTien": _money(total),
        "GiaTriHopDong": _money(hop_dong.gia_tri_hop_dong),
        "GiaTriHopDongBangChu": _number_to_words(total),
        "TuNgay": start_date["Full"],
        "DenNgay": end_date["Full"],
    }
    values.update({f"NgayKy_{key}": value for key, value in contract_date.items()})
    return values


def _annex_context(hop_dong, rows, ky=1):
    values = _allocation_context(hop_dong)
    values.update(
        {
            "KyPhanCong": ky,
            "NgayPhanCong": _one_date(rows),
        }
    )
    return values


def _fill_assignment_table(document, hop_dong, rows, ky=1):
    for table in document.tables:
        if not table.rows:
            continue
        header = " ".join(cell.text.replace("\n", " ") for cell in table.rows[0].cells)
        required = ("STT", "Họ tên trẻ", "Mã trẻ", "Loại can thiệp", "Số buổi")
        if not all(item in header for item in required) or len(table.rows) < 2:
            continue
        row_template = table.rows[1]
        _remove_row(row_template)
        base_context = _annex_context(hop_dong, rows, ky)
        for index, item in enumerate(rows, start=1):
            row = _clone_table_row(table, row_template)
            _replace_row(
                row,
                {
                    **base_context,
                    "STT": index,
                    "HoTenTre": _row_value(item, "ten_tre"),
                    "MaTre": _row_value(item, "ma_tre"),
                    "DichVu": _service_display(_row_value(item, "loai_dich_vu")),
                    "SoBuoi": _row_value(item, "so_buoi_du_kien"),
                },
            )
        return


def _build_annex_document(hop_dong, rows):
    document = _document(ANNEX_TEMPLATE)
    _fill_assignment_table(document, hop_dong, rows)
    _replace_document(document, _annex_context(hop_dong, rows))
    return document


def _extension_context(extension):
    """Tạo dữ liệu thay thế cho mẫu phụ lục gia hạn thời gian/khối lượng."""
    hop_dong = extension.hop_dong
    staff = hop_dong.can_bo
    if not staff:
        raise ValidationError("Phụ lục gia hạn theo mẫu này chỉ áp dụng cho hợp đồng CBCT.")
    contract_date = _date_parts(hop_dong.ngay_ky)
    extension_date = _date_parts(extension.ngay_lap)
    values = {
        "MaSoGVMN": staff.ma_can_bo,
        "HoTenGVMN": staff.ho_ten,
        "SoHopDong": hop_dong.so_hop_dong,
        "NgayKy_Ngay": contract_date["Ngay"],
        "NgayKy_Thang": contract_date["Thang"],
        "NgayKy_Nam": contract_date["Nam"],
        "NgayGiaHan_Ngay": extension_date["Ngay"],
        "NgayGiaHan_Thang": extension_date["Thang"],
        "NgayGiaHan_Nam": extension_date["Nam"],
        "NgayHieuLuc_Ngay": extension_date["Ngay"],
        "NgayHieuLuc_Thang": extension_date["Thang"],
        "NgayHieuLuc_Nam": extension_date["Nam"],
        "TuNgay": hop_dong.tu_ngay.strftime("%d/%m/%Y") if hop_dong.tu_ngay else "",
        "DenNgayMoi": extension.den_ngay_moi.strftime("%d/%m/%Y") if extension.den_ngay_moi else "",
        "DiaChi": staff.dia_chi or "",
        "CCCD": staff.cccd or "",
        "NgayCapCCCD": staff.ngay_cap.strftime("%d/%m/%Y") if staff.ngay_cap else "",
        "NoiCapCCCD": staff.noi_cap or "",
        "DienThoai": staff.dien_thoai or "",
        "Email": staff.email or "",
        "TongTienTangThem": _money(extension.tong_tien_tang_them),
        "TongTien": _money(extension.tong_tien_moi),
        "TongTienBangChu": _number_to_words(extension.tong_tien_moi),
    }
    for item in extension.chi_tiet_gia_han_khoi_luong.all():
        is_cs = PhanCongTre.service_group(item.loai_dich_vu) == "CS"
        suffix = "CSXH" if is_cs else "PHCN"
        values.update({
            f"TongTre{suffix}": item.so_tre_cu,
            f"TongTre{suffix}Moi": item.so_tre_moi,
            f"TongBuoi{suffix}Moi": item.so_buoi_moi,
            f"ThanhTien{suffix}Moi": _money(item.so_tre_moi * item.so_buoi_moi * item.don_gia_cong),
            f"HoTroDiLai{suffix}Moi": _money(item.so_tre_moi * item.so_buoi_moi * item.dinh_muc_di_lai),
        })
    for key in (
        "TongTrePHCN", "TongTrePHCNMoi", "TongBuoiPHCNMoi", "ThanhTienPHCNMoi", "HoTroDiLaiPHCNMoi",
        "TongTreCSXH", "TongTreCSXHMoi", "TongBuoiCSXHMoi", "ThanhTienCSXHMoi", "HoTroDiLaiCSXHMoi",
    ):
        values.setdefault(key, 0 if key.startswith(("TongTre", "TongBuoi")) else "0")
    return values


def export_extension_annex(extension):
    """Xuất phụ lục gia hạn theo đúng mẫu thời gian hoặc khối lượng."""
    template = (
        EXTENSION_TIME_TEMPLATE
        if extension.loai_phu_luc == "GIA_HAN_THOI_GIAN"
        else EXTENSION_VOLUME_TEMPLATE
    )
    if not template.exists():
        raise ValidationError(f"Chưa có template {template.name}.")
    document = _document(template)
    _replace_document(document, _extension_context(extension))
    title = f"GHHD - {extension.hop_dong.can_bo.ma_can_bo} {extension.hop_dong.can_bo.ho_ten}"
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith("GHHD -"):
            paragraph.text = title
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def _append_document(target, source):
    target.add_page_break()
    target_body = target.element.body
    for child in source.element.body.iterchildren():
        if child.tag.endswith("sectPr"):
            continue
        target_body.append(deepcopy(child))


def _snapshot_rows(hop_dong):
    snapshot = (
        ChiTietPhuLucPhanCong.objects.filter(phu_luc__hop_dong=hop_dong, phu_luc__loai_phu_luc="KY_1")
        .order_by("id")
    )
    if snapshot.exists():
        return list(snapshot)
    return list(
        PhanCongTre.objects.filter(phan_bo=hop_dong.de_xuat.phan_bo, ky_phan_cong=1)
        .select_related("tre")
        .order_by("id")
    )


def _row_value(row, name):
    if hasattr(row, name):
        return getattr(row, name)
    if name == "ma_tre":
        return row.tre.ma_tre
    if name == "ten_tre":
        return row.tre.ho_ten
    return None


def export_contract_bundle(hop_dong):
    if not CONTRACT_TEMPLATE.exists() or not ANNEX_TEMPLATE.exists():
        raise ValidationError("Chưa có đủ template Word hợp đồng và Phụ lục 2.")

    rows = _snapshot_rows(hop_dong)
    contract = _document(CONTRACT_TEMPLATE)
    _fill_assignment_table(contract, hop_dong, rows)
    _replace_document(contract, {**_allocation_context(hop_dong), **_annex_context(hop_dong, rows)})

    output = BytesIO()
    contract.save(output)
    output.seek(0)
    return output


def export_assignment_annex(hop_dong, ky):
    """Xuất Phụ lục phân công riêng cho Kỳ 2 trở đi, không ghi DB."""
    if ky < 2:
        raise ValidationError("Phụ lục xuất riêng chỉ áp dụng từ Kỳ 2 trở đi.")
    if not ANNEX_TEMPLATE.exists():
        raise ValidationError("Chưa có template Word Phụ lục phân công.")

    rows = list(
        PhanCongTre.objects.filter(phan_bo=hop_dong.de_xuat.phan_bo, ky_phan_cong=ky)
        .select_related("tre")
        .order_by("id")
    )
    if not rows:
        raise ValidationError(f"Chưa có phân công trẻ cho Kỳ {ky}.")

    document = _document(ANNEX_TEMPLATE)
    _fill_assignment_table(document, hop_dong, rows, ky=ky)
    _replace_document(document, _annex_context(hop_dong, rows, ky=ky))
    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output


def create_contract_from_proposal(proposal, cleaned_data):
    """Tạo HĐ từ proposal; tuyệt đối không lấy tiền/khối lượng từ client."""
    if proposal.trang_thai != "DA_DUYET":
        raise ValidationError("Chỉ được tạo hợp đồng từ đề xuất đã duyệt.")

    allocation = proposal.phan_bo
    with transaction.atomic():
        hop_dong = HopDong.objects.create(
            de_xuat=proposal,
            can_bo=allocation.can_bo,
            nhom_hd=allocation.nhom_hd,
            so_hop_dong=cleaned_data["so_hop_dong"],
            ngay_ky=cleaned_data.get("ngay_ky"),
            tu_ngay=cleaned_data["tu_ngay"],
            den_ngay=cleaned_data["den_ngay"],
            don_gia_cong=FinancialConfig.DON_GIA_CONG,
            dinh_muc_di_lai_phcn=allocation.dinh_muc_di_lai_phcn,
            dinh_muc_di_lai_cs=allocation.dinh_muc_di_lai_cs,
            gia_tri_hop_dong=proposal.gia_tri_du_kien,
            trang_thai=cleaned_data.get("trang_thai") or "DU_THAO",
            ghi_chu=cleaned_data.get("ghi_chu"),
        )
        hop_dong.full_clean()
        hop_dong.save()

        for service, count, sessions, travel in (
            ("VLTL", allocation.so_tre_phcn, allocation.so_buoi_phcn, allocation.dinh_muc_di_lai_phcn),
            ("CSXH", allocation.so_tre_cs, allocation.so_buoi_cs, allocation.dinh_muc_di_lai_cs),
        ):
            if count or sessions:
                ChiTietKhoiLuongHopDong.objects.create(
                    hop_dong=hop_dong,
                    loai_dich_vu=service,
                    so_tre=count,
                    so_buoi=sessions,
                    don_gia_cong=hop_dong.don_gia_cong,
                    dinh_muc_di_lai=travel,
                )

        phu_luc = PhuLucHopDong.objects.create(
            hop_dong=hop_dong,
            loai_phu_luc="KY_1",
            so_phu_luc=f"PL-K1-{hop_dong.so_hop_dong}",
            ngay_lap=timezone.localdate(),
        )
        for item in allocation.danh_sach_phan_cong.select_related("tre", "phan_bo__can_bo").filter(ky_phan_cong=1):
            ChiTietPhuLucPhanCong.objects.create(
                phu_luc=phu_luc,
                source_phan_cong=item,
                ma_tre=item.tre.ma_tre,
                ten_tre=item.tre.ho_ten,
                ma_can_bo=item.phan_bo.can_bo.ma_can_bo,
                ten_can_bo=item.phan_bo.can_bo.ho_ten,
                loai_dich_vu=item.loai_dich_vu,
                dot_phan_cong=item.dot_phan_cong,
                ky_phan_cong=item.ky_phan_cong,
                ngay_phan_cong=item.ngay_phan_cong,
                so_buoi_du_kien=item.so_buoi_du_kien,
                dinh_muc_di_lai=item.dinh_muc_di_lai,
                dia_diem_ct=item.dia_diem_ct,
                hinh_thuc_ct=item.hinh_thuc_ct,
                ghi_chu=item.ghi_chu,
            )
        proposal.trang_thai = "DA_TAO_HOP_DONG"
        proposal.save(update_fields=["trang_thai", "updated_at"])
    return hop_dong
