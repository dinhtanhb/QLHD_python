import os
from pathlib import Path
from decimal import Decimal
from datetime import date
from zipfile import ZipFile

os.environ['DJANGO_SETTINGS_MODULE'] = 'qlhd.settings'
os.environ['DB_ENGINE'] = 'sqlite'
os.environ['DEBUG'] = 'True'
import django
django.setup()
from django.test.runner import DiscoverRunner
from django.test.utils import setup_test_environment, teardown_test_environment
from django.test import override_settings
from django.urls import reverse
from quanly.test_review_fixes import ReviewFixTests
from quanly.models import CauHinhThue, Tre, NhatKyThucHien
from quanly.services.payment_ledger import snapshot_journals
from quanly.payment_export import export_journal_payment_request_excel, export_journal_account_list, export_journal_commitment
from quanly.document_export import export_journal_payment_request
from openpyxl import load_workbook

setup_test_environment()
runner = DiscoverRunner(verbosity=0, interactive=False)
old = runner.setup_databases()
case = ReviewFixTests('test_01_exports_use_finalized_tax_after_configuration_changes')
case._pre_setup()
out = Path('.codex_tmp/fix14_qa')
out.mkdir(exist_ok=True)
try:
    with override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher']):
        case.setUp()
        case.contract.gia_tri_hop_dong = 10000000
        case.contract.save()
        case.journal.don_gia_cong = 6000000
        case.journal.save()
        CauHinhThue.objects.create(tu_ngay=date(2020, 1, 1), nguong_thue=5000000, ty_le=Decimal('.05'))
        voucher = case.voucher()
        rows = snapshot_journals([case.journal])
        for name, exporter in [('DNTT.xlsx', export_journal_payment_request_excel),
                               ('DSTK.xlsx', export_journal_account_list), ('DNCK.xlsx', export_journal_commitment),
                               ('DNTT.docx', export_journal_payment_request)]:
            stream = exporter(rows)
            (out / name).write_bytes(stream.getvalue())
            with ZipFile(out / name) as package:
                assert package.testzip() is None
        assert load_workbook(out / 'DNTT.xlsx')['DNTT']['S12'].value == 300000
        assert load_workbook(out / 'DSTK.xlsx')['DSTK']['J5'].value == 5700000
        assert load_workbook(out / 'DNCK.xlsx')['DNCK']['G5'].value == 5700000
        child = Tre.objects.create(ma_tre='NEW', ho_ten='Fixture', ngay_sinh=case.child.ngay_sinh, gioi_tinh='Nam')
        pages = {
            'assignment-error.html': case.client.post(reverse('sua_phan_cong', args=[case.assignment.pk]), case._assignment_payload(tre=child.pk)),
            'journal-note.html': case.client.get(reverse('sua_nhat_ky_can_thiep', args=[case.journal.pk])),
            'extensions.html': case.client.get(reverse('danh_sach_gia_han_hop_dong')),
            'journals.html': case.client.get(reverse('nhat_ky_can_thiep'), {'q': case.staff.ma_can_bo, 'ky': 1}),
        }
        for name, response in pages.items():
            assert response.status_code == 200, (name, response.status_code)
            (out / name).write_bytes(response.content)
        print('Fixture QA: 4 export packages valid, tax 300000 / net 5700000; 4 HTML routes HTTP 200.')
finally:
    case._post_teardown()
    runner.teardown_databases(old)
    teardown_test_environment()
