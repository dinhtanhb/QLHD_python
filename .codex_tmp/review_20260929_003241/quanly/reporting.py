from collections import defaultdict
from decimal import Decimal

from django.db.models import Q

from .models import ChiTietPhieuThanhToan, NhatKyThucHien


def intervention_report_queryset(*, nhom_hd_id=None, ky=None, tu_ngay=None, den_ngay=None):
    queryset = NhatKyThucHien.objects.filter(hop_dong__isnull=False).select_related(
        "hop_dong__nhom_hd",
        "phan_cong__tre",
        "phan_cong__nhom_hd",
        "phan_cong__phan_bo__nhom_hd",
        "can_bo_nguon",
        "hop_dong__can_bo",
    ).order_by("ngay_thuc_hien", "id")
    if nhom_hd_id:
        queryset = queryset.filter(
            Q(nhom_hd_nguon_id=nhom_hd_id)
            | Q(nhom_hd_nguon__isnull=True, hop_dong__nhom_hd_id=nhom_hd_id)
            | Q(
                nhom_hd_nguon__isnull=True,
                hop_dong__nhom_hd__isnull=True,
                phan_cong__nhom_hd_id=nhom_hd_id,
            )
            | Q(
                nhom_hd_nguon__isnull=True,
                hop_dong__nhom_hd__isnull=True,
                phan_cong__nhom_hd__isnull=True,
                phan_cong__phan_bo__nhom_hd_id=nhom_hd_id,
            )
        )
    if ky:
        queryset = queryset.filter(ky_can_thiep=ky)
    if tu_ngay:
        queryset = queryset.filter(ngay_thuc_hien__gte=tu_ngay)
    if den_ngay:
        queryset = queryset.filter(ngay_thuc_hien__lte=den_ngay)
    return queryset


def build_intervention_report(queryset):
    journals = list(queryset)
    paid_details = {
        detail.nhat_ky_id: detail
        for detail in ChiTietPhieuThanhToan.objects.filter(
            nhat_ky_id__in=[item.pk for item in journals],
            hoat_dong=True,
        ).select_related("phieu")
    }
    by_group = defaultdict(lambda: {
        "nhom": "Chưa xác định",
        "ky": 0,
        "so_nhat_ky": 0,
        "so_buoi": 0,
        "di_lai": 0,
        "tien_cong": Decimal("0"),
        "tien_di_lai": Decimal("0"),
        "da_thanh_toan": Decimal("0"),
    })
    by_child = defaultdict(lambda: {
        "ma_tre": "",
        "ho_ten": "",
        "dich_vu": "",
        "nhom": "Chưa xác định",
        "ky": 0,
        "so_buoi": 0,
        "di_lai": 0,
        "tien_cong": Decimal("0"),
        "tien_di_lai": Decimal("0"),
    })
    missing_group = 0
    conflict_count = 0
    for journal in journals:
        group = journal.nhom_hd_hieu_luc
        group_code = group.ma_nhom_hd if group else "?"
        group_name = f"{group_code} - {group.ten_nhom_hd}" if group else "Chưa xác định"
        group_key = (group_code, journal.ky_can_thiep)
        group_row = by_group[group_key]
        group_row["nhom"] = group_name
        group_row["ky"] = journal.ky_can_thiep
        group_row["so_nhat_ky"] += 1
        group_row["so_buoi"] += journal.so_buoi_thuc_hien or 0
        group_row["di_lai"] += journal.so_luot_di_lai_cbct or 0
        current_labor = Decimal(journal.so_buoi_thuc_hien or 0) * Decimal(journal.don_gia_cong or 0)
        current_travel = Decimal(journal.so_luot_di_lai_cbct or 0) * Decimal(journal.dinh_muc_di_lai or 0)
        group_row["tien_cong"] += current_labor
        group_row["tien_di_lai"] += current_travel
        paid = paid_details.get(journal.pk)
        if paid:
            group_row["da_thanh_toan"] += Decimal(paid.tien_cong or 0) + Decimal(paid.tien_di_lai or 0)
        if not group:
            missing_group += 1
        if journal.canh_bao_trung:
            conflict_count += 1

        child = journal.phan_cong.tre
        child_key = (child.pk, journal.phan_cong.loai_dich_vu, group_code, journal.ky_can_thiep)
        child_row = by_child[child_key]
        child_row["ma_tre"] = child.ma_tre
        child_row["ho_ten"] = child.ho_ten
        child_row["dich_vu"] = journal.phan_cong.get_loai_dich_vu_display()
        child_row["nhom"] = group_name
        child_row["ky"] = journal.ky_can_thiep
        child_row["so_buoi"] += journal.so_buoi_thuc_hien or 0
        child_row["di_lai"] += journal.so_luot_di_lai_cbct or 0
        child_row["tien_cong"] += current_labor
        child_row["tien_di_lai"] += current_travel

    group_rows = sorted(by_group.values(), key=lambda row: (row["nhom"], row["ky"]))
    child_rows = sorted(
        by_child.values(),
        key=lambda row: (row["nhom"], row["ky"], row["ma_tre"], row["dich_vu"]),
    )
    total_labor = sum((row["tien_cong"] for row in group_rows), Decimal("0"))
    total_travel = sum((row["tien_di_lai"] for row in group_rows), Decimal("0"))
    return {
        "groups": group_rows,
        "children": child_rows,
        "summary": {
            "so_nhat_ky": len(journals),
            "so_buoi": sum(row["so_buoi"] for row in group_rows),
            "di_lai": sum(row["di_lai"] for row in group_rows),
            "tien_cong": total_labor,
            "tien_di_lai": total_travel,
            "tong_gross": total_labor + total_travel,
            "da_thanh_toan": sum((row["da_thanh_toan"] for row in group_rows), Decimal("0")),
            "thieu_nhom": missing_group,
            "trung_lich": conflict_count,
        },
    }


def export_intervention_report_xlsx(report):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("Thiếu openpyxl để xuất báo cáo Excel.") from exc

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "TongHop"
    summary_sheet.append(["BÁO CÁO TỔNG HỢP CAN THIỆP"])
    summary_sheet["A1"].font = Font(bold=True, size=14)
    summary_sheet.append([])
    summary_sheet.append(["Chỉ tiêu", "Giá trị"])
    summary_sheet["A3"].font = summary_sheet["B3"].font = Font(bold=True, color="FFFFFF")
    summary_sheet["A3"].fill = summary_sheet["B3"].fill = PatternFill("solid", fgColor="0D6EFD")
    labels = (
        ("Số nhật ký", "so_nhat_ky"),
        ("Số buổi", "so_buoi"),
        ("Lượt đi lại CBCT", "di_lai"),
        ("Tiền công", "tien_cong"),
        ("Tiền đi lại", "tien_di_lai"),
        ("Tổng trước thuế", "tong_gross"),
        ("Đã thanh toán", "da_thanh_toan"),
        ("Nhật ký trùng lịch", "trung_lich"),
    )
    for label, key in labels:
        summary_sheet.append([label, report["summary"][key]])

    group_sheet = workbook.create_sheet("TheoNhomKy")
    group_sheet.append(["Nhóm HĐ", "Kỳ", "Số nhật ký", "Số buổi", "Lượt đi lại", "Tiền công", "Tiền đi lại", "Đã thanh toán"])
    child_sheet = workbook.create_sheet("TheoTreDichVu")
    child_sheet.append(["Mã trẻ", "Họ tên", "Dịch vụ", "Nhóm HĐ", "Kỳ", "Số buổi", "Lượt đi lại", "Tiền công", "Tiền đi lại"])
    for sheet in (group_sheet, child_sheet):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="198754")
    for row in report["groups"]:
        group_sheet.append([
            row["nhom"], row["ky"], row["so_nhat_ky"], row["so_buoi"], row["di_lai"],
            row["tien_cong"], row["tien_di_lai"], row["da_thanh_toan"],
        ])
    for row in report["children"]:
        child_sheet.append([
            row["ma_tre"], row["ho_ten"], row["dich_vu"], row["nhom"], row["ky"],
            row["so_buoi"], row["di_lai"], row["tien_cong"], row["tien_di_lai"],
        ])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 32)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, Decimal):
                    cell.number_format = '#,##0'
    return workbook
