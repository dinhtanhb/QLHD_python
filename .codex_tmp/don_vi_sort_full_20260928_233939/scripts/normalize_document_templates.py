"""Chuẩn hóa khổ giấy và thiết lập in cho các mẫu tài liệu của QLHD."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins


ROOT = Path(__file__).resolve().parents[1] / "quanly" / "document_templates"
LANDSCAPE_XLSX_NAMES = {
    "Mau_CBCT.xlsx",
    "Mau_NhatKyCanThiep.xlsx",
    "Mau_PhanBo.xlsx",
    "Mau_PhanCong.xlsx",
    "Mau_Tre.xlsx",
}


def normalize_docx(path: Path) -> None:
    document = Document(str(path))
    for section in document.sections:
        is_landscape = section.orientation == WD_ORIENT.LANDSCAPE
        section.page_width = Mm(297 if is_landscape else 210)
        section.page_height = Mm(210 if is_landscape else 297)
        section.top_margin = Mm(20)
        section.bottom_margin = Mm(20)
        section.left_margin = Mm(20)
        section.right_margin = Mm(20)
        section.header_distance = Mm(10)
        section.footer_distance = Mm(10)
    document.save(str(path))


def xlsx_orientation(path: Path, worksheet) -> str:
    if path.name in LANDSCAPE_XLSX_NAMES:
        return "landscape"
    if worksheet.title.casefold() in {"dntt", "dntt_ncs"}:
        return "landscape"
    if worksheet.title.casefold() == "dstk" and "thanh_toan_cong_can_thiep" in str(path.parent):
        return "landscape"
    return "portrait"


def normalize_xlsx(path: Path) -> None:
    workbook = load_workbook(path, keep_links=False)
    for name in list(workbook.defined_names):
        defined_name = workbook.defined_names[name]
        if "[" in (defined_name.attr_text or ""):
            del workbook.defined_names[name]
    for worksheet in workbook.worksheets:
        if worksheet.title.casefold() in {"data staff", "datastaff", "dia diem thuc hien"}:
            worksheet.sheet_state = "hidden"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.orientation = xlsx_orientation(path, worksheet)
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.page_setup.scale = None
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_margins = PageMargins(
            left=0.5, right=0.5, top=0.5, bottom=0.5, header=0.2, footer=0.2
        )
        name = worksheet.title.casefold()
        if name == "dntt":
            end_column, end_row = 21, 58
        elif name == "dstk":
            end_column, end_row = 12, 27
        elif name == "dnck":
            end_column, end_row = 7, 30
        elif name == "dntt_ncs":
            end_column, end_row = 12, 108
        elif name == "dstk_ncs":
            end_column, end_row = 6, 128
        elif name == "dnck_ncs":
            end_column, end_row = 6, 128
        elif "import" in str(path.parent):
            end_column, end_row = worksheet.max_column, worksheet.max_row
        else:
            continue
        worksheet.print_area = f"A1:{get_column_letter(end_column)}{end_row}"
    workbook.save(path)


def main() -> None:
    for path in sorted(ROOT.rglob("*.docx")):
        normalize_docx(path)
        print(f"DOCX {path.relative_to(ROOT)}")
    for path in sorted(ROOT.rglob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        normalize_xlsx(path)
        print(f"XLSX {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
