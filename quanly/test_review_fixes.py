"""Regressions for the fourteen findings from the project review."""

from datetime import date, time
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.contrib.messages import get_messages
from django.test import TestCase, override_settings
from django.urls import reverse

from .assignment_import import import_assignment_workbook
from .forms import GiaHanKhoiLuongForm, GiaHanThoiGianForm, _current_contract_end
from .models import (CauHinhThue, ChiTietThanhToan, ChiTietThanhToanDiLaiPhuHuynh,
                     DotThanhToan, DotThanhToanDiLaiPhuHuynh, HopDong, NhatKyThucHien,
                     NhomHD, PhanCongTre, PhuLucHopDong, Tre, DonVi)
from .payment_export import (export_journal_account_list, export_journal_commitment,
                             export_journal_payment_request_excel)
from .document_export import export_journal_payment_request
from docx import Document
from .services.payment_ledger import (NoEligiblePaymentJournals, journal_payment_breakdown,
                                      snapshot_journals, tao_phieu_thanh_toan)
from . import test_review_workflows
from .views import _parent_travel_group_journals


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class ReviewFixTests(TestCase):
    setUp = test_review_workflows.ReviewWorkflowTests.setUp
    _journal_payload = test_review_workflows.ReviewWorkflowTests._journal_payload
    _assignment_payload = test_review_workflows.ReviewWorkflowTests._assignment_payload
    _extension_payload = test_review_workflows.ReviewWorkflowTests._extension_payload

    def voucher(self):
        self.contract.trang_thai = "DA_KY"
        self.contract.save(update_fields=["trang_thai"])
        return tao_phieu_thanh_toan(self.staff, self.contract, 1, self.admin)

    def workbook(self, rows):
        output = BytesIO()
        pd.DataFrame(rows).to_excel(output, index=False)
        return SimpleUploadedFile("fixture.xlsx", output.getvalue())

    def journal_row(self, **overrides):
        return {"MaCBCT": self.staff.ma_can_bo, "MaTre": self.child.ma_tre,
                "NhomHD": self.group.ma_nhom_hd, "MaLoaiDichVu": "VLTL",
                "NgayCanThiep": "2026-09-03", "KyCanThiep": 1, "SoBuoiThucTe": 1,
                **overrides}

    def contract_row(self, **overrides):
        return {"SoHopDong": self.contract.so_hop_dong, "MaDoiTac": self.staff.ma_can_bo,
                "NhomHD": self.group.ma_nhom_hd, "TuNgay": "2026-09-01",
                "DenNgay": "2026-09-30", **overrides}

    def parent_detail(self):
        NhatKyThucHien.objects.filter(pk=self.journal.pk).update(so_luot_di_lai=1)
        self.journal.refresh_from_db()
        dot = DotThanhToanDiLaiPhuHuynh.objects.create(hop_dong=self.contract, nam=2026, thang=9)
        return ChiTietThanhToanDiLaiPhuHuynh.objects.create(
            dot_thanh_toan=dot, nhat_ky=self.journal, so_luot_di_lai=1, dinh_muc_di_lai=50000)

    def test_01_exports_use_finalized_tax_after_configuration_changes(self):
        CauHinhThue.objects.create(tu_ngay=date(2020, 1, 1), nguong_thue=5000000, ty_le=Decimal("0.05"))
        self.contract.gia_tri_hop_dong = 10000000
        self.contract.save()
        self.journal.don_gia_cong = 6000000
        self.journal.save()
        voucher = self.voucher()
        CauHinhThue.objects.all().update(ty_le=Decimal("0.2"))
        rows = snapshot_journals([self.journal])
        self.assertEqual(voucher.thue_tncn, 300000)
        self.assertEqual(load_workbook(export_journal_payment_request_excel(rows))["DNTT"]["S12"].value, 300000)
        self.assertNotIn("10%", load_workbook(export_journal_payment_request_excel(rows))["DNTT"]["S11"].value)
        self.assertEqual(load_workbook(export_journal_account_list(rows))["DSTK"]["J5"].value, 5700000)
        self.assertEqual(load_workbook(export_journal_commitment(rows))["DNCK"]["G5"].value, 5700000)
        document = Document(export_journal_payment_request(rows))
        content = " ".join(p.text for p in document.paragraphs) + " ".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("300.000", content)
        self.assertNotIn("Thuế TNCN 10%", content)
        self.assertEqual(journal_payment_breakdown(rows)["thue_tncn"], 300000)

    def test_02_exports_sum_tax_per_contract_not_per_staff(self):
        self.journal.don_gia_cong = 3000000
        self.journal.save()
        self.voucher()
        other = HopDong.objects.create(can_bo=self.staff, nhom_hd=self.group, so_hop_dong="SECOND",
                                      tu_ngay=self.contract.tu_ngay, den_ngay=self.contract.den_ngay,
                                      gia_tri_hop_dong=5000000, trang_thai="DA_KY")
        journal = NhatKyThucHien.objects.create(hop_dong=other, phan_cong=self.assignment,
                                              ngay_thuc_hien=date(2026, 9, 4), don_gia_cong=3000000, dinh_muc_di_lai=0)
        tao_phieu_thanh_toan(self.staff, other, 1)
        rows = snapshot_journals([self.journal, journal])
        self.assertEqual(journal_payment_breakdown(rows)["thue_tncn"], 0)
        self.assertEqual(load_workbook(export_journal_account_list(rows))["DSTK"]["J5"].value, 6000000)

    def test_03_contract_reimport_preserves_finalized_financials_and_identity(self):
        self.voucher()
        self.contract.trang_thai = "THANH_LY"
        self.contract.is_locked = True
        self.contract.don_gia_cong = 350000
        self.contract.save()
        response = self.client.post(reverse("import_hop_dong"), {"file_excel": self.workbook([self.contract_row()])})
        self.assertTrue(any("cập nhật 0, bỏ qua 0, giữ nguyên 1" in str(item) for item in get_messages(response.wsgi_request)))
        self.contract.refresh_from_db()
        self.assertEqual(self.contract.trang_thai, "THANH_LY")
        self.assertEqual(self.contract.don_gia_cong, 350000)
        response = self.client.post(reverse("import_hop_dong"), {"file_excel": self.workbook([
            self.contract_row(MaDoiTac=self.staff.don_vi.ma_don_vi)])})
        self.assertTrue(any("định danh" in str(item) for item in get_messages(response.wsgi_request)))
        self.contract.refresh_from_db()
        self.journal.refresh_from_db()
        self.assertEqual(self.contract.can_bo_id, self.staff.pk)
        self.assertEqual(self.journal.thanh_tien, 200000)

    def test_04_assignment_identity_is_protected_but_notes_are_editable(self):
        self.voucher()
        child = Tre.objects.create(ma_tre="OTHER", ho_ten="Khác", ngay_sinh=self.child.ngay_sinh, gioi_tinh="Nam")
        response = self.client.post(reverse("sua_phan_cong", args=[self.assignment.pk]), self._assignment_payload(tre=child.pk))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Phân công đã có dữ liệu thanh toán")
        self.assignment.refresh_from_db()
        self.assertEqual(self.assignment.tre_id, self.child.pk)
        self.assignment.ghi_chu = "ghi chú mới"
        self.assignment.save()

    def test_05_legacy_payment_blocks_new_voucher_and_reverse(self):
        dot = DotThanhToan.objects.create(hop_dong=self.contract, nam=2026, thang=9)
        legacy = ChiTietThanhToan.objects.create(dot_thanh_toan=dot, nhat_ky=self.journal, so_buoi_thanh_toan=1)
        with self.assertRaises(NoEligiblePaymentJournals):
            self.voucher()
        legacy.delete()
        self.voucher()
        with self.assertRaises(ValidationError):
            ChiTietThanhToan.objects.create(dot_thanh_toan=dot, nhat_ky=self.journal, so_buoi_thanh_toan=1)

    def test_06_parent_claim_protects_journal_and_rebuild_preserves_count(self):
        detail = self.parent_detail()
        self.journal.dia_diem_ct = "Nhà"
        with self.assertRaises(ValidationError):
            self.journal.save()
        NhatKyThucHien.recalculate_day(self.journal.ngay_thuc_hien)
        self.journal.refresh_from_db()
        self.assertEqual(self.journal.so_luot_di_lai, detail.so_luot_di_lai)

    def test_07_manual_journal_routes_use_contract_rates(self):
        self.contract.don_gia_cong = 350000
        self.contract.dinh_muc_di_lai_phcn = 70000
        self.contract.save()
        for route, args in [("them_nhat_ky_can_thiep", []), ("them_nhat_ky_thuc_hien", [self.contract.pk])]:
            response = self.client.post(reverse(route, args=args), self._journal_payload())
            self.assertEqual(response.status_code, 302)
            item = NhatKyThucHien.objects.latest("pk")
            self.assertEqual(item.don_gia_cong, 350000)
            self.assertEqual(item.dinh_muc_di_lai, 70000)

    def test_08_earlier_adjacent_session_does_not_claim_paid_trip(self):
        self.journal.gio_bat_dau, self.journal.gio_ket_thuc = time(9), time(10)
        self.journal.dia_diem_ct = "Khác"
        self.journal.save()
        self.voucher()
        earlier = NhatKyThucHien.objects.create(hop_dong=self.contract, phan_cong=self.assignment,
            ngay_thuc_hien=self.journal.ngay_thuc_hien, gio_bat_dau=time(8), gio_ket_thuc=time(9),
            dia_diem_ct="Khác", ky_can_thiep=2, don_gia_cong=200000, dinh_muc_di_lai=50000)
        NhatKyThucHien.recalculate_day(earlier.ngay_thuc_hien)
        earlier.refresh_from_db()
        self.journal.refresh_from_db()
        self.assertEqual(earlier.so_luot_di_lai_cbct, 0)
        self.assertEqual(self.journal.so_luot_di_lai_cbct, 1)
        self.assertEqual(earlier.so_luot_di_lai, 1)  # CBCT claims do not consume PH trips.

    def test_09_parent_group_uses_contract_and_validation_rolls_back_batch(self):
        other = NhomHD.objects.create(ma_nhom_hd="OTHER", ten_nhom_hd="Khác")
        self.contract.nhom_hd = other
        self.contract.save()
        NhatKyThucHien.objects.filter(pk=self.journal.pk).update(so_luot_di_lai=1)
        self.assertEqual(_parent_travel_group_journals(other, 1, 2026, 9).count(), 1)
        self.assertEqual(_parent_travel_group_journals(self.group, 1, 2026, 9).count(), 0)
        with patch("quanly.views._sync_parent_travel_group_dot", side_effect=ValidationError("changed")):
            response = self.client.post(reverse("tao_dot_thanh_toan_di_lai_phu_huynh_theo_nhom"),
                {"nhom_hd": other.pk, "ky_can_thiep": 1, "nam": 2026, "thang": 9, "ngay_de_nghi": "2026-09-30"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(DotThanhToanDiLaiPhuHuynh.objects.count(), 0)

    def test_10_unsigned_extensions_do_not_change_current_end_or_volume(self):
        PhuLucHopDong.objects.create(hop_dong=self.contract, loai_phu_luc="GIA_HAN_THOI_GIAN",
                                    den_ngay_cu=self.contract.den_ngay, den_ngay_moi=date(2026, 12, 31), is_signed=False)
        self.assertEqual(_current_contract_end(self.contract), date(2026, 9, 30))
        form = GiaHanThoiGianForm(self._extension_payload(den_ngay_moi="2026-10-31", is_signed="1"))
        self.assertTrue(form.is_valid(), form.errors)
        PhuLucHopDong.objects.create(hop_dong=self.contract, loai_phu_luc="GIA_HAN_KHOI_LUONG", tong_tien_moi=90000000)
        self.assertEqual(GiaHanKhoiLuongForm.baseline(self.contract)["phcn"][0], 1)
        self.assertEqual(GiaHanKhoiLuongForm().calculate_amounts(self.contract)["tong_gia_tri_moi"], 5000000)

    def test_11_signed_volume_extension_updates_and_delete_restores_end(self):
        payload = self._extension_payload(den_ngay_moi="2026-10-31", is_signed="1",
            so_tre_phcn_moi=2, so_buoi_phcn_moi=self.allocation.so_buoi_phcn,
            so_tre_cs_moi=self.allocation.so_tre_cs, so_buoi_cs_moi=self.allocation.so_buoi_cs)
        response = self.client.post(reverse("them_gia_han_khoi_luong"), payload)
        self.assertEqual(response.status_code, 302)
        self.contract.refresh_from_db()
        self.assertEqual(self.contract.den_ngay, date(2026, 10, 31))
        extension = self.contract.phu_luc.get()
        self.client.post(reverse("xoa_gia_han_hop_dong", args=[extension.pk]))
        self.contract.refresh_from_db()
        self.assertEqual(self.contract.den_ngay, date(2026, 9, 30))

    def test_12_assignment_import_without_staff_creates_allocation(self):
        result = import_assignment_workbook(self.workbook([{"IDChild": self.child.ma_tre,
            "NhomHD": self.group.ma_nhom_hd, "LoaiDichVu": "VLTL", "SoBuoi": 3}]))
        self.assertEqual(result.errors, [])
        self.assertEqual(result.created, 1)
        self.assertEqual(result.phan_bo_tao_moi, 1)
        self.assertIsNone(PhanCongTre.objects.latest("pk").phan_bo.can_bo_id)
        self.assertEqual(PhanCongTre.objects.latest("pk").dinh_muc_di_lai,
                         PhanCongTre.objects.latest("pk").phan_bo.dinh_muc_di_lai_phcn)

    def test_13_normal_unit_contract_import_finds_matching_assignment(self):
        self.contract.can_bo = None
        self.contract.don_vi = self.staff.don_vi
        self.contract.de_xuat = None
        self.contract.save()
        other_unit = DonVi.objects.create(ma_don_vi="DV-OTHER", ten_don_vi="Đơn vị khác", nguoi_dai_dien="Fixture")
        HopDong.objects.create(so_hop_dong="NEWER-OTHER-UNIT", don_vi=other_unit, nhom_hd=self.group,
                              tu_ngay=self.contract.tu_ngay, den_ngay=self.contract.den_ngay)
        response = self.client.post(reverse("import_nhat_ky_can_thiep"), {"file_excel": self.workbook([self.journal_row()])})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(NhatKyThucHien.objects.count(), 2)
        self.assertEqual(NhatKyThucHien.objects.latest("pk").don_gia_cong, 0)
        self.assertEqual(NhatKyThucHien.objects.latest("pk").hop_dong_id, self.contract.pk)

    def test_14_failed_import_row_rolls_back_generated_assignment_and_occurrence(self):
        before = PhanCongTre.objects.count()
        rows = [self.journal_row(MaLoaiDichVu="HDTL", SoLuotDiLaiPH=-1)]
        response = self.client.post(reverse("import_nhat_ky_can_thiep"),
                                    {"historical_mode": "1", "file_excel": self.workbook(rows)})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(PhanCongTre.objects.count(), before)
        self.assertEqual(NhatKyThucHien.objects.count(), 1)
        rows.append(self.journal_row(MaLoaiDichVu="HDTL", SoLuotDiLaiPH=1))
        self.client.post(reverse("import_nhat_ky_can_thiep"),
                         {"historical_mode": "1", "file_excel": self.workbook(rows)})
        self.assertEqual(PhanCongTre.objects.count(), before + 1)
        self.assertEqual(NhatKyThucHien.objects.count(), 2)

    def test_partial_voucher_export_does_not_expand_selection(self):
        second = NhatKyThucHien.objects.create(hop_dong=self.contract, phan_cong=self.assignment,
            ngay_thuc_hien=date(2026, 9, 3), don_gia_cong=200000, dinh_muc_di_lai=50000)
        self.voucher()
        rows = snapshot_journals([self.journal])
        with self.assertRaises(ValidationError):
            export_journal_account_list(rows)
        self.assertEqual(len(rows), 1)

    def test_failed_row_does_not_consume_existing_journal_occurrence(self):
        self.journal.can_bo_nguon = self.staff
        self.journal.nhom_hd_nguon = self.group
        self.journal.save(recalculate_travel=False)
        rows = [self.journal_row(NgayCanThiep="2026-09-02", SoLuotDiLaiPH=-1),
                self.journal_row(NgayCanThiep="2026-09-02", SoLuotDiLaiPH=0, GhiChu="updated after failure")]
        response = self.client.post(reverse("import_nhat_ky_can_thiep"), {"file_excel": self.workbook(rows)})
        self.assertEqual(NhatKyThucHien.objects.count(), 1)
        self.journal.refresh_from_db()
        self.assertEqual(self.journal.ghi_chu, "updated after failure")
        self.assertTrue(any("thêm 0, cập nhật 1, bỏ qua 1" in str(item) for item in get_messages(response.wsgi_request)))

    def test_draft_does_not_override_contract_when_signed_extension_removed(self):
        extension = PhuLucHopDong.objects.create(hop_dong=self.contract, loai_phu_luc="GIA_HAN_THOI_GIAN",
            den_ngay_cu=date(2026, 9, 30), den_ngay_moi=date(2026, 10, 31), is_signed=True)
        self.contract.den_ngay = date(2026, 10, 31)
        self.contract.save()
        PhuLucHopDong.objects.create(hop_dong=self.contract, loai_phu_luc="GIA_HAN_KHOI_LUONG",
            den_ngay_cu=date(2026, 10, 31), den_ngay_moi=date(2026, 12, 31), tong_tien_moi=90000000)
        self.client.post(reverse("xoa_gia_han_hop_dong", args=[extension.pk]))
        self.contract.refresh_from_db()
        self.assertEqual(self.contract.den_ngay, date(2026, 9, 30))
        self.assertEqual(self.contract.gia_tri_hop_dong, 5000000)

    def test_paid_zero_trip_is_not_a_trip_anchor(self):
        self.journal.gio_bat_dau, self.journal.gio_ket_thuc = time(8), time(9)
        self.journal.dia_diem_ct = "Khác"
        self.journal.save()
        NhatKyThucHien.objects.filter(pk=self.journal.pk).update(so_luot_di_lai_cbct=0, thanh_tien=200000)
        self.voucher()
        later = NhatKyThucHien.objects.create(hop_dong=self.contract, phan_cong=self.assignment,
            ngay_thuc_hien=self.journal.ngay_thuc_hien, gio_bat_dau=time(9), gio_ket_thuc=time(10),
            dia_diem_ct="Khác", ky_can_thiep=2, don_gia_cong=200000, dinh_muc_di_lai=50000)
        self.assertEqual(later.so_luot_di_lai_cbct, 1)

    def test_legacy_amount_is_included_in_contract_budget(self):
        dot = DotThanhToan.objects.create(hop_dong=self.contract, nam=2026, thang=9)
        ChiTietThanhToan.objects.create(dot_thanh_toan=dot, nhat_ky=self.journal, so_buoi_thanh_toan=1)
        self.contract.gia_tri_hop_dong = 300000
        self.contract.trang_thai = "DA_KY"
        self.contract.save()
        NhatKyThucHien.objects.create(hop_dong=self.contract, phan_cong=self.assignment, ky_can_thiep=2,
            ngay_thuc_hien=date(2026, 9, 3), don_gia_cong=200000, dinh_muc_di_lai=0)
        with self.assertRaises(ValidationError):
            tao_phieu_thanh_toan(self.staff, self.contract, 2)
