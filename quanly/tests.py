# pyright: reportAttributeAccessIssue=false, reportOptionalMemberAccess=false, reportOptionalOperand=false, reportPossiblyUnboundVariable=false, reportUnboundVariable=false, reportFunctionMemberAccess=false, reportUnusedExpression=false
from decimal import Decimal
from datetime import date, time
from io import BytesIO
from types import SimpleNamespace

import pandas as pd
from django.contrib.auth.models import Group, User
from django.contrib.messages import get_messages
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from unittest.mock import patch

from django.test import RequestFactory, SimpleTestCase, TestCase
from django.urls import reverse

from .financial import FinancialConfig, calculate_payment_breakdown, calculate_tncn, calculate_travel_flags, journal_conflict_types, normalize_travel_location
from .document_export import (
    _allocation_context,
    _date_parts,
    _number_to_words,
    _payment_statement_total,
    export_acceptance_record,
    export_liquidation_record,
)
from .forms import GiaHanKhoiLuongForm, GiaHanThoiGianForm
from .models import (
    CanBo,
    DonVi,
    HopDong,
    NhatKyThucHien,
    NhomHD,
    PhanBoChiTieu,
    PhanCongTre,
    PhuLucHopDong,
    Tre,
    NghiemThu,
    ThanhLyHopDong,
    DotThanhToan,
    CauHinhThue,
    PhieuThanhToan,
    ChiTietPhieuThanhToan,
)
from .assignment_import import import_assignment_workbook
from .services.contract_status import STATUS_TRANSITIONS, transition_hop_dong_status
from .services.payment_ledger import huy_phieu, snapshot_journals, tao_phieu_thanh_toan, xac_nhan_chi
from .parsing import parse_decimal_strict, parse_money_vnd, parse_number
from .permissions import is_admin_user
from .payment_export import (
    INTERVENTION_COMMITMENT_TEMPLATE,
    PARENT_TRAVEL_TEMPLATE_ROOT,
    _blank_zero,
    _find_template_sheet,
    _group_journal_payment_rows,
    _payment_note,
    _selection_context,
    _workbook,
    normalize_parent_travel_category,
    parent_travel_category,
)
from . import views


class FinancialRulesTests(SimpleTestCase):
    def test_service_codes_accept_vietnamese_and_ascii_variants(self):
        self.assertEqual(views.normalize_service_code("GDĐB"), "GDDB")
        self.assertEqual(views.normalize_service_code("GDDB"), "GDDB")
        self.assertEqual(views.normalize_service_code("Giáo dục đặc biệt"), "GDDB")
        self.assertEqual(views.normalize_service_code("Vật lý trị liệu"), "VLTL")

    def test_parse_decimal_accepts_excel_number_variants_without_float_math(self):
        self.assertEqual(views.parse_decimal(50000), Decimal("50000"))
        self.assertEqual(views.parse_decimal(50000.0), Decimal("50000"))
        self.assertEqual(views.parse_decimal("50,000"), Decimal("50000"))
        self.assertEqual(views.parse_decimal("50.000"), Decimal("50000"))
        self.assertEqual(views.parse_decimal("50 000"), Decimal("50000"))
        self.assertEqual(views.parse_decimal("1.234,56"), Decimal("1234.56"))

    def test_import_occurrences_do_not_collapse_repeated_assignments(self):
        first = object()
        second = object()
        cache = {}
        occurrences = {}
        loader = lambda: [first, second]

        self.assertIs(views.take_import_occurrence(cache, occurrences, "same-key", loader)[0], first)
        self.assertIs(views.take_import_occurrence(cache, occurrences, "same-key", loader)[0], second)
        self.assertIsNone(views.take_import_occurrence(cache, occurrences, "same-key", loader)[0])

    def test_journal_occurrences_preserve_repeated_source_rows(self):
        first = object()
        second = object()
        cache = {}
        occurrences = {}
        loader = lambda: [first, second]

        self.assertIs(views.take_import_occurrence(cache, occurrences, "journal-key", loader)[0], first)
        self.assertIs(views.take_import_occurrence(cache, occurrences, "journal-key", loader)[0], second)
        self.assertIsNone(views.take_import_occurrence(cache, occurrences, "journal-key", loader)[0])

    def test_journal_import_identity_separates_period_service_and_location(self):
        args = (1, 2, "GDDB", 3, 13, date(2026, 7, 25), time(7, 30), time(8, 30), "Trường")
        base = views.journal_import_identity_key(*args)
        self.assertNotEqual(base, views.journal_import_identity_key(*args[:4], 14, *args[5:]))
        self.assertNotEqual(base, views.journal_import_identity_key(1, 2, "NNTL", *args[3:]))
        self.assertNotEqual(base, views.journal_import_identity_key(*args[:-1], "Nhà"))
        self.assertEqual(base, views.journal_import_identity_key(*args[:-1], " truong "))

    def test_reference_keys_accept_geo_names_without_ids(self):
        self.assertIn("dong nai", views.normalized_reference_keys("Tỉnh Đồng Nai", ("tỉnh",)))
        self.assertIn("ha noi", views.normalized_reference_keys("Hà Nội"))
        self.assertIn("tran bien", views.normalized_reference_keys("Phường Trấn Biên", ("phường",)))

    def test_intervention_journal_routes_are_registered(self):
        self.assertEqual(reverse("import_nhat_ky_can_thiep"), "/nhat-ky-can-thiep/import/")
        self.assertEqual(reverse("them_nhat_ky_can_thiep"), "/nhat-ky-can-thiep/them/")
        self.assertEqual(reverse("sua_nhat_ky_can_thiep", args=[7]), "/nhat-ky-can-thiep/7/sua/")
        self.assertEqual(reverse("xoa_nhat_ky_can_thiep", args=[7]), "/nhat-ky-can-thiep/7/xoa/")
        self.assertEqual(reverse("xuat_dntt_excel_nhat_ky"), "/nhat-ky-can-thiep/xuat-dntt-excel/")
        self.assertEqual(reverse("xuat_nhat_ky_trung"), "/nhat-ky-can-thiep/xuat-danh-sach-trung/")

    def test_admin_edit_delete_routes_are_registered(self):
        self.assertEqual(reverse("sua_de_xuat", args=[7]), "/de-xuat/7/sua/")
        self.assertEqual(reverse("xoa_de_xuat", args=[7]), "/de-xuat/7/xoa/")
        self.assertEqual(reverse("sua_dot_thanh_toan", args=[7]), "/thanh-toan/7/sua/")
        self.assertEqual(reverse("xoa_dot_thanh_toan", args=[7]), "/thanh-toan/7/xoa/")
        self.assertEqual(reverse("sua_dot_thanh_toan_phu_huynh", args=[7]), "/thanh-toan-di-lai-phu-huynh/7/sua/")
        self.assertEqual(reverse("xoa_dot_thanh_toan_phu_huynh", args=[7]), "/thanh-toan-di-lai-phu-huynh/7/xoa/")
        self.assertEqual(reverse("sua_gia_han_hop_dong", args=[7]), "/hop-dong/gia-han/7/sua/")

    def test_contract_date_parts_keep_day_month_year_separate(self):
        parts = _date_parts(date(2026, 9, 15))
        self.assertEqual(parts["Ngay"], "15")
        self.assertEqual(parts["Thang"], "9")
        self.assertEqual(parts["Nam"], "2026")
        self.assertEqual(parts["Full"], "15/09/2026")

    def test_contract_download_filename_is_windows_safe_and_keeps_vietnamese(self):
        filename = (
            f"HDDV - {views.safe_download_component('214-26/HĐDV-VH')} - "
            f"{views.safe_download_component('ADN0116')}_"
            f"{views.safe_download_component('Phan Thị Hằng')}.docx"
        )
        self.assertEqual(filename, "HDDV - 214-26_HĐDV-VH - ADN0116_Phan Thị Hằng.docx")
        header = views.content_disposition_filename(filename)
        self.assertIn('filename="HDDV - 214-26_HDDV-VH - ADN0116_Phan Thi Hang.docx"', header)
        self.assertIn("filename*=UTF-8''HDDV%20-%20214-26_H%C4%90DV-VH%20-%20ADN0116_Phan%20Th%E1%BB%8B%20H%E1%BA%B1ng.docx", header)

    def test_payment_statement_total_adds_current_request_without_double_counting(self):
        payment_details = [
            SimpleNamespace(nhat_ky_id=10, thanh_tien=Decimal("1000000")),
            SimpleNamespace(nhat_ky_id=20, thanh_tien=Decimal("2000000")),
        ]
        total = _payment_statement_total(
            payment_details,
            current_journal_ids={20},
            current_amount=Decimal("7500000"),
        )
        self.assertEqual(total, Decimal("8500000"))

    def test_contract_context_uses_full_dates_only_for_full_date_fields(self):
        allocation = SimpleNamespace(so_tre_phcn=1, so_buoi_phcn=20, dinh_muc_di_lai_phcn=50000, so_tre_cs=0, so_buoi_cs=10, dinh_muc_di_lai_cs=50000)
        staff = SimpleNamespace(ma_can_bo="CB001", ho_ten="Người thử", cccd="", ngay_cap=None, noi_cap="", dia_chi="", dien_thoai="", email="", ngan_hang="", tai_khoan="")
        contract = SimpleNamespace(de_xuat=SimpleNamespace(phan_bo=allocation), so_hop_dong="HD001", ngay_ky=date(2026, 9, 15), tu_ngay=date(2026, 9, 15), den_ngay=date(2026, 9, 30), can_bo=staff, don_gia_cong=200000, gia_tri_hop_dong=5000000)
        context = _allocation_context(contract)
        self.assertEqual(context["NgayKy_Ngay"], "15")
        self.assertEqual(context["NgayKy_Thang"], "9")
        self.assertEqual(context["NgayKy_Nam"], "2026")
        self.assertEqual(context["TuNgay"], "15/09/2026")
    def test_tax_uses_labor_only_below_threshold(self):
        self.assertEqual(calculate_tncn(Decimal("4999999")), Decimal("0"))
        self.assertEqual(calculate_tncn(Decimal("5000000")), Decimal("500000"))

    def test_travel_is_not_taxable_or_threshold_basis(self):
        result = calculate_payment_breakdown(Decimal("4000000"), Decimal("1000000"))
        self.assertEqual(result["thue_tncn"], Decimal("0"))
        self.assertEqual(result["thuc_linh"], Decimal("5000000"))

    def test_service_groups_are_explicit(self):
        self.assertEqual(PhanCongTre.service_group("VLTL"), "PHCN")
        self.assertEqual(PhanCongTre.service_group("HDTL"), "PHCN")
        self.assertEqual(PhanCongTre.service_group("CSYT"), "CS")
        self.assertIsNone(PhanCongTre.service_group("UNKNOWN"))

    def test_schedule_conflicts_distinguish_child_and_cbct(self):
        current = SimpleNamespace(child_id="T2", cb_id="CB1", service="CSXH", date=date(2026, 7, 4), start=__import__("datetime").time(10), end=__import__("datetime").time(11))
        other = SimpleNamespace(child_id="T1", cb_id="CB1", service="VLTL", date=current.date, start=__import__("datetime").time(10), end=__import__("datetime").time(11))
        self.assertEqual(journal_conflict_types(current, [other]), ["Trùng CBCT"])
        other.child_id = "T2"
        self.assertEqual(journal_conflict_types(current, [other]), ["Trùng lịch CT"])

    def test_adjacent_sessions_at_other_share_one_cbct_trip(self):
        current = SimpleNamespace(child_id="T2", cb_id="CB1", ace="A", service="CSXH", date=date(2026, 7, 4), start=__import__("datetime").time(11), end=__import__("datetime").time(12), location="Khác", record_id=2)
        other = SimpleNamespace(child_id="T1", cb_id="CB1", ace="B", service="VLTL", date=current.date, start=__import__("datetime").time(10), end=__import__("datetime").time(11), location="Khác", record_id=1)
        self.assertEqual(journal_conflict_types(current, [other]), [])
        travel = calculate_travel_flags(current, [other])
        self.assertEqual(travel, {"so_luot_di_lai_cbct": 0, "so_luot_di_lai_ph": 1})

    def test_three_continuous_school_sessions_are_one_parent_trip(self):
        current = SimpleNamespace(
            child_id="T1", cb_id="CB1", ace="", service="VLTL",
            date=date(2026, 7, 4), start=time(10), end=time(11), location="Trường", record_id=3,
        )
        previous = [
            SimpleNamespace(child_id="T1", cb_id="CB1", ace="", service="VLTL", date=current.date, start=time(8), end=time(9), location="Trường", record_id=1),
            SimpleNamespace(child_id="T1", cb_id="CB1", ace="", service="VLTL", date=current.date, start=time(9), end=time(10), location="Trường", record_id=2),
        ]
        self.assertEqual(calculate_travel_flags(current, previous), {"so_luot_di_lai_cbct": 0, "so_luot_di_lai_ph": 0})

        first = previous[0]
        second = previous[1]
        self.assertEqual(calculate_travel_flags(first, [second, current]), {"so_luot_di_lai_cbct": 0, "so_luot_di_lai_ph": 1})
        self.assertEqual(calculate_travel_flags(second, [first, current]), {"so_luot_di_lai_cbct": 0, "so_luot_di_lai_ph": 0})

    def test_two_morning_sessions_and_one_afternoon_session_are_two_trips(self):
        current = SimpleNamespace(
            child_id="T1", cb_id="CB1", ace="", service="VLTL",
            date=date(2026, 7, 4), start=time(13), end=time(14), location="Khác", record_id=3,
        )
        previous = [
            SimpleNamespace(child_id="T1", cb_id="CB1", ace="", service="VLTL", date=current.date, start=time(8), end=time(9), location="Khác", record_id=1),
            SimpleNamespace(child_id="T1", cb_id="CB1", ace="", service="VLTL", date=current.date, start=time(9), end=time(10), location="Khác", record_id=2),
        ]
        self.assertEqual(calculate_travel_flags(current, previous), {"so_luot_di_lai_cbct": 1, "so_luot_di_lai_ph": 1})

    def test_school_session_keeps_excel_time_and_only_counts_parent_travel(self):
        self.assertEqual(views.parse_time("07:30:00"), time(7, 30))
        self.assertEqual(views.parse_time("08:30:00"), time(8, 30))
        current = SimpleNamespace(
            child_id="CBP2341", cb_id="ABP0609", ace="", service="GDDB",
            date=date(2026, 7, 25), start=time(7, 30), end=time(8, 30),
            location="Trường", record_id=None,
        )
        self.assertEqual(
            calculate_travel_flags(current, []),
            {"so_luot_di_lai_cbct": 0, "so_luot_di_lai_ph": 1},
        )

    def test_unknown_location_does_not_create_travel_allowance(self):
        current = SimpleNamespace(
            child_id="T1", cb_id="CB1", ace="", service="VLTL",
            date=date(2026, 7, 4), start=time(8), end=time(9), location="Trung tâm", record_id=1,
        )
        self.assertEqual(normalize_travel_location("Trung tâm"), "")
        self.assertEqual(calculate_travel_flags(current, []), {"so_luot_di_lai_cbct": 0, "so_luot_di_lai_ph": 0})

    def test_missing_cbct_or_child_id_does_not_match_other_journal(self):
        current = SimpleNamespace(
            child_id=None, cb_id=None, ace="", service="VLTL",
            date=date(2026, 7, 4), start=time(8), end=time(9), location="Khác", record_id=1,
        )
        other = SimpleNamespace(
            child_id=None, cb_id=None, ace="", service="CSXH",
            date=current.date, start=time(8), end=time(9), location="Khác", record_id=2,
        )
        self.assertEqual(journal_conflict_types(current, [other]), [])
        self.assertEqual(calculate_travel_flags(current, [other]), {"so_luot_di_lai_cbct": 1, "so_luot_di_lai_ph": 1})

    def test_intervention_period_choices_cover_requested_range(self):
        self.assertEqual(list(FinancialConfig.KY_CAN_THIEP_CHOICES), list(range(1, 31)))
        self.assertEqual(list(FinancialConfig.NAM_CAN_THIEP_CHOICES), list(range(2024, 2031)))

    def test_payment_excel_groups_sessions_by_staff_and_assignment(self):
        staff = SimpleNamespace(pk=1, ho_ten="CBCT A")
        contract = SimpleNamespace(pk=10, so_hop_dong="HD-01", can_bo=staff)
        child = SimpleNamespace(pk=20, ho_ten="Trẻ A")
        assignment = SimpleNamespace(
            pk=30, tre=child, loai_dich_vu="NNTL", dia_diem_ct="Nhà",
            so_buoi_du_kien=25,
        )
        journals = [
            SimpleNamespace(
                hop_dong=contract, hop_dong_id=10, phan_cong=assignment,
                phan_cong_id=30, dia_diem_ct="Nhà", so_buoi_thuc_hien=1, so_luot_di_lai_cbct=1,
                don_gia_cong=Decimal("200000"), dinh_muc_di_lai=Decimal("50000"), ghi_chu="",
            ),
            SimpleNamespace(
                hop_dong=contract, hop_dong_id=10, phan_cong=assignment,
                phan_cong_id=30, dia_diem_ct="Nhà", so_buoi_thuc_hien=1, so_luot_di_lai_cbct=0,
                don_gia_cong=Decimal("200000"), dinh_muc_di_lai=Decimal("50000"), ghi_chu="",
            ),
        ]
        grouped = _group_journal_payment_rows(journals)
        self.assertEqual(len(grouped), 1)
        self.assertEqual(len(grouped[0]["details"]), 1)
        self.assertEqual(grouped[0]["actual_home"], 2)
        self.assertEqual(grouped[0]["actual_travel"], 1)
        self.assertEqual(grouped[0]["labor"], Decimal("400000"))
        self.assertEqual(grouped[0]["travel"], Decimal("50000"))

    def test_payment_excel_hides_zero_and_import_technical_notes(self):
        self.assertIsNone(_blank_zero(0))
        self.assertIsNone(_blank_zero(Decimal("0")))
        self.assertEqual(_blank_zero(200000), 200000)
        self.assertEqual(
            _payment_note("Ghi chú nghiệp vụ\nDữ liệu lịch sử: chưa có hợp đồng"),
            "Ghi chú nghiệp vụ",
        )

    @patch("quanly.payment_export._location_for_group", return_value="Đồng Nai")
    def test_payment_excel_uses_selected_dates_not_journal_dates(self, _location):
        journal = SimpleNamespace(
            ky_can_thiep=14,
            ngay_thuc_hien=date(2026, 7, 25),
            nhom_hd_hieu_luc=SimpleNamespace(ma_nhom_hd="24"),
        )
        context = _selection_context(
            [journal],
            ky="14",
            tu_ngay="01/08/2026",
            den_ngay="31/08/2026",
        )
        self.assertEqual(context["TuNgay"], "01/08/2026")
        self.assertEqual(context["DenNgay"], "31/08/2026")

    def test_dnck_workbook_has_no_external_links_and_number_words_are_local(self):
        workbook = _workbook(INTERVENTION_COMMITMENT_TEMPLATE)
        self.assertEqual(workbook._external_links, [])
        self.assertEqual(list(workbook.defined_names), [])
        self.assertEqual(_number_to_words(41050000), "Bốn mươi mốt triệu không trăm năm mươi nghìn đồng")

    def test_parent_travel_category_prefers_cbda_note_then_assignment_method(self):
        child = SimpleNamespace(ghi_chu="Phụ huynh chuyển CBDA ký nhận", pk=1)
        assignment = SimpleNamespace(tre=child, hinh_thuc_ct="Chuyên gia")
        self.assertEqual(parent_travel_category(SimpleNamespace(phan_cong=assignment)), "CBDA")

        child.ghi_chu = "Có tài khoản ngân hàng"
        self.assertEqual(parent_travel_category(SimpleNamespace(phan_cong=assignment)), "CG")
        assignment.hinh_thuc_ct = "Cá nhân"
        self.assertEqual(parent_travel_category(SimpleNamespace(phan_cong=assignment)), "NCS")
        child.ghi_chu = "Không CBDA"
        assignment.hinh_thuc_ct = "Không chuyên gia"
        self.assertEqual(parent_travel_category(SimpleNamespace(phan_cong=assignment)), "NCS")

    def test_parent_travel_category_normalizes_only_supported_codes(self):
        self.assertEqual(normalize_parent_travel_category(" cg "), "CG")
        self.assertEqual(normalize_parent_travel_category("NCS"), "NCS")
        with self.assertRaises(Exception):
            normalize_parent_travel_category("DSTK")

    def test_parent_travel_templates_have_expected_workbook_sheets(self):
        for prefix, category in (("DNTT", "NCS"), ("DNTT", "CG"), ("DNTT", "CBDA"), ("DSTK", "NCS"), ("DSTK", "CG"), ("DNCK", "NCS"), ("DNCK", "CG")):
            path = PARENT_TRAVEL_TEMPLATE_ROOT / f"Mau_{prefix}_DiLai_{category}.xlsx"
            if not path.exists():
                if category == "CBDA" and prefix != "DNTT":
                    continue
                self.fail(f"Thiếu template {path.name}")
            workbook = _workbook(path)
            try:
                self.assertIsNotNone(_find_template_sheet(workbook, prefix))
            finally:
                workbook.close()

    def test_document_templates_use_a4_and_consistent_margins(self):
        from docx import Document
        from openpyxl import load_workbook

        root = PARENT_TRAVEL_TEMPLATE_ROOT.parent
        for path in root.rglob("*.docx"):
            document = Document(str(path))
            for section in document.sections:
                dimensions = sorted((section.page_width / 914400, section.page_height / 914400))
                self.assertAlmostEqual(dimensions[0], 8.27, places=1)
                self.assertAlmostEqual(dimensions[1], 11.69, places=1)
                for margin in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin):
                    self.assertAlmostEqual(margin / 914400, 0.79, places=1)

        for path in root.rglob("*.xlsx"):
            if path.name.startswith("~$"):
                continue
            workbook = load_workbook(path, read_only=False, data_only=False)
            for worksheet in workbook.worksheets:
                self.assertEqual(str(worksheet.page_setup.paperSize), str(worksheet.PAPERSIZE_A4))
                self.assertIn(worksheet.page_setup.orientation, {"portrait", "landscape"})
                self.assertAlmostEqual(worksheet.page_margins.left, 0.5, places=2)
                self.assertAlmostEqual(worksheet.page_margins.right, 0.5, places=2)
                self.assertAlmostEqual(worksheet.page_margins.top, 0.5, places=2)
                self.assertAlmostEqual(worksheet.page_margins.bottom, 0.5, places=2)

    def test_parent_travel_export_routes_are_registered(self):
        self.assertEqual(reverse("danh_sach_thanh_toan_di_lai_phu_huynh"), "/thanh-toan-di-lai-phu-huynh/")
        self.assertEqual(
            reverse("tao_dot_thanh_toan_di_lai_phu_huynh_theo_nhom"),
            "/thanh-toan-di-lai-phu-huynh/them/",
        )
        self.assertEqual(reverse("xuat_dntt_di_lai_phu_huynh", args=[7]), "/thanh-toan-di-lai-phu-huynh/7/xuat-dntt/")
        self.assertEqual(reverse("xuat_dstk_di_lai_phu_huynh", args=[7]), "/thanh-toan-di-lai-phu-huynh/7/xuat-dstk/")
        self.assertEqual(reverse("xuat_dnck_di_lai_phu_huynh", args=[7]), "/thanh-toan-di-lai-phu-huynh/7/xuat-dnck/")

    def test_contract_extension_routes_and_types_are_registered(self):
        self.assertEqual(reverse("danh_sach_gia_han_hop_dong"), "/hop-dong/gia-han/")
        self.assertEqual(reverse("them_gia_han_thoi_gian"), "/hop-dong/gia-han/thoi-gian/them/")
        self.assertEqual(reverse("them_gia_han_khoi_luong"), "/hop-dong/gia-han/khoi-luong/them/")
        self.assertEqual(reverse("them_gia_han_thoi_gian_theo_hop_dong", args=[7]), "/hop-dong/7/gia-han/thoi-gian/them/")
        self.assertEqual(reverse("them_gia_han_khoi_luong_theo_hop_dong", args=[7]), "/hop-dong/7/gia-han/khoi-luong/them/")
        self.assertEqual(reverse("xuat_gia_han_hop_dong", args=[7]), "/hop-dong/gia-han/7/xuat/")
        self.assertIn(("GIA_HAN_THOI_GIAN", "Gia hạn thời gian"), PhuLucHopDong.LOAI_PHU_LUC_CHOICES)
        self.assertIn(("GIA_HAN_KHOI_LUONG", "Gia hạn khối lượng"), PhuLucHopDong.LOAI_PHU_LUC_CHOICES)

    def test_contract_extension_forms_use_current_end_and_signed_dropdown(self):
        self.assertNotIn("so_phu_luc", GiaHanThoiGianForm.base_fields)
        self.assertNotIn("so_phu_luc", GiaHanKhoiLuongForm.base_fields)
        self.assertIn("den_ngay_cu", GiaHanThoiGianForm.base_fields)
        self.assertIn("den_ngay_cu", GiaHanKhoiLuongForm.base_fields)
        self.assertEqual(GiaHanThoiGianForm.base_fields["is_signed"].choices, [("0", "Chưa ký"), ("1", "Đã ký")])
        self.assertEqual(GiaHanKhoiLuongForm.base_fields["is_signed"].choices, [("0", "Chưa ký"), ("1", "Đã ký")])
        for field in ("thanh_tien_phcn", "thanh_tien_cs", "gia_tri_tang_them", "tong_gia_tri_moi"):
            self.assertIn(field, GiaHanKhoiLuongForm.base_fields)

    def test_admin_permission_helper_requires_authenticated_admin_identity(self):
        anonymous = SimpleNamespace(is_authenticated=False)
        self.assertFalse(is_admin_user(anonymous))

        staff_admin = SimpleNamespace(
            is_authenticated=True,
            is_superuser=False,
            is_staff=True,
            get_username=lambda: "admin",
        )
        self.assertTrue(is_admin_user(staff_admin))

    def test_parent_travel_export_filters_categories_and_writes_each_template(self):
        from .payment_export import (
            export_parent_travel_account_list,
            export_parent_travel_commitment,
            export_parent_travel_payment_request,
        )
        from openpyxl import load_workbook

        class Details:
            def __init__(self, rows):
                self.rows = rows

            def select_related(self, *args):
                return self

            def order_by(self, *args):
                return self.rows

        provider = SimpleNamespace(ho_ten="CBCT Test", ten_don_vi="")
        contract = SimpleNamespace(
            so_hop_dong="HD-TEST",
            doi_tac=provider,
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30),
        )

        def row(pk, note, method, parent, account, bank):
            child = SimpleNamespace(
                pk=pk, ma_tre=f"T{pk}", ho_ten=f"Trẻ {pk}", ghi_chu=note,
                ten_phu_huynh=parent, ten_tai_khoan=parent,
                tai_khoan=account, ngan_hang=bank, chi_nhanh="Chi nhánh test",
            )
            assignment = SimpleNamespace(
                tre=child, hinh_thuc_ct=method, nhom_dich_vu="PHCN",
                so_buoi_du_kien=20,
            )
            journal = SimpleNamespace(
                phan_cong=assignment, hop_dong=contract,
                can_bo_hieu_luc=None, so_buoi_thuc_hien=3,
            )
            return SimpleNamespace(
                nhat_ky=journal, dinh_muc_di_lai=Decimal("50000"),
                so_luot_di_lai=2, thanh_tien=Decimal("100000"), ghi_chu="",
            )

        rows = [
            row(1, "", "Cá nhân", "Phụ huynh NCS", "111", "Bank"),
            row(2, "", "Chuyên gia", "Phụ huynh CG", "222", "Bank"),
            row(3, "CBDA", "Cá nhân", "Phụ huynh CBDA", "", ""),
        ]
        dot = SimpleNamespace(thang=9, nam=2026, hop_dong=contract, chi_tiet=Details(rows))

        exports = [
            (export_parent_travel_payment_request, "NCS", "DNTT"),
            (export_parent_travel_payment_request, "CG", "DNTT"),
            (export_parent_travel_payment_request, "CBDA", "DNTT"),
            (export_parent_travel_account_list, "NCS", "DSTK"),
            (export_parent_travel_account_list, "CG", "DSTK"),
            (export_parent_travel_commitment, "NCS", "DNCK"),
            (export_parent_travel_commitment, "CG", "DNCK"),
        ]
        for exporter, category, sheet_prefix in exports:
            workbook = load_workbook(exporter(dot, category), keep_links=False, data_only=False)
            try:
                sheet = next(ws for ws in workbook.worksheets if ws.title.startswith(sheet_prefix))
                expected_name = {"NCS": "Phụ huynh NCS", "CG": "Phụ huynh CG", "CBDA": "Phụ huynh CBDA"}[category]
                row_number = 5 if sheet_prefix == "DSTK" else 11 if sheet_prefix == "DNTT" else 4
                column_number = 2 if sheet_prefix != "DNTT" else 3
                self.assertIn(expected_name, str(sheet.cell(row_number, column_number).value or ""))
                if sheet_prefix == "DNTT":
                    self.assertEqual(sheet["I11"].value, 3)
                    self.assertEqual(sheet["J11"].value, 2)
                    self.assertEqual(sheet["H101"].value, "=SUM(H11:H11)")
                    self.assertEqual(sheet["K101"].value, "=SUM(K11:K11)")
                elif sheet_prefix == "DSTK":
                    self.assertIn("SUBTOTAL", str(sheet["F121"].value or ""))
            finally:
                workbook.close()


class AcceptanceLiquidationExportTests(SimpleTestCase):
    def _contract(self):
        staff = SimpleNamespace(
            ma_can_bo="ABP0001",
            ho_ten="Nguyễn Văn Test",
            dia_chi="Đồng Nai",
            cccd="012345678901",
            ngay_cap=date(2020, 1, 1),
            noi_cap="Cục CSQLHC",
            dien_thoai="0900000000",
            email="test@example.com",
            tai_khoan="123456",
            ngan_hang="Ngân hàng Test",
            chi_nhanh="Đồng Nai",
        )
        return SimpleNamespace(
            can_bo=staff,
            de_xuat_id=None,
            so_hop_dong="01-26/HĐDV-VH",
            ngay_ky=date(2026, 1, 1),
            gia_tri_hop_dong=Decimal("1000000"),
            chi_tiet_khoi_luong=SimpleNamespace(all=lambda: []),
        )

    def test_acceptance_export_fills_workload_table(self):
        from . import document_export

        contract = self._contract()
        record = SimpleNamespace(
            ngay_nghiem_thu=date(2026, 9, 28),
            gia_tri_nghiem_thu=Decimal("1000000"),
        )
        rows = [{
            "ma_tre": "CBP0001",
            "ten_tre": "Trẻ Test",
            "service": "GDDB",
            "planned": 10,
            "actual": 10,
            "home": 4,
            "school": 6,
            "other": 0,
        }]
        workload = {
            "PHCN": {"so_tre": 1, "so_buoi": 10, "don_gia": Decimal("200000"), "dinh_muc": Decimal("50000")},
            "CS": {"so_tre": 0, "so_buoi": 0, "don_gia": Decimal("200000"), "dinh_muc": Decimal("50000")},
        }
        with patch.object(document_export, "_journal_workload_rows", return_value=rows), patch.object(
            document_export, "_contract_workload", return_value=workload
        ), patch.object(document_export, "_contract_journals", return_value=[]), patch.object(
            document_export, "_payment_history", return_value=[]
        ):
            output = export_acceptance_record(contract, record)

        from docx import Document

        document = Document(output)
        table = document.tables[1]
        self.assertEqual(len(table.rows), 2)
        self.assertIn("CBP0001", table.rows[1].cells[1].text)
        self.assertNotIn("{{", "\n".join(paragraph.text for paragraph in document.paragraphs))

    def test_liquidation_export_removes_unused_payment_rows(self):
        from . import document_export

        contract = self._contract()
        record = SimpleNamespace(
            ngay_thanh_ly=date(2026, 9, 28),
            gia_tri_thanh_ly=Decimal("1000000"),
        )
        workload = {
            "PHCN": {"so_tre": 1, "so_buoi": 10, "don_gia": Decimal("200000"), "dinh_muc": Decimal("50000")},
            "CS": {"so_tre": 0, "so_buoi": 0, "don_gia": Decimal("200000"), "dinh_muc": Decimal("50000")},
        }
        with patch.object(document_export, "_contract_workload", return_value=workload), patch.object(
            document_export, "_contract_journals", return_value=[]
        ), patch.object(document_export, "_payment_history", return_value=[]):
            output = export_liquidation_record(contract, record)

        from docx import Document

        document = Document(output)
        payment_table = document.tables[1]
        self.assertEqual(len(payment_table.rows), 4)
        self.assertNotIn("{{", "\n".join(paragraph.text for paragraph in document.paragraphs))


class DashboardViewTests(SimpleTestCase):
    def test_homepage_returns_response_after_login(self):
        class EmptyQuerySet:
            def __iter__(self):
                return iter(())

            def count(self):
                return 0

            def aggregate(self, **kwargs):
                return {key: 0 for key in kwargs}

            def filter(self, **kwargs):
                return self

        empty = EmptyQuerySet()
        expected = object()
        with (
            patch.object(views.PhanBoChiTieu.objects, "all", return_value=empty),
            patch.object(views.PhanCongTre.objects, "filter", return_value=empty),
            patch.object(views.HopDong.objects, "all", return_value=empty),
            patch.object(views.Tre.objects, "filter", return_value=empty),
            patch.object(views.CanBo.objects, "filter", return_value=empty),
            patch.object(views.DonVi.objects, "filter", return_value=empty),
            patch.object(views.NhomHD.objects, "filter", return_value=empty),
            patch.object(views.DeXuatHopDong.objects, "count", return_value=0),
            patch.object(views, "render", return_value=expected),
        ):
            response = views.trang_chu.__wrapped__(RequestFactory().get("/"))
        self.assertIs(response, expected)


class DatabaseRegressionTests(TestCase):
    """DB-backed regression coverage for the P0 import and contract lifecycle rules."""

    def setUp(self):
        self.admin_group = Group.objects.create(name="Admin")
        self.admin = User.objects.create_user(username="p0-admin", password="secret")
        self.admin.groups.add(self.admin_group)
        self.viewer = User.objects.create_user(username="p0-viewer", password="secret")
        self.don_vi = DonVi.objects.create(ma_don_vi="DV-TEST", ten_don_vi="Đơn vị test", nguoi_dai_dien="Người test")
        self.cb1 = CanBo.objects.create(ma_can_bo="ABP0001", ho_ten="Cán bộ Một", don_vi=self.don_vi)
        self.cb2 = CanBo.objects.create(ma_can_bo="ABP0002", ho_ten="Cán bộ Hai", don_vi=self.don_vi)
        self.group = NhomHD.objects.create(ma_nhom_hd="1", ten_nhom_hd="Nhóm HĐ 1")
        self.child = Tre.objects.create(ma_tre="CBP0001", ho_ten="Trẻ Một", ngay_sinh=date(2015, 1, 1), gioi_tinh="Nam")
        self.allocation = PhanBoChiTieu.objects.create(can_bo=self.cb1, nhom_hd=self.group, so_tre_phcn=2, so_buoi_phcn=20)

    def _workbook(self, rows):
        stream = BytesIO()
        pd.DataFrame(rows).to_excel(stream, index=False)
        return SimpleUploadedFile("phan_cong_test.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def _base_row(self, **overrides):
        row = {
            "Mã CB": "ABP0001",
            "IDChild": "CBP0001",
            "Nhóm HĐ": "1",
            "Loại dịch vụ": "VLTL",
            "Số buổi dự kiến": 20,
            "Đợt phân công": 1,
            "Kỳ phân công": 1,
            "Ngày phân công": "2026-09-01",
        }
        row.update(overrides)
        return row

    def test_catalog_pages_expose_global_kpis(self):
        self.client.force_login(self.admin)
        expected_keys = {
            "danh_sach_tre": {"total", "dong_nai", "binh_phuoc", "nam", "nu", "other_gender"},
            "danh_sach_can_bo": {"total", "active", "dong_nai", "binh_phuoc", "nam", "nu", "other_gender"},
            "danh_sach_nhom_hd": {"total", "active", "with_allocation", "with_contract", "without_contract"},
            "danh_sach_don_vi": {"total", "active", "with_staff", "with_contract", "without_contract"},
        }
        for url_name, keys in expected_keys.items():
            with self.subTest(url_name=url_name):
                response = self.client.get(reverse(url_name))
                self.assertEqual(response.status_code, 200)
                self.assertTrue(keys.issubset(response.context["kpi"].keys()))

        tre_kpi = self.client.get(reverse("danh_sach_tre")).context["kpi"]
        self.assertEqual(
            tre_kpi,
            {"total": 1, "dong_nai": 0, "binh_phuoc": 1, "nam": 1, "nu": 0, "other_gender": 0},
        )
        cb_kpi = self.client.get(reverse("danh_sach_can_bo")).context["kpi"]
        self.assertEqual(
            cb_kpi,
            {"total": 2, "active": 2, "dong_nai": 0, "binh_phuoc": 2, "nam": 0, "nu": 0, "other_gender": 2},
        )
        nhom_kpi = self.client.get(reverse("danh_sach_nhom_hd")).context["kpi"]
        self.assertEqual(nhom_kpi, {"total": 1, "active": 1, "with_allocation": 1, "with_contract": 0, "without_contract": 1})
        don_vi_kpi = self.client.get(reverse("danh_sach_don_vi")).context["kpi"]
        self.assertEqual(don_vi_kpi, {"total": 1, "active": 1, "with_staff": 1, "with_contract": 0, "without_contract": 1})

    def test_intervention_report_page_and_excel_export(self):
        assignment = PhanCongTre.objects.create(
            phan_bo=self.allocation,
            can_bo_nguon=self.cb1,
            nhom_hd=self.group,
            tre=self.child,
            loai_dich_vu="VLTL",
            so_buoi_du_kien=20,
            dinh_muc_di_lai=Decimal("50000"),
        )
        contract = HopDong.objects.create(
            can_bo=self.cb1,
            nhom_hd=self.group,
            so_hop_dong="HD-REPORT-001",
            ngay_ky=date(2026, 9, 1),
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30),
            gia_tri_hop_dong=Decimal("10000000"),
            trang_thai="DA_KY",
        )
        NhatKyThucHien.objects.create(
            hop_dong=contract,
            phan_cong=assignment,
            ngay_thuc_hien=date(2026, 9, 10),
            ky_can_thiep=1,
            so_buoi_thuc_hien=2,
            don_gia_cong=Decimal("200000"),
            dinh_muc_di_lai=Decimal("50000"),
        )
        self.client.force_login(self.admin)
        response = self.client.get(reverse("bao_cao_tong_hop"), {"nhom_hd": self.group.pk, "ky": 1})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["report"]["summary"]["so_nhat_ky"], 1)
        report = response.context["report"]
        self.assertEqual(report["contracts"][0]["so_hop_dong"], "HD-REPORT-001")
        self.assertEqual(report["year_cumulative"]["so_buoi"], 2)
        self.assertEqual(report["since_signing"]["so_buoi"], 2)
        contract_response = self.client.get(reverse("bao_cao_thanh_toan"), {"hop_dong": contract.pk, "nam": 2026})
        self.assertEqual(contract_response.status_code, 200)
        self.assertEqual(contract_response.context["report"]["contracts"][0]["hop_dong_id"], contract.pk)
        workbook_response = self.client.get(reverse("bao_cao_tong_hop"), {"format": "xlsx"})
        self.assertEqual(workbook_response.status_code, 200)
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(workbook_response.content), read_only=True)
        self.assertEqual(workbook.sheetnames, ["TongHop", "TheoNhomKy", "TheoTreDichVu"])
        workbook.close()
        self.assertEqual(self.client.get(reverse("bao_cao_hoat_dong")).status_code, 200)

    def test_batch_payment_skips_contracts_without_new_journals(self):
        assignment = PhanCongTre.objects.create(
            phan_bo=self.allocation,
            can_bo_nguon=self.cb1,
            nhom_hd=self.group,
            tre=self.child,
            loai_dich_vu="VLTL",
            so_buoi_du_kien=20,
            dinh_muc_di_lai=Decimal("50000"),
        )
        active_contract = HopDong.objects.create(
            can_bo=self.cb1,
            nhom_hd=self.group,
            so_hop_dong="HD-BATCH-VALID",
            ngay_ky=date(2026, 9, 1),
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30),
            gia_tri_hop_dong=Decimal("10000000"),
            trang_thai="DA_KY",
        )
        no_cost_contract = HopDong.objects.create(
            can_bo=self.cb2,
            nhom_hd=self.group,
            so_hop_dong="HD-BATCH-EMPTY",
            ngay_ky=date(2026, 9, 1),
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30),
            gia_tri_hop_dong=Decimal("10000000"),
            trang_thai="DA_KY",
        )
        NhatKyThucHien.objects.create(
            hop_dong=active_contract,
            phan_cong=assignment,
            ngay_thuc_hien=date(2026, 9, 10),
            ky_can_thiep=13,
            so_buoi_thuc_hien=1,
            don_gia_cong=Decimal("200000"),
            dinh_muc_di_lai=Decimal("50000"),
        )

        self.client.force_login(self.admin)
        response = self.client.post(reverse("tao_phieu_thanh_toan"), {"nhom_hd": self.group.pk, "ky": 13})
        self.assertEqual(response.status_code, 302)
        self.assertEqual(PhieuThanhToan.objects.filter(hop_dong=active_contract).count(), 1)
        self.assertEqual(PhieuThanhToan.objects.filter(hop_dong=no_cost_contract).count(), 0)
        message_text = " ".join(str(message) for message in get_messages(response.wsgi_request))
        self.assertIn("Đã tạo 1 phiếu thanh toán", message_text)
        self.assertIn("Không có nhật ký mới đủ điều kiện", message_text)

    def test_readonly_workflows_are_hidden_from_unapproved_users(self):
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(reverse("bao_cao_tong_hop")).status_code, 302)
        self.assertEqual(self.client.get(reverse("thanh_quyet_toan")).status_code, 302)

        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("bao_cao_tong_hop")).status_code, 200)
        self.assertEqual(self.client.get(reverse("thanh_quyet_toan")).status_code, 200)

    def test_don_vi_list_is_sorted_by_code_descending(self):
        DonVi.objects.create(ma_don_vi="ZZZ-001", ten_don_vi="Đơn vị Z")
        DonVi.objects.create(ma_don_vi="AAA-999", ten_don_vi="Đơn vị A")
        self.client.force_login(self.admin)
        response = self.client.get(reverse("danh_sach_don_vi"))
        codes = [item.ma_don_vi for item in response.context["page_obj"].object_list]
        self.assertEqual(codes, sorted(codes, reverse=True))

    def test_strict_payment_parsing_and_tax_threshold(self):
        self.assertEqual(parse_number("0.125"), Decimal("0.125"))
        self.assertEqual(parse_number("0,125"), Decimal("0.125"))
        self.assertEqual(parse_money_vnd("1.234.567"), Decimal("1234567"))
        self.assertEqual(parse_money_vnd("1,234,567"), Decimal("1234567"))
        with self.assertRaises(ValueError):
            parse_decimal_strict("1.234")
        self.assertEqual(calculate_tncn(Decimal("5000000")), Decimal("500000"))
        self.assertEqual(calculate_tncn(Decimal("4999999")), Decimal("0"))

    def test_payment_voucher_is_snapshotted_and_protects_journal(self):
        assignment = PhanCongTre.objects.create(
            phan_bo=self.allocation,
            can_bo_nguon=self.cb1,
            nhom_hd=self.group,
            tre=self.child,
            loai_dich_vu="VLTL",
            so_buoi_du_kien=20,
            dinh_muc_di_lai=Decimal("50000"),
        )
        contract = HopDong.objects.create(
            can_bo=self.cb1,
            nhom_hd=self.group,
            so_hop_dong="HD-PAY-001",
            ngay_ky=date(2026, 9, 1),
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30),
            gia_tri_hop_dong=Decimal("10000000"),
            trang_thai="DA_KY",
        )
        journal = NhatKyThucHien.objects.create(
            hop_dong=contract,
            phan_cong=assignment,
            ngay_thuc_hien=date(2026, 9, 10),
            ky_can_thiep=1,
            so_buoi_thuc_hien=10,
            so_luot_di_lai_cbct=1,
            don_gia_cong=Decimal("200000"),
            dinh_muc_di_lai=Decimal("50000"),
        )
        NhatKyThucHien.objects.filter(pk=journal.pk).update(so_luot_di_lai_cbct=1)
        voucher = tao_phieu_thanh_toan(self.cb1, contract, 1, self.admin)
        self.assertEqual(voucher.lan_thanh_toan, 1)
        self.assertEqual(voucher.tong_tien_cong, Decimal("2000000"))
        self.assertEqual(voucher.tong_tien_di_lai, Decimal("50000"))
        self.assertEqual(voucher.thue_tncn, Decimal("0"))
        self.assertEqual(ChiTietPhieuThanhToan.objects.filter(phieu=voucher).count(), 1)
        snapshot = snapshot_journals([journal])[0]
        self.assertEqual(snapshot.so_buoi_thuc_hien, 10)
        self.assertEqual(snapshot.don_gia_cong, Decimal("200000"))
        self.assertEqual(snapshot.lan_thanh_toan, 1)
        with self.assertRaises(ValidationError):
            tao_phieu_thanh_toan(self.cb1, contract, 1, self.admin)
        journal.refresh_from_db()
        journal.so_buoi_thuc_hien = 11
        with self.assertRaises(ValidationError):
            journal.save()
        with self.assertRaises(ValidationError):
            journal.delete()

    def test_payment_voucher_round_reuse_after_latest_cancel(self):
        assignment = PhanCongTre.objects.create(
            phan_bo=self.allocation,
            can_bo_nguon=self.cb1,
            nhom_hd=self.group,
            tre=self.child,
            loai_dich_vu="VLTL",
            so_buoi_du_kien=20,
            dinh_muc_di_lai=Decimal("50000"),
        )
        contract = HopDong.objects.create(
            can_bo=self.cb1,
            nhom_hd=self.group,
            so_hop_dong="HD-PAY-002",
            ngay_ky=date(2026, 9, 1),
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30),
            gia_tri_hop_dong=Decimal("10000000"),
            trang_thai="DA_KY",
        )
        for period in (1, 2, 3):
            NhatKyThucHien.objects.create(
                hop_dong=contract, phan_cong=assignment, ngay_thuc_hien=date(2026, 9, period),
                ky_can_thiep=period, so_buoi_thuc_hien=1, don_gia_cong=Decimal("200000"),
                dinh_muc_di_lai=Decimal("50000"),
            )
            NhatKyThucHien.objects.filter(phan_cong=assignment, ky_can_thiep=period).update(so_luot_di_lai_cbct=1)
        first = tao_phieu_thanh_toan(self.cb1, contract, 1, self.admin)
        second = tao_phieu_thanh_toan(self.cb1, contract, 2, self.admin)
        third = tao_phieu_thanh_toan(self.cb1, contract, 3, self.admin)
        self.assertEqual((first.lan_thanh_toan, second.lan_thanh_toan, third.lan_thanh_toan), (1, 2, 3))
        huy_phieu(third, "Tạo nhầm kỳ", self.admin)
        replacement = tao_phieu_thanh_toan(self.cb1, contract, 3, self.admin)
        self.assertEqual(replacement.lan_thanh_toan, 3)
        xac_nhan_chi(replacement, date(2026, 9, 30), self.admin)
        replacement.refresh_from_db()
        self.assertEqual(replacement.trang_thai, "DA_CHI")

    def test_payment_ledger_pages_are_available_to_admin(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("danh_sach_phieu_thanh_toan")).status_code, 200)
        self.assertEqual(self.client.get(reverse("tao_phieu_thanh_toan")).status_code, 200)

    def test_khoa_phan_bo_requires_post_and_admin(self):
        self.client.force_login(self.admin)
        get_response = self.client.get(reverse("khoa_phan_bo", args=[self.allocation.pk]))
        self.allocation.refresh_from_db()
        self.assertEqual(get_response.status_code, 302)
        self.assertFalse(self.allocation.is_locked)

        post_response = self.client.post(reverse("khoa_phan_bo", args=[self.allocation.pk]))
        self.allocation.refresh_from_db()
        self.assertEqual(post_response.status_code, 302)
        self.assertTrue(self.allocation.is_locked)

        other = PhanBoChiTieu.objects.create(can_bo=self.cb2, nhom_hd=self.group)
        self.client.force_login(self.viewer)
        self.client.post(reverse("khoa_phan_bo", args=[other.pk]))
        other.refresh_from_db()
        self.assertFalse(other.is_locked)

        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(reverse("import_phan_cong")).status_code, 302)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("import_phan_cong")).status_code, 200)

    def test_admin_can_edit_delete_and_unlock_locked_allocation(self):
        self.client.force_login(self.admin)
        self.allocation.is_locked = True
        self.allocation.save(update_fields=["is_locked", "updated_at"])
        payload = {
            "can_bo": self.cb1.pk,
            "cbda_quan_ly": "CBDA test",
            "nhom_hd": self.group.pk,
            "tham_gia_ct": "on",
            "ngay_lap": "2026-09-01",
            "so_tre_phcn": "3",
            "so_buoi_phcn": "20",
            "dinh_muc_di_lai_phcn": "50000",
            "so_tre_cs": "0",
            "so_buoi_cs": "10",
            "dinh_muc_di_lai_cs": "50000",
            "ghi_chu": "",
        }
        response = self.client.post(reverse("sua_phan_bo_chi_tieu", args=[self.allocation.pk]), payload)
        self.assertEqual(response.status_code, 302)
        self.allocation.refresh_from_db()
        self.assertEqual(self.allocation.so_tre_phcn, 3)

        unlock = self.client.post(reverse("mo_khoa_phan_bo", args=[self.allocation.pk]))
        self.assertEqual(unlock.status_code, 302)
        self.allocation.refresh_from_db()
        self.assertFalse(self.allocation.is_locked)

        locked_for_delete = PhanBoChiTieu.objects.create(can_bo=self.cb2, nhom_hd=self.group, is_locked=True)
        delete_response = self.client.post(reverse("xoa_phan_bo", args=[locked_for_delete.pk]))
        self.assertEqual(delete_response.status_code, 302)
        self.assertFalse(PhanBoChiTieu.objects.filter(pk=locked_for_delete.pk).exists())

    def test_non_admin_cannot_mutate_assignment_or_workflow_under_locked_records(self):
        assignment = PhanCongTre.objects.create(
            phan_bo=self.allocation,
            can_bo_nguon=self.cb1,
            nhom_hd=self.group,
            tre=self.child,
            loai_dich_vu="VLTL",
            so_buoi_du_kien=20,
            dinh_muc_di_lai=Decimal("50000"),
        )
        self.allocation.is_locked = True
        self.allocation.save(update_fields=["is_locked", "updated_at"])
        contract = HopDong.objects.create(
            can_bo=self.cb1,
            nhom_hd=self.group,
            so_hop_dong="HD-P0-LOCKED",
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30),
            gia_tri_hop_dong=Decimal("1000"),
            is_locked=True,
        )
        self.client.force_login(self.viewer)
        self.assertEqual(self.client.get(reverse("sua_phan_cong", args=[assignment.pk])).status_code, 302)
        self.assertEqual(
            self.client.get(reverse("them_gia_han_thoi_gian_theo_hop_dong", args=[contract.pk])).status_code,
            302,
        )
        self.assertEqual(self.client.get(reverse("them_nhat_ky_thuc_hien", args=[contract.pk])).status_code, 302)
        self.assertEqual(self.client.get(reverse("tao_dot_thanh_toan", args=[contract.pk])).status_code, 302)

    def test_signed_extension_cannot_be_deleted_after_payment(self):
        contract = HopDong.objects.create(
            can_bo=self.cb1,
            nhom_hd=self.group,
            so_hop_dong="HD-P0-008",
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 10, 31),
            gia_tri_hop_dong=Decimal("1000"),
            is_locked=True,
        )
        extension = PhuLucHopDong.objects.create(
            hop_dong=contract,
            loai_phu_luc="GIA_HAN_THOI_GIAN",
            so_phu_luc="GH-TG-008",
            ngay_lap=date(2026, 9, 28),
            is_signed=True,
            den_ngay_cu=date(2026, 9, 30),
            den_ngay_moi=date(2026, 10, 31),
        )
        DotThanhToan.objects.create(hop_dong=contract, nam=2026, thang=9)
        self.client.force_login(self.admin)
        response = self.client.post(reverse("xoa_gia_han_hop_dong", args=[extension.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(PhuLucHopDong.objects.filter(pk=extension.pk).exists())

    def test_acceptance_get_does_not_create_and_value_is_user_confirmed(self):
        contract = HopDong.objects.create(
            can_bo=self.cb1,
            nhom_hd=self.group,
            so_hop_dong="HD-P0-001",
            ngay_ky=date(2026, 9, 1),
            tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30),
            gia_tri_hop_dong=Decimal("1000"),
        )
        self.client.force_login(self.admin)
        response = self.client.get(reverse("cap_nhat_nghiem_thu", args=[contract.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(NghiemThu.objects.count(), 0)
        invalid = self.client.post(reverse("cap_nhat_nghiem_thu", args=[contract.pk]), {
            "ngay_nghiem_thu": "2026-09-30", "ket_qua": "DAT", "gia_tri_nghiem_thu": "1001", "bien_ban_so": "BB-1",
        })
        self.assertEqual(invalid.status_code, 200)
        self.assertEqual(NghiemThu.objects.count(), 0)
        valid = self.client.post(reverse("cap_nhat_nghiem_thu", args=[contract.pk]), {
            "ngay_nghiem_thu": "2026-09-30", "ket_qua": "DAT", "gia_tri_nghiem_thu": "750", "bien_ban_so": "BB-1",
        })
        self.assertEqual(valid.status_code, 302)
        acceptance = NghiemThu.objects.get(hop_dong=contract)
        self.assertEqual(acceptance.gia_tri_nghiem_thu, Decimal("750"))
        contract.refresh_from_db()
        self.assertEqual(contract.trang_thai, "NGHIEM_THU")
        changed = self.client.post(reverse("cap_nhat_nghiem_thu", args=[contract.pk]), {
            "ngay_nghiem_thu": "2026-09-30", "ket_qua": "KHONG_DAT", "gia_tri_nghiem_thu": "500", "bien_ban_so": "BB-1",
        })
        self.assertEqual(changed.status_code, 302)
        contract.refresh_from_db()
        self.assertEqual(contract.trang_thai, "DANG_THUC_HIEN")

    def test_liquidation_requires_real_acceptance_and_locks_contract(self):
        contract = HopDong.objects.create(
            can_bo=self.cb1, nhom_hd=self.group, so_hop_dong="HD-P0-002", ngay_ky=date(2026, 9, 1),
            tu_ngay=date(2026, 9, 1), den_ngay=date(2026, 9, 30), gia_tri_hop_dong=Decimal("1000"),
        )
        self.client.force_login(self.admin)
        blocked = self.client.get(reverse("cap_nhat_thanh_ly", args=[contract.pk]))
        self.assertEqual(blocked.status_code, 302)
        NghiemThu.objects.create(hop_dong=contract, ngay_nghiem_thu=date(2026, 9, 30), ket_qua="DAT", gia_tri_nghiem_thu=Decimal("800"))
        saved = self.client.post(reverse("cap_nhat_thanh_ly", args=[contract.pk]), {
            "ngay_thanh_ly": "2026-09-30", "bien_ban_so": "TL-1",
        })
        self.assertEqual(saved.status_code, 302)
        contract.refresh_from_db()
        self.assertEqual(contract.trang_thai, "THANH_LY")
        self.assertTrue(contract.is_locked)
        self.assertEqual(ThanhLyHopDong.objects.get(hop_dong=contract).gia_tri_thanh_ly, Decimal("800"))

    def test_status_transition_does_not_allow_liquidation_rollback(self):
        contract = HopDong.objects.create(
            can_bo=self.cb1, nhom_hd=self.group, so_hop_dong="HD-P0-003", tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30), gia_tri_hop_dong=Decimal("1000"), trang_thai="THANH_LY",
        )
        with self.assertRaises(Exception):
            transition_hop_dong_status(contract, "DANG_THUC_HIEN")
        self.assertIn("THANH_LY", STATUS_TRANSITIONS)

    def test_contract_edit_rejects_illegal_status_and_locked_contract(self):
        contract = HopDong.objects.create(
            can_bo=self.cb1, nhom_hd=self.group, so_hop_dong="HD-P0-005", tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30), gia_tri_hop_dong=Decimal("1000"),
        )
        self.client.force_login(self.admin)
        payload = {
            "so_hop_dong": contract.so_hop_dong, "ngay_ky": "2026-09-01", "tu_ngay": "2026-09-01",
            "den_ngay": "2026-09-30", "don_gia_cong": "200", "dinh_muc_di_lai_phcn": "50",
            "dinh_muc_di_lai_cs": "50", "gia_tri_hop_dong": "1000", "trang_thai": "THANH_LY", "ghi_chu": "",
        }
        response = self.client.post(reverse("sua_hop_dong", args=[contract.pk]), payload)
        self.assertEqual(response.status_code, 200, response.headers.get("Location", ""))
        contract.refresh_from_db()
        self.assertEqual(contract.trang_thai, "DU_THAO")
        contract.is_locked = True
        contract.save(update_fields=["is_locked", "updated_at"])
        self.assertEqual(self.client.get(reverse("sua_hop_dong", args=[contract.pk])).status_code, 200)

    def test_admin_can_edit_and_delete_locked_contract(self):
        contract = HopDong.objects.create(
            can_bo=self.cb1, nhom_hd=self.group, so_hop_dong="HD-P0-006", tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30), gia_tri_hop_dong=Decimal("1000"), is_locked=True,
        )
        self.client.force_login(self.admin)
        payload = {
            "so_hop_dong": contract.so_hop_dong, "ngay_ky": "2026-09-01", "tu_ngay": "2026-09-01",
            "den_ngay": "2026-09-30", "don_gia_cong": "200", "dinh_muc_di_lai_phcn": "50",
            "dinh_muc_di_lai_cs": "50", "gia_tri_hop_dong": "1200", "trang_thai": "DU_THAO", "ghi_chu": "test",
        }
        response = self.client.post(reverse("sua_hop_dong", args=[contract.pk]), payload)
        self.assertEqual(response.status_code, 302)
        contract.refresh_from_db()
        self.assertEqual(contract.gia_tri_hop_dong, Decimal("1200"))

        delete_response = self.client.post(reverse("xoa_hop_dong", args=[contract.pk]))
        self.assertEqual(delete_response.status_code, 302)
        self.assertFalse(HopDong.objects.filter(pk=contract.pk).exists())

    def test_admin_can_delete_signed_time_extension_and_restore_contract_end(self):
        contract = HopDong.objects.create(
            can_bo=self.cb1, nhom_hd=self.group, so_hop_dong="HD-P0-007", tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 10, 31), gia_tri_hop_dong=Decimal("1000"), is_locked=True,
        )
        extension = PhuLucHopDong.objects.create(
            hop_dong=contract,
            loai_phu_luc="GIA_HAN_THOI_GIAN",
            so_phu_luc="GH-TG-007",
            ngay_lap=date(2026, 9, 28),
            is_signed=True,
            den_ngay_cu=date(2026, 9, 30),
            den_ngay_moi=date(2026, 10, 31),
        )
        self.client.force_login(self.admin)
        response = self.client.post(reverse("xoa_gia_han_hop_dong", args=[extension.pk]))
        self.assertEqual(response.status_code, 302)
        contract.refresh_from_db()
        self.assertEqual(contract.den_ngay, date(2026, 9, 30))
        self.assertFalse(PhuLucHopDong.objects.filter(pk=extension.pk).exists())

    def test_import_is_idempotent_and_identity_contains_cbct(self):
        first = import_assignment_workbook(self._workbook([self._base_row()]))
        second = import_assignment_workbook(self._workbook([self._base_row()]))
        self.assertEqual(first.created, 1)
        self.assertEqual(PhanCongTre.objects.filter(tre=self.child).count(), 1)
        self.assertEqual(second.created, 0)
        self.assertEqual(second.unchanged, 1)

        PhanBoChiTieu.objects.create(can_bo=self.cb2, nhom_hd=self.group, so_tre_phcn=2, so_buoi_phcn=20)
        other = import_assignment_workbook(self._workbook([self._base_row(**{"Mã CB": "ABP0002"})]))
        self.assertEqual(other.created, 1)
        self.assertEqual(PhanCongTre.objects.filter(tre=self.child, loai_dich_vu="VLTL").count(), 2)

    def test_import_invalid_row_rolls_back_and_never_creates_fake_child(self):
        result = import_assignment_workbook(self._workbook([
            self._base_row(**{"IDChild": "CBP-NEW", "Tên trẻ": ""}),
            self._base_row(**{"IDChild": "CBP0001", "Số buổi dự kiến": 0}),
        ]))
        self.assertTrue(result.errors)
        self.assertEqual(PhanCongTre.objects.count(), 0)
        self.assertFalse(Tre.objects.filter(ma_tre="CBP-NEW").exists())

    def test_import_validate_only_rolls_back(self):
        result = import_assignment_workbook(self._workbook([self._base_row()]), validate_only=True)
        self.assertEqual(result.created, 1)
        self.assertEqual(PhanCongTre.objects.count(), 0)

    def test_import_deduplicates_new_child_identity_within_one_file(self):
        rows = [
            self._base_row(**{"IDChild": "CBP-NEW", "Tên trẻ": "Trẻ Mới", "Ngày sinh": "2015-01-01", "Giới tính": "Nam"}),
            self._base_row(**{"IDChild": "CBP-NEW", "Tên trẻ": "Trẻ Mới", "Ngày sinh": "2015-01-01", "Giới tính": "Nam"}),
        ]
        result = import_assignment_workbook(self._workbook(rows))
        self.assertEqual(result.created, 1)
        self.assertEqual(result.unchanged, 1)
        self.assertEqual(Tre.objects.filter(ma_tre="CBP-NEW").count(), 1)
        self.assertEqual(PhanCongTre.objects.filter(tre__ma_tre="CBP-NEW").count(), 1)

    def test_journal_save_uses_decimal_formula(self):
        assignment = PhanCongTre.objects.create(
            phan_bo=self.allocation, can_bo_nguon=self.cb1, nhom_hd=self.group, tre=self.child,
            loai_dich_vu="VLTL", so_buoi_du_kien=20, dinh_muc_di_lai=Decimal("50"),
        )
        contract = HopDong.objects.create(
            can_bo=self.cb1, nhom_hd=self.group, so_hop_dong="HD-P0-004", tu_ngay=date(2026, 9, 1),
            den_ngay=date(2026, 9, 30), gia_tri_hop_dong=Decimal("1000"),
        )
        journal = NhatKyThucHien(
            hop_dong=contract, phan_cong=assignment, ngay_thuc_hien=date(2026, 9, 1),
            so_buoi_thuc_hien=2, so_luot_di_lai_cbct=1, don_gia_cong=Decimal("200"), dinh_muc_di_lai=Decimal("50"),
        )
        journal.save(recalculate_travel=False)
        self.assertEqual(journal.thanh_tien, Decimal("450"))



# =============================================================================
# Hồi quy cho đợt rà soát 29/09/2026: báo cáo GROUP BY, Thanh quyết toán,
# hủy phiếu, import phân bổ (session JSON) và tham số GET xấu.
# =============================================================================
from django.db import connection
from django.db.models import F
from django.test.utils import CaptureQueriesContext

from .parsing import parse_get_int
from .reporting import build_intervention_report, intervention_report_queryset
from .services.payment_ledger import lay_cau_hinh_thue


class ParseGetIntTests(SimpleTestCase):
    def test_accepts_plain_ascii_digits_only(self):
        self.assertEqual(parse_get_int("12"), 12)
        self.assertEqual(parse_get_int(" 7 "), 7)
        for bad in ("", None, "abc", "1.5", "-3", "²", "١٢", "0", "１２"):
            with self.subTest(value=bad):
                self.assertIsNone(parse_get_int(bad))

    def test_bounds_are_enforced(self):
        self.assertIsNone(parse_get_int("0", min_value=1900, max_value=2100))
        self.assertIsNone(parse_get_int("99999", min_value=1900, max_value=2100))
        self.assertEqual(parse_get_int("2026", min_value=1900, max_value=2100), 2026)
        self.assertIsNone(parse_get_int("9" * 40))


class SettlementReportingTestBase(TestCase):
    """Bộ dữ liệu có đủ các tầng nhóm/CBCT để đối chiếu GROUP BY với thuộc tính của model."""

    def setUp(self):
        self.admin_group = Group.objects.create(name="Admin")
        self.admin = User.objects.create_user(username="rpt-admin", password="secret")
        self.admin.groups.add(self.admin_group)
        self.don_vi = DonVi.objects.create(ma_don_vi="DV-RPT", ten_don_vi="Đơn vị báo cáo", nguoi_dai_dien="Người test")
        self.cb1 = CanBo.objects.create(ma_can_bo="RPT0001", ho_ten="CB Một", don_vi=self.don_vi)
        self.cb2 = CanBo.objects.create(ma_can_bo="RPT0002", ho_ten="CB Hai", don_vi=self.don_vi)
        self.g1 = NhomHD.objects.create(ma_nhom_hd="1", ten_nhom_hd="Nhóm 1")
        self.g12 = NhomHD.objects.create(ma_nhom_hd="12", ten_nhom_hd="Nhóm 12")
        self.child1 = Tre.objects.create(ma_tre="RPT-T1", ho_ten="Trẻ 1", ngay_sinh=date(2015, 1, 1), gioi_tinh="Nam")
        self.child2 = Tre.objects.create(ma_tre="RPT-T2", ho_ten="Trẻ 2", ngay_sinh=date(2016, 2, 2), gioi_tinh="Nữ")
        self.alloc1 = PhanBoChiTieu.objects.create(can_bo=self.cb1, nhom_hd=self.g1, so_tre_phcn=2, so_buoi_phcn=50)
        self.alloc2 = PhanBoChiTieu.objects.create(can_bo=self.cb2, nhom_hd=self.g12, so_tre_phcn=2, so_buoi_phcn=50)

    # -- helpers ---------------------------------------------------------
    def _assignment(self, alloc, child, service="VLTL", nhom=None):
        return PhanCongTre.objects.create(
            phan_bo=alloc, can_bo_nguon=alloc.can_bo, nhom_hd=nhom, tre=child, loai_dich_vu=service,
            so_buoi_du_kien=500, dinh_muc_di_lai=Decimal("50000"),
        )

    def _contract(self, so, can_bo, nhom, value="1000000000", status="DA_KY"):
        return HopDong.objects.create(
            can_bo=can_bo, nhom_hd=nhom, so_hop_dong=so, ngay_ky=date(2026, 8, 1),
            tu_ngay=date(2026, 8, 1), den_ngay=date(2026, 12, 31),
            gia_tri_hop_dong=Decimal(value), trang_thai=status,
        )

    def _journal(self, contract, assignment, day, ky=1, buoi=1, di_lai=0, cong="200000", dm="50000", **extra):
        journal = NhatKyThucHien(
            hop_dong=contract, phan_cong=assignment, ngay_thuc_hien=day, ky_can_thiep=ky,
            so_buoi_thuc_hien=buoi, so_luot_di_lai_cbct=di_lai,
            don_gia_cong=Decimal(cong), dinh_muc_di_lai=Decimal(dm), du_lieu_lich_su=True, **extra,
        )
        journal.save(recalculate_travel=False)
        return journal

    def _scenario(self):
        """Hợp đồng nhóm 12 nhưng phân công nhóm 1 (ca gây lỗi #2), nhóm nguồn lịch sử và nhật ký không hợp đồng."""
        self.a1 = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        self.a3 = self._assignment(self.alloc1, self.child1, "HDTL", nhom=self.g1)
        self.a2 = self._assignment(self.alloc2, self.child2, "CSXH")
        self.cA = self._contract("HD-A", self.cb1, self.g12)
        self.cB = self._contract("HD-B", self.cb2, self.g1)
        self.j1 = self._journal(self.cA, self.a1, date(2026, 9, 10), ky=1, buoi=2, di_lai=1)
        self.j2 = self._journal(self.cA, self.a1, date(2026, 9, 11), ky=1, buoi=3)
        self.j3 = self._journal(self.cA, self.a3, date(2026, 9, 12), ky=2, buoi=1, di_lai=1)
        self.j4 = self._journal(self.cB, self.a2, date(2026, 9, 10), ky=1, buoi=4, di_lai=2)
        self.j5 = self._journal(self.cB, self.a2, date(2026, 9, 9), ky=1, buoi=1, nhom_hd_nguon=self.g12)
        self.j6 = self._journal(None, self.a2, date(2026, 9, 13), ky=1, buoi=5)
        NhatKyThucHien.objects.filter(pk=self.j4.pk).update(canh_bao_trung=True)

    def _all_journals(self):
        return list(NhatKyThucHien.objects.select_related(
            "hop_dong__can_bo", "hop_dong__nhom_hd", "nhom_hd_nguon", "can_bo_nguon", "phan_cong__tre",
            "phan_cong__nhom_hd", "phan_cong__phan_bo__can_bo", "phan_cong__phan_bo__nhom_hd",
        ).order_by("ngay_thuc_hien", "id"))

    @staticmethod
    def _label(group):
        return f"{group.ma_nhom_hd} - {group.ten_nhom_hd}" if group else "Chưa xác định"

    @staticmethod
    def _paid(journal):
        total = Decimal("0")
        for detail in journal.chi_tiet_phieu_thanh_toan.filter(hoat_dong=True):
            total += detail.tien_cong + detail.tien_di_lai
        return total


class InterventionReportEquivalenceTests(SettlementReportingTestBase):
    """Báo cáo GROUP BY phải khớp tuyệt đối với cách tính từng nhật ký qua thuộc tính model."""

    def _reference(self):
        contracts, groups, children = {}, {}, {}
        summary = {"contracts": set(), "so_nhat_ky": 0, "so_buoi": 0, "di_lai": 0, "tien_cong": Decimal("0"),
                   "tien_di_lai": Decimal("0"), "da_thanh_toan": Decimal("0"), "thieu_nhom": 0, "trung_lich": 0}

        def bucket(store, key, **initial):
            return store.setdefault(key, {"so_nhat_ky": 0, "so_buoi": 0, "di_lai": 0, "tien_cong": Decimal("0"),
                                          "tien_di_lai": Decimal("0"), "da_thanh_toan": Decimal("0"), **initial})

        for j in self._all_journals():
            if not j.hop_dong_id:
                continue
            group = j.nhom_hd_hieu_luc
            labor = Decimal(j.so_buoi_thuc_hien) * j.don_gia_cong
            travel = Decimal(j.so_luot_di_lai_cbct) * j.dinh_muc_di_lai
            paid = self._paid(j)
            service = j.phan_cong.get_loai_dich_vu_display()
            rows = (
                bucket(contracts, (j.hop_dong.so_hop_dong, j.ky_can_thiep), nhom=self._label(group)),
                bucket(groups, (self._label(group), j.ky_can_thiep)),
                bucket(children, (j.phan_cong.tre.ma_tre, service, self._label(group), j.ky_can_thiep)),
            )
            for index, row in enumerate(rows):
                row["so_nhat_ky"] += 1
                row["so_buoi"] += j.so_buoi_thuc_hien
                row["di_lai"] += j.so_luot_di_lai_cbct
                row["tien_cong"] += labor
                row["tien_di_lai"] += travel
                if index < 2:
                    row["da_thanh_toan"] += paid
            summary["contracts"].add(j.hop_dong_id)
            summary["so_nhat_ky"] += 1
            summary["so_buoi"] += j.so_buoi_thuc_hien
            summary["di_lai"] += j.so_luot_di_lai_cbct
            summary["tien_cong"] += labor
            summary["tien_di_lai"] += travel
            summary["da_thanh_toan"] += paid
            summary["thieu_nhom"] += 0 if group else 1
            summary["trung_lich"] += 1 if j.canh_bao_trung else 0
        return contracts, groups, children, summary

    def test_report_matches_model_based_reference_at_every_level(self):
        self._scenario()
        # Lập phiếu cho HD-A kỳ 1 để có số "đã thanh toán".
        tao_phieu_thanh_toan(self.cb1, self.cA, 1)
        contracts, groups, children, summary = self._reference()

        report = build_intervention_report(intervention_report_queryset())

        got_contracts = {(r["so_hop_dong"], r["ky"]): r for r in report["contracts"]}
        self.assertEqual(set(got_contracts), set(contracts))
        for key, expected in contracts.items():
            row = got_contracts[key]
            for field in ("nhom", "so_nhat_ky", "so_buoi", "di_lai", "tien_cong", "tien_di_lai", "da_thanh_toan"):
                self.assertEqual(row[field], expected[field], f"contract {key} field {field}")

        got_groups = {(r["nhom"], r["ky"]): r for r in report["groups"]}
        self.assertEqual(set(got_groups), set(groups))
        for key, expected in groups.items():
            for field in ("so_nhat_ky", "so_buoi", "di_lai", "tien_cong", "tien_di_lai", "da_thanh_toan"):
                self.assertEqual(got_groups[key][field], expected[field], f"group {key} field {field}")

        got_children = {(r["ma_tre"], r["dich_vu"], r["nhom"], r["ky"]): r for r in report["children"]}
        self.assertEqual(set(got_children), set(children))
        for key, expected in children.items():
            for field in ("so_buoi", "di_lai", "tien_cong", "tien_di_lai"):
                self.assertEqual(got_children[key][field], expected[field], f"child {key} field {field}")

        s = report["summary"]
        self.assertEqual(s["so_hop_dong"], len(summary["contracts"]))
        for field in ("so_nhat_ky", "so_buoi", "di_lai", "tien_cong", "tien_di_lai", "da_thanh_toan", "thieu_nhom", "trung_lich"):
            self.assertEqual(s[field], summary[field], field)
        self.assertEqual(s["tong_gross"], summary["tien_cong"] + summary["tien_di_lai"])
        self.assertGreater(s["da_thanh_toan"], 0)

    def test_contract_row_uses_group_of_earliest_journal(self):
        self._scenario()
        row = next(r for r in build_intervention_report(intervention_report_queryset())["contracts"]
                   if r["so_hop_dong"] == "HD-B" and r["ky"] == 1)
        # j5 (09-09, nhóm nguồn 12) sớm hơn j4 (09-10, nhóm hợp đồng 1).
        self.assertEqual(row["nhom"], "12 - Nhóm 12")
        self.assertEqual(row["so_nhat_ky"], 2)

    def test_group_filter_follows_effective_group_priority(self):
        self._scenario()
        in_g12 = intervention_report_queryset(nhom_hd_id=self.g12.pk)
        in_g1 = intervention_report_queryset(nhom_hd_id=self.g1.pk)
        # j1-j3: hợp đồng nhóm 12 dù phân công nhóm 1; j5: nhóm nguồn 12; j4: hợp đồng nhóm 1.
        self.assertEqual(set(in_g12.values_list("pk", flat=True)), {self.j1.pk, self.j2.pk, self.j3.pk, self.j5.pk})
        self.assertEqual(set(in_g1.values_list("pk", flat=True)), {self.j4.pk})

    def test_year_and_signing_cumulatives_are_computed(self):
        self._scenario()
        report = build_intervention_report(
            intervention_report_queryset(nam=2026),
            year_queryset=intervention_report_queryset(tu_ngay=date(2026, 1, 1), den_ngay=date(2026, 12, 31)),
            since_signing_queryset=intervention_report_queryset().filter(ngay_thuc_hien__gte=F("hop_dong__ngay_ky")),
            report_year=2026,
        )
        self.assertEqual(report["summary"]["so_nhat_ky"], 5)
        self.assertEqual(report["year_cumulative"]["so_nhat_ky"], 5)
        self.assertEqual(report["since_signing"]["so_nhat_ky"], 5)

    def test_empty_report_is_well_formed(self):
        report = build_intervention_report(intervention_report_queryset())
        self.assertEqual(report["contracts"], [])
        self.assertEqual(report["groups"], [])
        self.assertEqual(report["children"], [])
        self.assertEqual(report["summary"]["so_nhat_ky"], 0)
        self.assertEqual(report["summary"]["tong_gross"], Decimal("0"))


class ReportQueryCountTests(SettlementReportingTestBase):
    """Số truy vấn không được tăng theo số nhật ký (không còn nạp từng nhật ký vào Python)."""

    def _count(self, callable_):
        with CaptureQueriesContext(connection) as ctx:
            callable_()
        return len(ctx)

    def _add_many_journals(self, n=30):
        for i in range(n):
            contract = self._contract(f"HD-BULK-{i}", self.cb1 if i % 2 else self.cb2, self.g1 if i % 3 else self.g12)
            assignment = self._assignment(self.alloc1 if i % 2 else self.alloc2, self.child1 if i % 2 else self.child2)
            for day in range(1, 4):
                self._journal(contract, assignment, date(2026, 9, day), ky=1 + i % 3, buoi=1 + day)

    def test_report_builder_query_count_is_constant(self):
        self._scenario()
        build = lambda: build_intervention_report(
            intervention_report_queryset(), year_queryset=intervention_report_queryset(),
            since_signing_queryset=intervention_report_queryset(),
        )
        small = self._count(build)
        self._add_many_journals()
        large = self._count(build)
        self.assertEqual(small, large)
        self.assertLessEqual(large, 20)

    def test_report_page_query_count_is_constant(self):
        self._scenario()
        self.client.force_login(self.admin)
        url = reverse("bao_cao_tong_hop")
        self.client.get(url)  # làm nóng cache nội bộ (content types, session...)
        small = self._count(lambda: self.client.get(url))
        self._add_many_journals()
        large = self._count(lambda: self.client.get(url))
        self.assertEqual(small, large)

    def test_thanh_quyet_toan_query_count_is_constant(self):
        self._scenario()
        self.client.force_login(self.admin)
        url = reverse("thanh_quyet_toan")
        self.client.get(url)
        small = self._count(lambda: self.client.get(url))
        self._add_many_journals()
        large = self._count(lambda: self.client.get(url))
        self.assertEqual(small, large)


class ThanhQuyetToanTests(SettlementReportingTestBase):
    def _rows(self, **params):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("thanh_quyet_toan"), params)
        self.assertEqual(response.status_code, 200)
        return response.context["danh_sach"]

    def test_group_filter_matches_report_page(self):
        """Ca của review: hợp đồng nhóm 12, phân công nhóm 1 → TQT và Báo cáo phải cùng kết luận."""
        self.a1 = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        contract = self._contract("HD-X", self.cb1, self.g12)
        self._journal(contract, self.a1, date(2026, 9, 10), ky=1, buoi=2)
        self.client.force_login(self.admin)
        for group, expected in ((self.g1, 0), (self.g12, 1)):
            with self.subTest(group=group.ma_nhom_hd):
                tqt = self.client.get(reverse("thanh_quyet_toan"), {"nhom_hd": group.pk}).context["danh_sach"]
                report = self.client.get(reverse("bao_cao_tong_hop"), {"nhom_hd": group.pk}).context["report"]
                self.assertEqual(sum(r["journal_count"] for r in tqt), expected)
                self.assertEqual(report["summary"]["so_nhat_ky"], expected)

    def test_totals_match_model_based_reference(self):
        self._scenario()
        config = lay_cau_hinh_thue(timezone_localdate())
        expected = {}
        for j in self._all_journals():
            group, staff = j.nhom_hd_hieu_luc, j.can_bo_hieu_luc
            if not group or not staff:
                continue
            item = expected.setdefault((group.pk, j.ky_can_thiep), {"journals": 0, "hd": set(), "staff": {}, "so_buoi": 0, "di_lai": 0})
            item["journals"] += 1
            if j.hop_dong_id:
                item["hd"].add(j.hop_dong_id)
            item["so_buoi"] += j.so_buoi_thuc_hien
            item["di_lai"] += j.so_luot_di_lai_cbct
            amounts = item["staff"].setdefault((staff.pk, j.hop_dong_id), [Decimal("0"), Decimal("0")])
            amounts[0] += Decimal(j.so_buoi_thuc_hien) * j.don_gia_cong
            amounts[1] += Decimal(j.so_luot_di_lai_cbct) * j.dinh_muc_di_lai
        rows = {(r["nhom"].pk, r["ky"]): r for r in self._rows()}
        self.assertEqual(set(rows), set(expected))
        for key, item in expected.items():
            breakdowns = [calculate_payment_breakdown(c, t, config) for c, t in item["staff"].values()]
            row = rows[key]
            self.assertEqual(row["journal_count"], item["journals"])
            self.assertEqual(row["hop_dong_count"], len(item["hd"]))
            self.assertEqual(row["can_bo_count"], len({staff_pk for staff_pk, _ in item["staff"]}))
            self.assertEqual(row["so_buoi"], item["so_buoi"])
            self.assertEqual(row["di_lai"], item["di_lai"])
            for field in ("tien_cong", "tien_di_lai", "tong_truoc_thue", "thue_tncn", "thuc_linh"):
                self.assertEqual(row[field], sum((b[field] for b in breakdowns), Decimal("0")), f"{key} {field}")

    def test_tax_follows_cau_hinh_thue_in_database(self):
        self.a1 = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        contract = self._contract("HD-TAX", self.cb1, self.g1)
        self._journal(contract, self.a1, date(2026, 9, 10), ky=1, buoi=130)  # 26.000.000 tiền công
        # Không có cấu hình: dùng mặc định 5 triệu / 10%.
        self.assertEqual(self._rows()[0]["thue_tncn"], Decimal("2600000"))
        CauHinhThue.objects.create(tu_ngay=date(2020, 1, 1), nguong_thue=Decimal("20000000"), ty_le=Decimal("0.05"), can_cu_phap_ly="Test")
        self.assertEqual(self._rows()[0]["thue_tncn"], Decimal("1300000"))
        CauHinhThue.objects.create(tu_ngay=date(2021, 1, 1), nguong_thue=Decimal("30000000"), ty_le=Decimal("0.05"), can_cu_phap_ly="Test 2")
        row = self._rows()[0]
        self.assertEqual(row["thue_tncn"], Decimal("0"))
        self.assertEqual(row["thuc_linh"], row["tong_truoc_thue"])

    def test_single_contract_tax_matches_voucher(self):
        self.a1 = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        contract = self._contract("HD-VOUCHER", self.cb1, self.g1)
        self._journal(contract, self.a1, date(2026, 9, 10), ky=1, buoi=40, di_lai=3)
        CauHinhThue.objects.create(tu_ngay=date(2020, 1, 1), nguong_thue=Decimal("1000000"), ty_le=Decimal("0.07"), can_cu_phap_ly="Test")
        voucher = tao_phieu_thanh_toan(self.cb1, contract, 1)
        row = self._rows()[0]
        self.assertEqual(row["thue_tncn"], voucher.thue_tncn)
        self.assertEqual(row["thuc_linh"], voucher.thuc_nhan)

    def test_tax_is_per_contract_not_accumulated_across_contracts(self):
        """Quy tắc nghiệp vụ: ngưỡng TNCN áp cho từng lần thanh toán của một số HĐ."""
        self.a1 = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        first = self._contract("HD-T1", self.cb1, self.g1)
        second = self._contract("HD-T2", self.cb1, self.g1)
        self._journal(first, self.a1, date(2026, 9, 10), ky=1, buoi=15)   # 3.000.000 < 5.000.000
        self._journal(second, self.a1, date(2026, 9, 11), ky=1, buoi=15)  # 3.000.000 < 5.000.000
        row = self._rows()[0]
        self.assertEqual(row["hop_dong_count"], 2)
        self.assertEqual(row["can_bo_count"], 1)
        self.assertEqual(row["tien_cong"], Decimal("6000000"))
        self.assertEqual(row["thue_tncn"], Decimal("0"))  # gộp theo CBCT sẽ ra 600.000 (sai)
        v1 = tao_phieu_thanh_toan(self.cb1, first, 1)
        v2 = tao_phieu_thanh_toan(self.cb1, second, 1)
        self.assertEqual(row["thue_tncn"], v1.thue_tncn + v2.thue_tncn)
        # Một hợp đồng vượt ngưỡng thì chỉ hợp đồng đó chịu thuế.
        self._journal(first, self.a1, date(2026, 9, 12), ky=2, buoi=30)  # 6.000.000 >= 5.000.000
        self._journal(second, self.a1, date(2026, 9, 13), ky=2, buoi=5)  # 1.000.000
        ky2 = next(r for r in self._rows() if r["ky"] == 2)
        self.assertEqual(ky2["thue_tncn"], Decimal("600000"))

    def test_can_bo_filter_uses_contract_staff(self):
        """CBCT hiệu lực = nguồn → hợp đồng → phân bổ; bộ lọc không được bỏ qua tầng hợp đồng."""
        self.a1 = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        self.a1.can_bo_nguon = None
        self.a1.save()
        contract = self._contract("HD-STAFF", self.cb2, self.g1)  # hợp đồng của cb2, phân bổ của cb1
        journal = self._journal(contract, self.a1, date(2026, 9, 10), ky=1, buoi=1)
        self.assertEqual(journal.can_bo_hieu_luc, self.cb2)
        self.assertEqual(sum(r["journal_count"] for r in self._rows(can_bo=self.cb2.pk)), 1)
        self.assertEqual(sum(r["journal_count"] for r in self._rows(can_bo=self.cb1.pk)), 0)


def timezone_localdate():
    from django.utils import timezone
    return timezone.localdate()


class HuyPhieuTests(SettlementReportingTestBase):
    def _two_vouchers(self):
        self.a1 = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        contract = self._contract("HD-HUY", self.cb1, self.g1)
        self._journal(contract, self.a1, date(2026, 9, 10), ky=1, buoi=2)
        self._journal(contract, self.a1, date(2026, 9, 20), ky=2, buoi=3)
        first = tao_phieu_thanh_toan(self.cb1, contract, 1)
        second = tao_phieu_thanh_toan(self.cb1, contract, 2)
        self.assertEqual((first.lan_thanh_toan, second.lan_thanh_toan), (1, 2))
        return first, second

    def test_cannot_cancel_older_voucher_while_newer_is_active(self):
        first, second = self._two_vouchers()
        with self.assertRaisesMessage(ValidationError, "mới nhất"):
            huy_phieu(first, "Sai kỳ")
        first.refresh_from_db()
        self.assertEqual(first.trang_thai, "CHO_CHI")

    def test_can_cancel_previous_voucher_after_newest_was_cancelled(self):
        """Lỗi #5: phiếu đã hủy không được tính là 'lần mới nhất'."""
        first, second = self._two_vouchers()
        huy_phieu(second, "Hủy lần 2")
        first.refresh_from_db()
        huy_phieu(first, "Hủy lần 1")
        first.refresh_from_db()
        self.assertEqual(first.trang_thai, "HUY")
        self.assertIsNone(first.hoat_dong)
        self.assertFalse(first.chi_tiet.filter(hoat_dong=True).exists())

    def test_cancel_requires_reason_and_pending_state(self):
        first, second = self._two_vouchers()
        with self.assertRaises(ValidationError):
            huy_phieu(second, "   ")
        xac_nhan_chi(second)
        with self.assertRaises(ValidationError):
            huy_phieu(second, "Đã chi rồi")


class BadQueryParameterTests(SettlementReportingTestBase):
    BAD_VALUES = ({"nam": "0"}, {"nam": "99999"}, {"nam": "²"}, {"ky": "²"}, {"ky": "0"}, {"thang": "²"}, {"thang": "13"},
                  {"nhom_hd": "²"}, {"hop_dong": "²"}, {"can_bo": "²"}, {"nam": "9" * 40}, {"tu_ngay": "not-a-date"})

    def test_pages_do_not_crash_on_malformed_filters(self):
        self._scenario()
        self.client.force_login(self.admin)
        for name in ("bao_cao_tong_hop", "thanh_quyet_toan", "de_nghi_thanh_toan"):
            for params in self.BAD_VALUES:
                with self.subTest(view=name, params=params):
                    self.assertEqual(self.client.get(reverse(name), params).status_code, 200)

    def test_excel_export_survives_malformed_filters(self):
        self._scenario()
        self.client.force_login(self.admin)
        response = self.client.get(reverse("bao_cao_tong_hop"), {"format": "xlsx", "nam": "²", "ky": "²"})
        self.assertEqual(response.status_code, 200)

    def test_malformed_filter_is_ignored_not_applied(self):
        self._scenario()
        self.client.force_login(self.admin)
        baseline = self.client.get(reverse("bao_cao_tong_hop")).context["report"]["summary"]["so_nhat_ky"]
        got = self.client.get(reverse("bao_cao_tong_hop"), {"nam": "99999", "ky": "²"}).context["report"]["summary"]["so_nhat_ky"]
        self.assertEqual(got, baseline)

    def test_valid_filters_still_apply(self):
        self._scenario()
        self.client.force_login(self.admin)
        response = self.client.get(reverse("bao_cao_tong_hop"), {"ky": "2", "nam": "2026"})
        self.assertEqual(response.context["report"]["summary"]["so_nhat_ky"], 1)


class ImportPhanBoSessionTests(SettlementReportingTestBase):
    """Lỗi #7: preview lưu date/Decimal vào session JSON làm TypeError."""

    def _upload(self, **overrides):
        row = {"MaCBCT": "RPT0001", "NhomHD": "1", "SoTrePHCN": 2, "SoBuoiPHCN": 10, "SoTreCS": 1, "SoBuoiCS": 5,
               "DMDL_PHCN": "2", "DMDL_CS": "1", "NgayLap": "2026-09-01"}
        row.update(overrides)
        stream = BytesIO()
        pd.DataFrame([row]).to_excel(stream, index=False)
        return SimpleUploadedFile("pb.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_preview_then_confirm_creates_allocation(self):
        self.client.force_login(self.admin)
        before = PhanBoChiTieu.objects.count()
        preview = self.client.post(reverse("import_phan_bo"), {"excel_file": self._upload()})
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.context["valid_count"], 1)
        stored = self.client.session["import_phan_bo_valid_data"]
        self.assertEqual(stored[0]["ngay_lap"], "2026-09-01")
        confirm = self.client.post(reverse("confirm_import_phan_bo"))
        self.assertEqual(confirm.status_code, 302)
        self.assertEqual(PhanBoChiTieu.objects.count(), before + 1)
        created = PhanBoChiTieu.objects.order_by("-pk").first()
        self.assertEqual((created.can_bo_id, created.nhom_hd_id), (self.cb1.pk, self.g1.pk))
        self.assertEqual(created.ngay_lap, date(2026, 9, 1))
        self.assertEqual(created.dinh_muc_di_lai_phcn, Decimal(str(FinancialConfig.DON_GIA_DI_LAI_DM2)))
        self.assertEqual(created.dinh_muc_di_lai_cs, Decimal(str(FinancialConfig.DON_GIA_DI_LAI_DM1)))

    def test_invalid_rows_are_not_stored_for_confirmation(self):
        self.client.force_login(self.admin)
        response = self.client.post(reverse("import_phan_bo"), {"excel_file": self._upload(MaCBCT="KHONG-CO")})
        self.assertEqual(response.context["invalid_count"], 1)
        self.assertEqual(self.client.session["import_phan_bo_valid_data"], [])
        before = PhanBoChiTieu.objects.count()
        self.client.post(reverse("confirm_import_phan_bo"))
        self.assertEqual(PhanBoChiTieu.objects.count(), before)


class PaymentExportQueryCountTests(SettlementReportingTestBase):
    """Lỗi N+1 khi xuất DSTK/ĐNTT: số truy vấn không được tăng theo số nhật ký."""

    def _vouchered_journals(self, n):
        assignment = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        contract = self._contract(f"HD-N1-{n}-{NhatKyThucHien.objects.count()}", self.cb1, self.g1)
        for day in range(1, n + 1):
            self._journal(contract, assignment, date(2026, 9, day), ky=1, buoi=1)
        tao_phieu_thanh_toan(self.cb1, contract, 1)
        return list(NhatKyThucHien.objects.filter(hop_dong=contract).order_by("id"))

    def _snapshot_queries(self, journals):
        with CaptureQueriesContext(connection) as ctx:
            rows = snapshot_journals(journals)
            for row in rows:  # các thuộc tính mà exporter đọc
                row.can_bo_hieu_luc, row.nhom_hd_hieu_luc, row.hop_dong_hieu_luc, row.phan_cong.tre
        return len(ctx)

    def test_snapshot_journals_query_count_is_constant(self):
        few = self._vouchered_journals(3)
        small = self._snapshot_queries(few)
        many = self._vouchered_journals(25)
        large = self._snapshot_queries(many)
        self.assertEqual(small, large)
        self.assertLessEqual(large, 3)

    def test_creating_voucher_does_not_query_per_journal(self):
        assignment = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        counts = []
        for size in (2, 20):
            contract = self._contract(f"HD-EX-{size}", self.cb1, self.g1)
            for day in range(1, size + 1):
                self._journal(contract, assignment, date(2026, 9, day), ky=1, buoi=1)
            with CaptureQueriesContext(connection) as ctx:
                tao_phieu_thanh_toan(self.cb1, contract, 1)
            counts.append(len(ctx))
        self.assertEqual(counts[0], counts[1])



class PhieuStaleStateTests(SettlementReportingTestBase):
    """Hai kế toán thao tác cùng lúc: đối tượng cũ không được ghi đè trạng thái mới trong CSDL."""
    def _two_vouchers(self):
        self.a1 = self._assignment(self.alloc1, self.child1, "VLTL", nhom=self.g1)
        contract = self._contract("HD-HUY", self.cb1, self.g1)
        self._journal(contract, self.a1, date(2026, 9, 10), ky=1, buoi=2)
        self._journal(contract, self.a1, date(2026, 9, 20), ky=2, buoi=3)
        first = tao_phieu_thanh_toan(self.cb1, contract, 1)
        second = tao_phieu_thanh_toan(self.cb1, contract, 2)
        self.assertEqual((first.lan_thanh_toan, second.lan_thanh_toan), (1, 2))
        return first, second

    def test_confirm_payment_with_stale_object_after_cancel_is_rejected(self):
        first, second = self._two_vouchers()
        stale = PhieuThanhToan.objects.get(pk=second.pk)
        huy_phieu(second, "Hủy trước")
        with self.assertRaises(ValidationError):
            xac_nhan_chi(stale)
        second.refresh_from_db()
        self.assertEqual(second.trang_thai, "HUY")
        self.assertIsNone(second.ngay_chi)
        self.assertFalse(second.chi_tiet.filter(hoat_dong=True).exists())

    def test_cancel_with_stale_object_after_payment_is_rejected(self):
        first, second = self._two_vouchers()
        stale = PhieuThanhToan.objects.get(pk=second.pk)
        xac_nhan_chi(second)
        with self.assertRaises(ValidationError):
            huy_phieu(stale, "Muộn rồi")
        second.refresh_from_db()
        self.assertEqual(second.trang_thai, "DA_CHI")
        self.assertTrue(second.chi_tiet.filter(hoat_dong=True).exists())

    def test_instances_are_synced_after_success(self):
        first, second = self._two_vouchers()
        xac_nhan_chi(second, user=self.admin)
        self.assertEqual((second.trang_thai, second.xac_nhan_chi_boi), ("DA_CHI", self.admin))
        self.assertIsNotNone(second.ngay_chi)
        first_again = PhieuThanhToan.objects.get(pk=first.pk)
        with self.assertRaises(ValidationError):  # phiếu lần 2 đã chi, lần 1 không phải mới nhất
            huy_phieu(first_again, "x")


class PermissionQueryCountTests(SettlementReportingTestBase):
    """Kiểm tra quyền (decorator + context processor) chỉ tốn một truy vấn nhóm cho mỗi request."""

    def test_role_checks_use_single_group_query(self):
        keto = User.objects.create_user(username="keto", password="secret")
        keto.groups.add(Group.objects.create(name="KeToan"))
        self.client.force_login(keto)
        url = reverse("thanh_quyet_toan")
        self.client.get(url)
        with CaptureQueriesContext(connection) as ctx:
            response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        group_queries = [q for q in ctx.captured_queries if "auth_group" in q["sql"]]
        self.assertEqual(len(group_queries), 1, [q["sql"] for q in group_queries])

    def test_role_flags_and_access_rules_unchanged(self):
        keto = User.objects.create_user(username="keto2", password="secret")
        keto.groups.add(Group.objects.create(name="KeToan"))
        self.client.force_login(keto)
        response = self.client.get(reverse("thanh_quyet_toan"))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["is_ketoan"])
        self.assertFalse(response.context["is_admin"])
        self.assertFalse(response.context["is_cbda"])
        nobody = User.objects.create_user(username="nobody", password="secret")
        self.client.force_login(nobody)
        self.assertEqual(self.client.get(reverse("thanh_quyet_toan")).status_code, 302)
        self.client.force_login(self.admin)
        self.assertTrue(self.client.get(reverse("thanh_quyet_toan")).context["is_admin"])


class ReportChildrenPaginationTests(SettlementReportingTestBase):
    def _many_children(self, count):
        contract = self._contract("HD-PAGE", self.cb1, self.g1)
        for i in range(count):
            child = Tre.objects.create(ma_tre=f"PG{i:04d}", ho_ten=f"Trẻ {i}", ngay_sinh=date(2015, 1, 1), gioi_tinh="Nam")
            assignment = self._assignment(self.alloc1, child, "VLTL", nhom=self.g1)
            self._journal(contract, assignment, date(2026, 9, 1 + i % 25), ky=1, buoi=1)

    def test_children_table_is_paginated_but_totals_and_excel_are_complete(self):
        self._many_children(130)
        self.client.force_login(self.admin)
        url = reverse("bao_cao_tong_hop")
        page1 = self.client.get(url, {"ky": "1"})
        self.assertEqual(len(page1.context["page_obj"]), 100)
        self.assertEqual(page1.context["page_obj"].paginator.count, 130)
        self.assertEqual(page1.context["report"]["summary"]["so_nhat_ky"], 130)
        self.assertContains(page1, "ky=1")  # bộ lọc được giữ khi sang trang
        page2 = self.client.get(url, {"ky": "1", "page": "2"})
        self.assertEqual(len(page2.context["page_obj"]), 30)
        self.assertEqual(self.client.get(url, {"page": "abc"}).status_code, 200)
        self.assertEqual(self.client.get(url, {"page": "9999"}).status_code, 200)
        from openpyxl import load_workbook
        export = self.client.get(url, {"format": "xlsx", "ky": "1"})
        workbook = load_workbook(BytesIO(export.content), read_only=True)
        self.assertEqual(sum(1 for _ in workbook["TheoTreDichVu"].iter_rows()) - 1, 130)
        workbook.close()