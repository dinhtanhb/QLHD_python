from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from .financial import FinancialConfig
from .models import (
    ChiTietKhoiLuongHopDong,
    ChiTietPhuLucPhanCong,
    DeXuatHopDong,
    HopDong,
    PhanCongTre,
    PhuLucHopDong,
)


TEMPLATE_ROOT = Path(settings.BASE_DIR) / "quanly" / "document_templates"
CONTRACT_TEMPLATE = TEMPLATE_ROOT / "hop_dong" / "Mau_HopDong.docx"
ANNEX_TEMPLATE = TEMPLATE_ROOT / "phu_luc" / "Mau_PhuLucPhanCong.docx"
PLACEHOLDER_PATTERN = re.compile(r"\{\{[^{}]+\}\}")


def _document(path):
    """Nạp python-docx khi thực sự xuất file, không chặn Django khởi động."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ValidationError("Thiếu thư viện python-docx. Hãy cài dependencies trong requirements.txt.") from exc
    return Document(str(path))


def _money(value):
    return f"{Decimal(value or 0):,.0f}".replace(",", ".")


def _number_to_words(value):
    """Đổi số nguyên tiền thành chữ ở mức đủ dùng cho mẫu hợp đồng."""
    units = ["", "nghìn", "triệu", "tỷ"]
    digits = ["không", "một", "hai", "ba", "bốn", "năm", "sáu", "bảy", "tám", "chín"]
    number = int(Decimal(value or 0))
    if number == 0:
        return "Không đồng"

    def group_words(group, full=False):
        hundred, remainder = divmod(group, 100)
        ten, one = divmod(remainder, 10)
        words = []
        if hundred:
            words.extend([digits[hundred], "trăm"])
        elif full and remainder:
            words.extend(["không", "trăm"])
        if ten > 1:
            words.extend([digits[ten], "mươi"])
            if one == 1:
                words.append("mốt")
            elif one == 5:
                words.append("lăm")
            elif one:
                words.append(digits[one])
        elif ten == 1:
            words.append("mười")
            if one == 5:
                words.append("lăm")
            elif one:
                words.append(digits[one])
        elif one:
            words.append(digits[one])
        return words

    groups = []
    while number:
        number, group = divmod(number, 1000)
        groups.append(group)
    words = []
    for index in range(len(groups) - 1, -1, -1):
        group = groups[index]
        if not group:
            continue
        words.extend(group_words(group, full=index < len(groups) - 1))
        if index:
            words.append(units[index])
    return " ".join(words).capitalize() + " đồng"


def _date_parts(value):
    value = value or date.today()
    return {
        "Ngay": value.strftime("%d/%m/%Y"),
        "Thang": str(value.month),
        "Nam": str(value.year),
    }


def _service_display(value):
    return dict(PhanCongTre.LOAI_DV_CHOICES).get(value, value or "")


def _replace_text_in_paragraph(paragraph, context):
    full_text = paragraph.text
    if not full_text:
        return
    replaced = full_text
    for key, value in context.items():
        replaced = replaced.replace("{{" + key + "}}", str(value if value is not None else ""))
    if replaced == full_text:
        return
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(replaced)
        return
    runs[0].text = replaced
    for run in runs[1:]:
        run.text = ""


def _iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            yield paragraph
        for paragraph in section.footer.paragraphs:
            yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                for nested_table in cell.tables:
                    for row2 in nested_table.rows:
                        for cell2 in row2.cells:
                            for paragraph in cell2.paragraphs:
                                yield paragraph


def _replace_document(document, context):
    for paragraph in _iter_paragraphs(document):
        _replace_text_in_paragraph(paragraph, context)


def _clone_table_row(table, source_row):
    new_row = deepcopy(source_row._tr)
    table._tbl.append(new_row)
    return table.rows[-1]


def _replace_row(row, context):
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            _replace_text_in_paragraph(paragraph, context)


def _remove_row(row):
    row._tr.getparent().remove(row._tr)


def _one_date(rows):
    dates = {row.ngay_phan_cong for row in rows if row.ngay_phan_cong}
    if len(dates) == 1:
        return next(iter(dates)).strftime("%d/%m/%Y")
    if len(dates) > 1:
        return "Nhiều ngày"
    return ""


def _allocation_context(hop_dong):
    allocation = hop_dong.de_xuat.phan_bo
    cong = Decimal(hop_dong.don_gia_cong or FinancialConfig.DON_GIA_CONG)
    phcn_work = Decimal(allocation.so_tre_phcn) * Decimal(allocation.so_buoi_phcn) * cong
    phcn_travel = Decimal(allocation.so_tre_phcn) * Decimal(allocation.so_buoi_phcn) * Decimal(
        allocation.dinh_muc_di_lai_phcn or 0
    )
    cs_work = Decimal(allocation.so_tre_cs) * Decimal(allocation.so_buoi_cs) * cong
    cs_travel = Decimal(allocation.so_tre_cs) * Decimal(allocation.so_buoi_cs) * Decimal(
        allocation.dinh_muc_di_lai_cs or 0
    )
    contract_date = _date_parts(hop_dong.ngay_ky)
    start_date = _date_parts(hop_dong.tu_ngay)
    end_date = _date_parts(hop_dong.den_ngay)
    can_bo = hop_dong.can_bo
    total = phcn_work + phcn_travel + cs_work + cs_travel
    values = {
        "SoHopDong": hop_dong.so_hop_dong,
        "MaSoGVMN": hop_dong.can_bo.ma_can_bo,
        "HoTenGVMN": hop_dong.can_bo.ho_ten,
        "DanhXung": "Ông/Bà",
        "CCCD": can_bo.cccd or "",
        "NgayCapCCCD": can_bo.ngay_cap.strftime("%d/%m/%Y") if can_bo.ngay_cap else "",
        "NoiCapCCCD": can_bo.noi_cap or "",
        "DiaChi": can_bo.dia_chi or "",
        "DienThoai": can_bo.dien_thoai or "",
        "Email": can_bo.email or "",
        "NganHang": can_bo.ngan_hang or "",
        "SoTaiKhoan": can_bo.tai_khoan or "",
        "TongTrePHCN": allocation.so_tre_phcn,
        "TongBuoiPHCN": allocation.so_buoi_phcn,
        "ThanhTienPHCN": _money(phcn_work),
        "HoTroDiLaiPHCN": _money(phcn_travel),
        "TongTreCSXH": allocation.so_tre_cs,
        "TongBuoiCSXH": allocation.so_buoi_cs,
        "ThanhTienCSXH": _money(cs_work),
        "HoTroDiLaiCSXH": _money(cs_travel),
        "TongTien": _money(total),
        "GiaTriHopDong": _money(hop_dong.gia_tri_hop_dong),
        "GiaTriHopDongBangChu": _number_to_words(total),
        "TuNgay": start_date["Ngay"],
        "DenNgay": end_date["Ngay"],
    }
    values.update({f"NgayKy_{key}": value for key, value in contract_date.items()})
    return values


def _annex_context(hop_dong, rows):
    values = _allocation_context(hop_dong)
    values.update(
        {
            "KyPhanCong": 1,
            "NgayPhanCong": _one_date(rows),
        }
    )
    return values


def _fill_assignment_table(document, hop_dong, rows):
    for table in document.tables:
        if not table.rows:
            continue
        header = " ".join(cell.text.replace("\n", " ") for cell in table.rows[0].cells)
        required = ("STT", "Họ tên trẻ", "Mã trẻ", "Loại can thiệp", "Số buổi")
        if not all(item in header for item in required) or len(table.rows) < 2:
            continue
        row_template = table.rows[1]
        _remove_row(row_template)
        base_context = _annex_context(hop_dong, rows)
        for index, item in enumerate(rows, start=1):
            row = _clone_table_row(table, row_template)
            _replace_row(
                row,
                {
                    **base_context,
                    "STT": index,
                    "HoTenTre": _row_value(item, "ten_tre"),
                    "MaTre": _row_value(item, "ma_tre"),
                    "DichVu": _service_display(_row_value(item, "loai_dich_vu")),
                    "SoBuoi": _row_value(item, "so_buoi_du_kien"),
                },
            )
        return


def _build_annex_document(hop_dong, rows):
    document = _document(ANNEX_TEMPLATE)
    _fill_assignment_table(document, hop_dong, rows)
    _replace_document(document, _annex_context(hop_dong, rows))
    return document


def _append_document(target, source):
    target.add_page_break()
    target_body = target.element.body
    for child in source.element.body.iterchildren():
        if child.tag.endswith("sectPr"):
            continue
        target_body.append(deepcopy(child))


def _snapshot_rows(hop_dong):
    snapshot = (
        ChiTietPhuLucPhanCong.objects.filter(phu_luc__hop_dong=hop_dong, phu_luc__loai_phu_luc="KY_1")
        .order_by("id")
    )
    if snapshot.exists():
        return list(snapshot)
    return list(
        PhanCongTre.objects.filter(phan_bo=hop_dong.de_xuat.phan_bo, ky_phan_cong=1)
        .select_related("tre")
        .order_by("id")
    )


def _row_value(row, name):
    if hasattr(row, name):
        return getattr(row, name)
    if name == "ma_tre":
        return row.tre.ma_tre
    if name == "ten_tre":
        return row.tre.ho_ten
    return None


def export_contract_bundle(hop_dong):
    if not CONTRACT_TEMPLATE.exists() or not ANNEX_TEMPLATE.exists():
        raise ValidationError("Chưa có đủ template Word hợp đồng và Phụ lục 2.")

    rows = _snapshot_rows(hop_dong)
    contract = _document(CONTRACT_TEMPLATE)
    _fill_assignment_table(contract, hop_dong, rows)
    _replace_document(contract, {**_allocation_context(hop_dong), **_annex_context(hop_dong, rows)})

    output = BytesIO()
    contract.save(output)
    output.seek(0)
    return output


def create_contract_from_proposal(proposal, cleaned_data):
    """Tạo HĐ từ proposal; tuyệt đối không lấy tiền/khối lượng từ client."""
    if proposal.trang_thai != "DA_DUYET":
        raise ValidationError("Chỉ được tạo hợp đồng từ đề xuất đã duyệt.")

    allocation = proposal.phan_bo
    with transaction.atomic():
        hop_dong = HopDong.objects.create(
            de_xuat=proposal,
            can_bo=allocation.can_bo,
            nhom_hd=allocation.nhom_hd,
            so_hop_dong=cleaned_data["so_hop_dong"],
            ngay_ky=cleaned_data.get("ngay_ky"),
            tu_ngay=cleaned_data["tu_ngay"],
            den_ngay=cleaned_data["den_ngay"],
            don_gia_cong=FinancialConfig.DON_GIA_CONG,
            dinh_muc_di_lai_phcn=allocation.dinh_muc_di_lai_phcn,
            dinh_muc_di_lai_cs=allocation.dinh_muc_di_lai_cs,
            gia_tri_hop_dong=proposal.gia_tri_du_kien,
            trang_thai=cleaned_data.get("trang_thai") or "DU_THAO",
            ghi_chu=cleaned_data.get("ghi_chu"),
        )
        hop_dong.full_clean()
        hop_dong.save()

        for service, count, sessions, travel in (
            ("VLTL", allocation.so_tre_phcn, allocation.so_buoi_phcn, allocation.dinh_muc_di_lai_phcn),
            ("CSXH", allocation.so_tre_cs, allocation.so_buoi_cs, allocation.dinh_muc_di_lai_cs),
        ):
            if count or sessions:
                ChiTietKhoiLuongHopDong.objects.create(
                    hop_dong=hop_dong,
                    loai_dich_vu=service,
                    so_tre=count,
                    so_buoi=sessions,
                    don_gia_cong=hop_dong.don_gia_cong,
                    dinh_muc_di_lai=travel,
                )

        phu_luc = PhuLucHopDong.objects.create(
            hop_dong=hop_dong,
            loai_phu_luc="KY_1",
            so_phu_luc=f"PL-K1-{hop_dong.so_hop_dong}",
            ngay_lap=timezone.localdate(),
        )
        for item in allocation.danh_sach_phan_cong.select_related("tre", "phan_bo__can_bo").filter(ky_phan_cong=1):
            ChiTietPhuLucPhanCong.objects.create(
                phu_luc=phu_luc,
                source_phan_cong=item,
                ma_tre=item.tre.ma_tre,
                ten_tre=item.tre.ho_ten,
                ma_can_bo=item.phan_bo.can_bo.ma_can_bo,
                ten_can_bo=item.phan_bo.can_bo.ho_ten,
                loai_dich_vu=item.loai_dich_vu,
                dot_phan_cong=item.dot_phan_cong,
                ky_phan_cong=item.ky_phan_cong,
                ngay_phan_cong=item.ngay_phan_cong,
                so_buoi_du_kien=item.so_buoi_du_kien,
                dinh_muc_di_lai=item.dinh_muc_di_lai,
                dia_diem_ct=item.dia_diem_ct,
                hinh_thuc_ct=item.hinh_thuc_ct,
                ghi_chu=item.ghi_chu,
            )
        proposal.trang_thai = "DA_TAO_HOP_DONG"
        proposal.save(update_fields=["trang_thai", "updated_at"])
    return hop_dong
