from copy import copy
from collections import OrderedDict
from decimal import Decimal
from functools import lru_cache
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.core.exceptions import ValidationError

TEMPLATE_ROOT = Path(settings.BASE_DIR) / "quanly" / "document_templates"
INTERVENTION_PAYMENT_TEMPLATE = TEMPLATE_ROOT / "thanh_toan_cong_can_thiep" / "Mau_DNTT.xlsx"
INTERVENTION_ACCOUNT_TEMPLATE = TEMPLATE_ROOT / "thanh_toan_cong_can_thiep" / "Mau_DSTK.xlsx"
INTERVENTION_COMMITMENT_TEMPLATE = TEMPLATE_ROOT / "thanh_toan_cong_can_thiep" / "Mau_DNCK.xlsx"
PARENT_TRAVEL_TEMPLATE_ROOT = TEMPLATE_ROOT / "thanh_toan_di_lai_phu_huynh"
FIRST_DETAIL_ROW = 12
LAST_TEMPLATE_DETAIL_ROW = 50
TOTAL_ROW = 51


def _workbook(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError("Thiếu thư viện openpyxl để xuất Excel.") from exc
    return load_workbook(path)


def _location_bucket(value):
    text = (value or "").strip().lower()
    if "nhà" in text or "nha" in text:
        return "home"
    if "trường" in text or "truong" in text or "cơ sở" in text or "co so" in text:
        return "facility"
    return "other"


def _copy_row_style(ws, source_row, target_row):
    for source, target in zip(ws[source_row], ws[target_row]):
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.protection = copy(source.protection)
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _clear_detail_rows(ws):
    for row in range(FIRST_DETAIL_ROW, LAST_TEMPLATE_DETAIL_ROW + 1):
        for cell in ws[row]:
            cell.value = None


def _fill_placeholders(workbook, context):
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for key, value in context.items():
                        cell.value = cell.value.replace("{{" + key + "}}", str(value or ""))


@lru_cache(maxsize=64)
def _location_for_group(group_code):
    workbook = _workbook(INTERVENTION_PAYMENT_TEMPLATE)
    try:
        if "Dia diem thuc hien" not in workbook.sheetnames:
            return "Theo nhật ký thực hiện"
        ws = workbook["Dia diem thuc hien"]
        for code, location in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            if str(code or "").strip().replace(".0", "") == str(group_code or "").strip():
                return location or "Theo nhật ký thực hiện"
    finally:
        workbook.close()
    return "Theo nhật ký thực hiện"


def export_intervention_payment_request(dot):
    """Xuất đề nghị thanh toán công cán bộ theo một đợt thanh toán."""
    if not INTERVENTION_PAYMENT_TEMPLATE.exists():
        raise ValidationError("Chưa có template đề nghị thanh toán công cán bộ.")
    rows = list(dot.chi_tiet.select_related("nhat_ky__phan_cong__tre", "nhat_ky__hop_dong").order_by("id"))
    if not rows:
        raise ValidationError("Đợt thanh toán chưa có chi tiết để xuất.")

    workbook = _workbook(INTERVENTION_PAYMENT_TEMPLATE)
    ws = workbook["DNTT"]
    _clear_detail_rows(ws)
    context = {
        "KyThanhToan": f"{dot.thang}/{dot.nam}",
        "TuNgay": dot.hop_dong.tu_ngay.strftime("%d/%m/%Y"),
        "DenNgay": dot.hop_dong.den_ngay.strftime("%d/%m/%Y"),
        "DiaDiemThucHien": "Theo nhật ký thực hiện",
    }

    for index, item in enumerate(rows):
        row_number = FIRST_DETAIL_ROW + index
        if row_number > LAST_TEMPLATE_DETAIL_ROW:
            ws.insert_rows(row_number)
            _copy_row_style(ws, FIRST_DETAIL_ROW, row_number)
        journal = item.nhat_ky
        assignment = journal.phan_cong
        bucket = _location_bucket(assignment.dia_diem_ct)
        planned = assignment.so_buoi_du_kien
        actual = item.so_buoi_thanh_toan
        travel = item.so_luot_di_lai
        values = {
            "A": index + 1,
            "B": f"{journal.hop_dong.can_bo.ho_ten} - {assignment.tre.ho_ten}",
            "C": journal.hop_dong.so_hop_dong,
            "D": assignment.nhom_dich_vu,
            "E": planned if bucket == "facility" else 0,
            "F": planned if bucket == "home" else 0,
            "G": planned if bucket == "other" else 0,
            "H": travel if bucket == "home" else 0,
            "I": journal.don_gia_cong,
            "J": journal.dinh_muc_di_lai,
            "K": f"=E{row_number}*I{row_number}+H{row_number}*J{row_number}",
            "L": actual if bucket == "facility" else 0,
            "M": actual if bucket == "home" else 0,
            "N": actual if bucket == "other" else 0,
            "O": travel,
            "P": f"=(L{row_number}+M{row_number}+N{row_number})*I{row_number}",
            "Q": f"=O{row_number}*J{row_number}",
            "R": f"=P{row_number}+Q{row_number}",
            "S": f"=IF(P{row_number}>=5000000,P{row_number}*10%,0)",
            "T": f"=R{row_number}-S{row_number}",
            "U": item.ghi_chu or "",
        }
        for column, value in values.items():
            ws[f"{column}{row_number}"] = value

    end_row = FIRST_DETAIL_ROW + len(rows) - 1
    for column in ("K", "P", "Q", "R", "S", "T"):
        ws[f"{column}{TOTAL_ROW}"] = f"=SUM({column}{FIRST_DETAIL_ROW}:{column}{end_row})"
    _fill_placeholders(workbook, context)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def export_intervention_account_list(dot):
    """Xuất danh sách tài khoản và số tiền thực lĩnh của một đợt thanh toán."""
    if not INTERVENTION_ACCOUNT_TEMPLATE.exists():
        raise ValidationError("Chưa có template danh sách tài khoản thanh toán.")
    rows = list(dot.chi_tiet.select_related("nhat_ky__hop_dong__can_bo").order_by("id"))
    if not rows:
        raise ValidationError("Đợt thanh toán chưa có chi tiết để xuất danh sách tài khoản.")

    workbook = _workbook(INTERVENTION_ACCOUNT_TEMPLATE)
    ws = workbook["DSTK"]
    for row in range(5, 20):
        for cell in ws[row]:
            cell.value = None

    staff = dot.hop_dong.can_bo
    gross = sum(item.thanh_tien for item in rows)
    net = sum(item.thuc_linh for item in rows)
    values = {
        "A5": 1,
        "B5": staff.ho_ten,
        "C5": staff.dien_thoai or "",
        "D5": staff.email or "",
        "E5": staff.dia_chi or "",
        "F5": staff.tai_khoan or "",
        "G5": staff.ngan_hang or "",
        "H5": staff.chi_nhanh or "",
        "I5": gross,
        "J5": net,
        "K5": staff.mst or "",
        "L5": staff.cccd or "",
        "I20": "=SUM(I5:I19)",
        "J20": "=SUM(J5:J19)",
    }
    for coordinate, value in values.items():
        ws[coordinate] = value
    _fill_placeholders(workbook, {"DiaDiemThucHien": "Theo hồ sơ thanh toán"})
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


PARENT_TRAVEL_CATEGORIES = {"CG", "NCS", "CBDA"}


def _parent_template(category, prefix):
    category = (category or "NCS").upper()
    if category not in PARENT_TRAVEL_CATEGORIES:
        raise ValidationError("Nhóm mẫu phải là CG, NCS hoặc CBDA.")
    path = PARENT_TRAVEL_TEMPLATE_ROOT / f"Mau_{prefix}_DiLai_{category}.xlsx"
    if not path.exists():
        raise ValidationError(f"Chưa có template {path.name}.")
    return path


def export_parent_travel_payment_request(dot, category="NCS"):
    rows = list(dot.chi_tiet.select_related("nhat_ky__hop_dong__can_bo", "nhat_ky__phan_cong__tre").order_by("id"))
    if not rows:
        raise ValidationError("Đợt thanh toán chưa có chi tiết để xuất.")
    workbook = _workbook(_parent_template(category, "DNTT"))
    ws = workbook["DNTT_NCS"]
    first_row, last_row = 11, ws.max_row - 8
    for row in range(first_row, last_row + 1):
        for cell in ws[row]:
            cell.value = None
    for index, item in enumerate(rows, start=1):
        row = first_row + index - 1
        if row > last_row:
            raise ValidationError("Số dòng thanh toán vượt giới hạn template.")
        journal = item.nhat_ky
        assignment = journal.phan_cong
        child = assignment.tre
        ws.cell(row, 1).value = index
        ws.cell(row, 2).value = f"{journal.hop_dong.can_bo.ho_ten} - {child.ho_ten}"
        ws.cell(row, 3).value = child.ten_phu_huynh or ""
        ws.cell(row, 4).value = journal.hop_dong.so_hop_dong
        ws.cell(row, 5).value = assignment.nhom_dich_vu
        ws.cell(row, 6).value = assignment.so_buoi_du_kien
        ws.cell(row, 7).value = item.dinh_muc_di_lai
        ws.cell(row, 8).value = f"=F{row}*G{row}"
        ws.cell(row, 9).value = item.so_luot_di_lai
        ws.cell(row, 10).value = item.so_luot_di_lai
        ws.cell(row, 11).value = f"=J{row}*G{row}"
        ws.cell(row, 12).value = item.ghi_chu or ""
    _fill_placeholders(workbook, {"KyThanhToan": f"{dot.thang}/{dot.nam}", "TuNgay": dot.hop_dong.tu_ngay.strftime("%d/%m/%Y"), "DenNgay": dot.hop_dong.den_ngay.strftime("%d/%m/%Y")})
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def export_parent_travel_account_list(dot, category="NCS"):
    rows = list(dot.chi_tiet.select_related("nhat_ky__phan_cong__tre").order_by("id"))
    if not rows:
        raise ValidationError("Đợt thanh toán chưa có chi tiết để xuất danh sách tài khoản.")
    workbook = _workbook(_parent_template(category, "DSTK"))
    ws = workbook["DSTK_NCS"]
    first_row, last_row = 5, ws.max_row - 5
    for row in range(first_row, last_row + 1):
        for cell in ws[row]:
            cell.value = None
    grouped = {}
    for item in rows:
        child = item.nhat_ky.phan_cong.tre
        key = (child.ten_phu_huynh or "", child.ten_tai_khoan or "", child.tai_khoan or "", child.ngan_hang or "", child.chi_nhanh or "")
        grouped[key] = grouped.get(key, 0) + item.thanh_tien
    for index, (key, amount) in enumerate(grouped.items(), start=1):
        row = first_row + index - 1
        if row > last_row:
            raise ValidationError("Số người nhận vượt giới hạn template.")
        parent, account_name, account, bank, branch = key
        for column, value in enumerate((index, parent, account, bank, branch, amount), start=1):
            ws.cell(row, column).value = value
    _fill_placeholders(workbook, {"DiaDiemThucHien": "Theo hồ sơ thanh toán"})
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _journal_totals(journals):
    labor = sum((Decimal(row.so_buoi_thuc_hien) * Decimal(row.don_gia_cong) for row in journals), Decimal("0"))
    travel = sum((Decimal(row.so_luot_di_lai_cbct) * Decimal(row.dinh_muc_di_lai) for row in journals), Decimal("0"))
    return labor, travel


def _group_journal_payment_rows(journals):
    """Gom nhật ký theo CBCT, sau đó theo đúng dòng phân công của trẻ.

    Một dòng nhật ký là một phiên can thiệp. ĐNTT chỉ hiển thị một dòng cho
    mỗi phân công (trẻ + dịch vụ + đợt phân công), và một dòng tổng cho CBCT.
    """
    grouped = OrderedDict()
    for journal in journals:
        staff = getattr(journal, "can_bo_hieu_luc", None) or getattr(getattr(journal, "hop_dong", None), "can_bo", None)
        if not staff:
            continue
        staff_key = staff.pk
        if staff_key not in grouped:
            grouped[staff_key] = {
                "staff": staff,
                "contracts": OrderedDict(),
                "details": OrderedDict(),
            }
        staff_group = grouped[staff_key]
        contract = getattr(journal, "hop_dong_hieu_luc", None) or getattr(journal, "hop_dong", None)
        if contract:
            staff_group["contracts"][contract.pk] = contract
        detail_key = journal.phan_cong_id
        if detail_key not in staff_group["details"]:
            assignment = journal.phan_cong
            bucket = _location_bucket(assignment.dia_diem_ct or journal.dia_diem_ct)
            planned = int(assignment.so_buoi_du_kien or 0)
            staff_group["details"][detail_key] = {
                "assignment": assignment,
                "child": assignment.tre,
                "service": assignment.loai_dich_vu,
                "planned_facility": planned if bucket == "facility" else 0,
                "planned_home": planned if bucket == "home" else 0,
                "planned_other": planned if bucket == "other" else 0,
                "planned_travel": planned if bucket in {"home", "other"} else 0,
                "actual_facility": 0,
                "actual_home": 0,
                "actual_other": 0,
                "actual_travel": 0,
                "labor": Decimal("0"),
                "travel": Decimal("0"),
                "labor_rate": Decimal(journal.don_gia_cong),
                "travel_rate": Decimal(journal.dinh_muc_di_lai),
                "notes": [],
            }
        detail = staff_group["details"][detail_key]
        bucket = _location_bucket(journal.dia_diem_ct or journal.phan_cong.dia_diem_ct)
        detail[f"actual_{bucket}"] += int(journal.so_buoi_thuc_hien or 0)
        detail["actual_travel"] += int(journal.so_luot_di_lai_cbct or 0)
        detail["labor"] += Decimal(journal.so_buoi_thuc_hien or 0) * Decimal(journal.don_gia_cong)
        detail["travel"] += Decimal(journal.so_luot_di_lai_cbct or 0) * Decimal(journal.dinh_muc_di_lai)
        if journal.ghi_chu and journal.ghi_chu not in detail["notes"]:
            detail["notes"].append(journal.ghi_chu)

    for staff_group in grouped.values():
        details = list(staff_group["details"].values())
        staff_group["details"] = details
        staff_group["planned_facility"] = sum(x["planned_facility"] for x in details)
        staff_group["planned_home"] = sum(x["planned_home"] for x in details)
        staff_group["planned_other"] = sum(x["planned_other"] for x in details)
        staff_group["planned_travel"] = sum(x["planned_travel"] for x in details)
        staff_group["actual_facility"] = sum(x["actual_facility"] for x in details)
        staff_group["actual_home"] = sum(x["actual_home"] for x in details)
        staff_group["actual_other"] = sum(x["actual_other"] for x in details)
        staff_group["actual_travel"] = sum(x["actual_travel"] for x in details)
        staff_group["labor"] = sum((x["labor"] for x in details), Decimal("0"))
        staff_group["travel"] = sum((x["travel"] for x in details), Decimal("0"))
    return list(grouped.values())


def _selection_context(journals, ky="", thang="", nam=""):
    first = journals[0]
    dates = [item.ngay_thuc_hien for item in journals if item.ngay_thuc_hien]
    group = first.nhom_hd_hieu_luc
    location = _location_for_group(group.ma_nhom_hd)
    return {
        "KyThanhToan": ky or first.ky_can_thiep,
        "TuNgay": min(dates).strftime("%d/%m/%Y") if dates else "",
        "DenNgay": max(dates).strftime("%d/%m/%Y") if dates else "",
        "DiaDiemThucHien": location,
        "NhomHD": group.ma_nhom_hd,
        "Thang": thang,
        "Nam": nam,
    }


def export_journal_account_list(journals):
    journals = list(journals)
    if not journals:
        raise ValidationError("Không có nhật ký phù hợp để xuất DSTK.")
    workbook = _workbook(INTERVENTION_ACCOUNT_TEMPLATE)
    ws = workbook["DSTK"]
    for row in range(5, 20):
        for cell in ws[row]:
            cell.value = None
    grouped = {}
    for journal in journals:
        staff = getattr(journal, "can_bo_hieu_luc", None) or getattr(getattr(journal, "hop_dong", None), "can_bo", None)
        if staff:
            grouped.setdefault(staff.pk, (staff, []))[1].append(journal)
    if len(grouped) > 15:
        raise ValidationError("DSTK hiện hỗ trợ tối đa 15 CBCT trong một Nhóm HĐ/kỳ.")
    total_net = Decimal("0")
    for index, (staff, items) in enumerate(grouped.values(), start=1):
        row = 4 + index
        labor, travel = _journal_totals(items)
        from .financial import calculate_payment_breakdown
        net = calculate_payment_breakdown(labor, travel)["thuc_linh"]
        total_net += net
        for column, value in enumerate((index, staff.ho_ten, staff.dien_thoai or "", staff.email or "", staff.dia_chi or "", staff.tai_khoan or "", staff.ngan_hang or "", staff.chi_nhanh or "", labor + travel, net, staff.mst or "", staff.cccd or ""), start=1):
            ws.cell(row, column).value = value
    ws["J20"] = total_net
    _fill_placeholders(workbook, _selection_context(journals))
    if "DataStaff" in workbook.sheetnames:
        workbook["DataStaff"].sheet_state = "hidden"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def export_journal_payment_request_excel(journals, ky="", thang="", nam=""):
    journals = list(journals)
    if not journals:
        raise ValidationError("Không có nhật ký phù hợp để xuất ĐNTT Excel.")
    workbook = _workbook(INTERVENTION_PAYMENT_TEMPLATE)
    ws = workbook["DNTT"]
    _clear_detail_rows(ws)
    staff_groups = _group_journal_payment_rows(journals)
    output_rows = sum(1 + len(group["details"]) for group in staff_groups)
    total_row = TOTAL_ROW
    extra_rows = max(0, output_rows - (LAST_TEMPLATE_DETAIL_ROW - FIRST_DETAIL_ROW + 1))
    for _ in range(extra_rows):
        ws.insert_rows(total_row)
        _copy_row_style(ws, LAST_TEMPLATE_DETAIL_ROW, total_row)
        total_row += 1
    from openpyxl.styles import PatternFill
    from .financial import calculate_payment_breakdown

    summary_fill = PatternFill(fill_type="solid", fgColor="FFF200")
    row = FIRST_DETAIL_ROW
    for index, group in enumerate(staff_groups, start=1):
        contracts = list(group["contracts"].values())
        labor_rate = group["details"][0]["labor_rate"] if group["details"] else Decimal("0")
        travel_rate = group["details"][0]["travel_rate"] if group["details"] else Decimal("0")
        planned_total = sum(
            (
                Decimal(detail["planned_facility"] + detail["planned_home"] + detail["planned_other"]) * detail["labor_rate"]
                + Decimal(detail["planned_travel"]) * detail["travel_rate"]
                for detail in group["details"]
            ),
            Decimal("0"),
        )
        breakdown = calculate_payment_breakdown(group["labor"], group["travel"])
        summary_values = (
            index,
            group["staff"].ho_ten,
            ", ".join(item.so_hop_dong for item in contracts),
            "",
            group["planned_facility"], group["planned_home"], group["planned_other"], group["planned_travel"],
            labor_rate, travel_rate, planned_total,
            group["actual_facility"], group["actual_home"], group["actual_other"], group["actual_travel"],
            group["labor"], group["travel"], breakdown["tong_truoc_thue"], breakdown["thue_tncn"], breakdown["thuc_linh"], "",
        )
        for column, value in enumerate(summary_values, start=1):
            ws.cell(row, column).value = value
            ws.cell(row, column).fill = copy(summary_fill)
            font = copy(ws.cell(row, column).font)
            font.bold = True
            ws.cell(row, column).font = font
        row += 1
        for detail in group["details"]:
            planned_total = (
                Decimal(detail["planned_facility"] + detail["planned_home"] + detail["planned_other"]) * detail["labor_rate"]
                + Decimal(detail["planned_travel"]) * detail["travel_rate"]
            )
            gross = detail["labor"] + detail["travel"]
            detail_values = (
                "", detail["child"].ho_ten, "", detail["service"],
                detail["planned_facility"], detail["planned_home"], detail["planned_other"], detail["planned_travel"],
                detail["labor_rate"], detail["travel_rate"], planned_total,
                detail["actual_facility"], detail["actual_home"], detail["actual_other"], detail["actual_travel"],
                detail["labor"], detail["travel"], gross, "", gross, "; ".join(detail["notes"]),
            )
            for column, value in enumerate(detail_values, start=1):
                ws.cell(row, column).value = value
            row += 1

    totals = {
        "K": sum((
            Decimal(detail["planned_facility"] + detail["planned_home"] + detail["planned_other"]) * detail["labor_rate"]
            + Decimal(detail["planned_travel"]) * detail["travel_rate"]
            for group in staff_groups for detail in group["details"]
        ), Decimal("0")),
        "P": sum((group["labor"] for group in staff_groups), Decimal("0")),
        "Q": sum((group["travel"] for group in staff_groups), Decimal("0")),
    }
    totals["R"] = totals["P"] + totals["Q"]
    totals["S"] = sum((calculate_payment_breakdown(group["labor"], group["travel"])["thue_tncn"] for group in staff_groups), Decimal("0"))
    totals["T"] = totals["R"] - totals["S"]
    for column, value in totals.items():
        ws[f"{column}{total_row}"] = value
    context = _selection_context(journals, ky=ky, thang=thang, nam=nam)
    _fill_placeholders(workbook, context)
    ws.print_title_rows = "8:11"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    for helper_name in ("Dia diem thuc hien", "DataStaff"):
        if helper_name in workbook.sheetnames:
            workbook[helper_name].sheet_state = "hidden"
    output = BytesIO(); workbook.save(output); output.seek(0); return output


def export_journal_commitment(journals, tu_ngay="", den_ngay=""):
    journals = list(journals)
    if not journals:
        raise ValidationError("Không có nhật ký phù hợp để xuất ĐNCK.")
    workbook = _workbook(INTERVENTION_COMMITMENT_TEMPLATE)
    ws = workbook["DNCK"]
    for row in range(5, 26):
        for cell in ws[row]:
            try:
                cell.value = None
            except AttributeError:
                # Một số template gộp ô ở khu vực cuối bảng.
                pass
    grouped = {}
    for journal in journals:
        staff = getattr(journal, "can_bo_hieu_luc", None) or getattr(getattr(journal, "hop_dong", None), "can_bo", None)
        if staff:
            grouped.setdefault(staff.pk, (staff, []))[1].append(journal)
    if len(grouped) > 14:
        raise ValidationError("ĐNCK hiện hỗ trợ tối đa 14 CBCT trong một Nhóm HĐ/kỳ.")
    total_net = Decimal("0")
    for index, (staff, items) in enumerate(grouped.values(), start=1):
        row = 4 + index
        labor, travel = _journal_totals(items)
        from .financial import calculate_payment_breakdown
        net = calculate_payment_breakdown(labor, travel)["thuc_linh"]
        total_net += net
        for column, value in enumerate((index, staff.ho_ten, staff.dia_chi or "", staff.tai_khoan or "", staff.ngan_hang or "", staff.chi_nhanh or "", net), start=1):
            ws.cell(row, column).value = value
    ws["G19"] = total_net
    _fill_placeholders(workbook, {"TuNgay": tu_ngay, "DenNgay": den_ngay})
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
